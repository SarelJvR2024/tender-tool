"""
Company-owned letters and documents: upload your own, or generate it from
your own guideline. One mechanism for every letter-type document a tender
needs, starting with the letter of authority to sign.

Sarel, 2026-10-17 (Hardcoded-Assumptions-Audit Q7): "each company must have
the option to upload their own docs or templates, or provide a guideline of
how AI engine must generate their docs - not just this letter but all docs and
letters that might be needed."

Before this, fill_authority_letter.py wrote Klipcorp's own historical
"Letter of Authority to Sign" wording for EVERY company. Now each document
type has a per-company choice (app/supabase/migrations/0037: table
company_document_templates), the same two-way choice the CV engine already
offers on upload (cv_uploads.mode 'as_is' | 'ai_extract'):

  uploaded      The company's own real document. Used AS-IS: the exact file,
                byte for byte, every time this document is needed. No AI, no
                re-typing, no reformatting (the pack converts a copy to PDF only
                to bind it in, as it does for an as-is CV).
  ai_guideline  The company's guideline: ITS OWN wording for this document,
                with {{field}} marks where the tender's details go (e.g.
                "{{signatory_name}} is authorised to sign for
                {{tender_reference}}"). The generated document is that text,
                verbatim, with the marked fields filled from the tender/company
                records. Same discipline as companies.how_we_work and the
                spec's "No AI-drafted narrative sections" rule: nothing is
                paraphrased, "improved" or padded with generic wording. A field
                with no value on file prints as a visible [NOT ON FILE: ...]
                and is logged `missing`.
  (no row)      NOT SET UP. Nothing is produced - the pack lists the document
                as not included, with "set it up on the Company page". It never
                falls back to another company's wording.

Optional AI placement. A guideline written without {{field}} marks can still
need the tender's details placed into it. When an LLM client is configured
(ANTHROPIC_API_KEY on the worker), the model is asked to place the given
facts into the company's text and change nothing else - and its output is
then CHECKED, not trusted: every word must come from the guideline or from a
supplied fact (guard_only_company_wording). Output that adds any other word is
rejected and the guideline is used verbatim instead, with the reason logged.
Without a key, a field-less guideline is reproduced verbatim and logged
assumed_default ("no fields marked - check the tender's details are in it").

Each document type's fields are listed in DOCUMENT_TYPES; the app's copy of
the list (labels, help text) is app/lib/companyDocuments.ts - keep in step.
"""
import datetime
import hashlib
import os
import re
import shutil

DOCUMENT_TYPES = {
    "authority_to_sign": {
        "label": "Letter of authority to sign",
        "used_for": "SANRAL Form A2 (Certificate of Authority for Signatory) and any buyer that asks for a "
                    "board resolution or letter authorising the signatory",
        "fields": {
            "company_legal_name": "Company legal name",
            "signatory_name": "Signatory's name",
            "signatory_capacity": "Signatory's capacity / position",
            "director_name": "Director's name",
            "tender_reference": "Tender / RFP reference",
            "tender_description": "Tender description",
            "contract_number": "Contract number",
            "nra_reference": "SANRAL NRA reference",
            "closing_date": "Closing date",
            "date": "Today's date",
        },
    },
}

FIELD_RE = re.compile(r"\{\{\s*([A-Za-z_][A-Za-z0-9_]*)\s*\}\}")
MARKER_RE = re.compile(r"\[(?:NOT ON FILE|UNKNOWN FIELD): [^\]]*\]")
WORD_RE = re.compile(r"[0-9A-Za-zÀ-ɏ]+")


def label(document_type):
    return DOCUMENT_TYPES.get(document_type, {}).get("label") or document_type.replace("_", " ")


def not_set_up_reason(document_type, company):
    return (f"{label(document_type)} not set up for {company or 'this company'} yet - on the Company page, "
            f"upload your own {label(document_type).lower()} (used as-is) or write the guideline it should be "
            "generated from. Nothing was generated in its place.")


def fill_fields(guideline, facts, document_type, log=None):
    """The guideline verbatim, each {{field}} replaced by its fact. Returns
    (text, missing_field_names, unknown_field_names)."""
    known = DOCUMENT_TYPES.get(document_type, {}).get("fields", {})
    missing, unknown = [], []

    def sub(m):
        name = m.group(1).lower()
        if name not in known:
            unknown.append(name)
            return f"[UNKNOWN FIELD: {m.group(1)}]"
        v = facts.get(name)
        if v in (None, ""):
            missing.append(name)
            return f"[NOT ON FILE: {known[name]}]"
        if log is not None:
            log.confirmed(f"{document_type}.{name}", v)
        return str(v)

    return FIELD_RE.sub(sub, guideline), missing, unknown


def _words(text):
    return {w.casefold() for w in WORD_RE.findall(text or "")}


def guard_only_company_wording(output, guideline, facts):
    """Words in `output` that are neither in the company's guideline nor in a
    supplied fact value - i.e. anything an AI step invented. Empty = clean.
    The [NOT ON FILE: ...] / [UNKNOWN FIELD: ...] markers and the field labels
    they carry are allowed (they are the tool's own visible flags)."""
    allowed = _words(FIELD_RE.sub(" ", guideline))
    for v in facts.values():
        allowed |= _words(str(v)) if v not in (None, "") else set()
    return sorted(_words(MARKER_RE.sub(" ", output)) - allowed)


LLM_SYSTEM = (
    "You place facts into a company's own document wording. The user gives you the company's text and a list of "
    "facts. Return the company's text with the facts inserted where the text plainly refers to them. Rules: use the "
    "company's words exactly; do not add, remove, rephrase, summarise or polish any wording; do not add greetings, "
    "headings, sign-offs or explanations; if a fact has no obvious place, leave it out. Output only the document text."
)


def llm_place_facts(guideline, facts, client, model=None):
    """One Messages API call; returns the model's text. `client` is anything
    with .messages.create(...) (anthropic.Anthropic(), or a test stub)."""
    model = model or os.environ.get("DOCUMENT_GUIDELINE_MODEL") or "claude-opus-5-5"
    fact_lines = "\n".join(f"- {k}: {v}" for k, v in facts.items() if v not in (None, ""))
    resp = client.messages.create(
        model=model, max_tokens=4000, system=LLM_SYSTEM,
        messages=[{"role": "user", "content": f"<company_text>\n{guideline}\n</company_text>\n\n<facts>\n{fact_lines}\n</facts>"}],
    )
    return "".join(getattr(b, "text", "") for b in resp.content).strip()


def default_llm_client():
    """anthropic.Anthropic() when ANTHROPIC_API_KEY is set, else None (the
    deterministic path still works without it)."""
    if not os.environ.get("ANTHROPIC_API_KEY"):
        return None
    try:
        import anthropic
    except ImportError:
        return None
    # Explicit, bounded timeout - see extract_cv_fields.default_llm_client's
    # own comment (UX audit, 2026-10-23) for why: the SDK's 600s default left
    # a single hung connection able to wedge a job far longer than it should.
    return anthropic.Anthropic(timeout=180)


def render_guideline(guideline, facts, document_type, log, llm=None):
    """Text of the document generated from the company's guideline. Returns
    (text, method) - method 'fields' | 'ai_placed' | 'verbatim'."""
    if FIELD_RE.search(guideline):
        text, missing, unknown = fill_fields(guideline, facts, document_type, log)
        for name in missing:
            log.missing(f"{document_type}.{name}",
                        f"{label(document_type)}: your guideline asks for {DOCUMENT_TYPES[document_type]['fields'][name]}, "
                        "which isn't on file for this tender - printed as [NOT ON FILE]; fill it in before signing.")
        for name in unknown:
            log.missing(f"{document_type}.{name}",
                        f"{label(document_type)}: your guideline uses {{{{{name}}}}}, which isn't a field Tender Tool "
                        f"knows ({', '.join(DOCUMENT_TYPES[document_type]['fields'])}) - printed as [UNKNOWN FIELD]; "
                        "correct the guideline on the Company page.")
        method = "fields"
    elif llm is not None:
        try:
            text = llm_place_facts(guideline, facts, llm)
        except Exception as e:  # an AI failure must not lose the document
            text = None
            log.assumed_default(f"{document_type}.ai_placement", "guideline used verbatim",
                                f"{label(document_type)}: the AI step failed ({type(e).__name__}) - your guideline "
                                "was used word for word; check the tender's details are in it before signing.")
        if text is not None:
            extra = guard_only_company_wording(text, guideline, facts)
            if extra or not text:
                log.assumed_default(f"{document_type}.ai_placement", "guideline used verbatim",
                                    f"{label(document_type)}: the AI step added wording that isn't in your guideline "
                                    f"({', '.join(extra[:12]) or 'empty output'}) - rejected; your guideline was used "
                                    "word for word instead. Mark {{fields}} in it so the tender's details are placed exactly.")
                text = None
            else:
                log.computed(f"{document_type}.ai_placement", "facts placed by AI",
                             "AI placed the tender's details into your own wording; every word was checked against "
                             "your guideline and the supplied facts.")
                return text, "ai_placed"
        text = guideline
        method = "verbatim"
    else:
        text, method = guideline, "verbatim"
        log.assumed_default(f"{document_type}.fields", "guideline used verbatim",
                            f"{label(document_type)}: your guideline has no {{{{fields}}}} marked, so it was used word "
                            "for word and no tender details were placed in it - check it is right for this tender, or "
                            "mark fields like {{signatory_name}} and {{tender_reference}} on the Company page.")
    return text, method


def write_docx(text, out_path, logo_path=None):
    """The text as a plain document on the company's letterhead (logo only).
    Blank lines separate paragraphs; a line starting '# ' is a bold, centred
    heading; nothing else is added - no heading, greeting or sign-off of our own."""
    import docx
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    d = docx.Document()
    if logo_path and os.path.exists(logo_path):
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(logo_path, width=Inches(1.8))
        d.add_paragraph()
    for block in re.split(r"\n\s*\n", text.replace("\r\n", "\n").strip()):
        if block.startswith("# ") and "\n" not in block:
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(block[2:].strip())
            r.bold = True
            r.font.size = Pt(13)
            continue
        p = d.add_paragraph()
        for i, line in enumerate(block.split("\n")):
            if i:
                p.add_run().add_break()
            p.add_run(line)
    d.save(out_path)
    return out_path


def _sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 16), b""):
            h.update(chunk)
    return h.hexdigest()


def produce(document_type, template, facts, out_dir, log, company=None, logo_path=None, llm=None):
    """One company's document of `document_type`.

    template: None (not set up), or {"mode": "uploaded", "path": local file,
              "original_filename"?} or {"mode": "ai_guideline", "guideline": text}
              (process_pack_jobs.load_document_templates builds this from the
              company's company_document_templates row).
    facts:    {field: value} for this tender - see DOCUMENT_TYPES[...]["fields"].
    Returns {"status": "uploaded"|"generated"|"not_set_up"|"failed", "path",
             "reason", "method"}."""
    os.makedirs(out_dir, exist_ok=True)
    name = label(document_type)
    if not template or template.get("mode") not in ("uploaded", "ai_guideline"):
        reason = not_set_up_reason(document_type, company)
        log.missing(f"{document_type}.setup", reason)
        return {"status": "not_set_up", "path": None, "reason": reason, "method": None}
    if template["mode"] == "uploaded":
        src = template.get("path")
        if not src or not os.path.exists(src):
            reason = f"{name}: the uploaded copy could not be fetched - upload it again on the Company page."
            log.missing(f"{document_type}.uploaded", reason)
            return {"status": "failed", "path": None, "reason": reason, "method": "uploaded"}
        fname = re.sub(r"[^A-Za-z0-9._ ()-]+", "_", template.get("original_filename") or os.path.basename(src))
        dest = os.path.join(out_dir, f"{document_type}__{fname}")
        shutil.copyfile(src, dest)
        if _sha256(dest) != _sha256(src):
            raise RuntimeError(f"{name}: the as-is copy differs from the upload")
        log.confirmed(f"{document_type}.uploaded", template.get("original_filename") or os.path.basename(src))
        return {"status": "uploaded", "path": dest, "reason": None, "method": "uploaded", "sha256": _sha256(dest)}
    guideline = (template.get("guideline") or "").strip()
    if not guideline:
        reason = not_set_up_reason(document_type, company)
        log.missing(f"{document_type}.setup", reason)
        return {"status": "not_set_up", "path": None, "reason": reason, "method": None}
    text, method = render_guideline(guideline, facts, document_type, log, llm)
    out = write_docx(text, os.path.join(out_dir, f"{document_type}.docx"), logo_path)
    return {"status": "generated", "path": out, "reason": None, "method": method, "text": text}


def authority_facts(profile, signatory, tender, rfp_number=None, today=None):
    """The authority-to-sign fields for one tender, from the company's own
    records. Nothing is inferred: a value not on file stays None."""
    comp = profile.get("company") or {}
    director = comp.get("director") or {}
    closing = tender.get("closing_date")
    return {
        "company_legal_name": comp.get("legal_name"),
        "signatory_name": (signatory or {}).get("name"),
        "signatory_capacity": (signatory or {}).get("signatory_capacity"),
        "director_name": director.get("name") if isinstance(director, dict) else None,
        "tender_reference": rfp_number,
        "tender_description": tender.get("description"),
        "contract_number": tender.get("contract_no"),
        "nra_reference": tender.get("nra_reference"),
        "closing_date": closing,
        "date": (today or datetime.date.today()).strftime("%-d %B %Y"),
    }
