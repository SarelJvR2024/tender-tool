"""
Focused tests for classify_tender.py's classify_documents() - specifically its
.docx branch, which had no dedicated test file before this (test_manual_intake.py
needs Postgres so it isn't in CI; it also only covers the happy path end to end).

Finding (2026-10-20, thirty-fifth run, docs/specs/QA-Edge-Case-Findings.md): the
.xlsm branch already logs a note when a workbook can't be opened ("unreadable" ->
notes.append(...)), but the parallel .docx branch's `except Exception: continue`
silently dropped a corrupt/mislabeled SANRAL .docx form with no trace at all. The
downstream "missing documents" message then told the user the document was simply
"not found" even though they HAD uploaded it - it just couldn't be read. Fixed by
logging a note in the .docx branch's exception handler too, mirroring the .xlsm
branch's existing pattern.

Run: python3 engine/intake/test_classify_tender.py
"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import classify_tender as ct  # noqa: E402
import docx  # noqa: E402

FAILS = []

# A real Technical Submission workbook, so classify_documents() actually
# routes "sanral" (the route the .docx tests below need to exercise
# sanral_files/sanral_documents - those are deliberately {}/[] on the
# general route, by design, when no workbook is present at all).
TECH_WORKBOOK = os.path.join(os.path.dirname(HERE), "..", "templates", "samples", "R.352-020-2025.1F_ran.xlsm")

# Finding, 2026-10-27: a real A-Forms.docx that genuinely needs more than the
# old 60-paragraph head to find its own "Part T2: Returnable Schedules"
# heading (it sits at paragraph 521 of 2104, behind a long notes/TOC
# section) - see classify_tender.py's own _docx_head() comment.
REAL_A_FORMS = os.path.join(os.path.dirname(HERE), "..", "templates", "forms", "R.352-020-2025.1F", "A_Forms.docx")


def check(name, cond, detail=""):
    print(("  ok   " if cond else "  FAIL ") + name + (f"  ({detail})" if detail and not cond else ""))
    if not cond:
        FAILS.append(name)


def _doc_entry(name, local, fmt="docx"):
    return {"name": name, "format": fmt, "_local": local, "unreadable": None}


def test_corrupt_docx_is_noted_not_silently_dropped(tmp_path="/tmp"):
    print("classify_documents: a corrupt/garbage .docx SANRAL form")
    path = os.path.join(tmp_path, "C.1.1.1 Form of Offer (corrupt).docx")
    with open(path, "wb") as f:
        f.write(b"not really a docx, just garbage bytes")
    decision = ct.classify_documents([_doc_entry("C.1.1.1 Form of Offer (corrupt).docx", path)],
                                      buyer="SANRAL", all_text="SANRAL tender")
    check("a note is logged for the unreadable .docx (not silently dropped)",
          any("C.1.1.1 Form of Offer (corrupt).docx" in n and "could not be opened" in n
              for n in decision["notes"]), decision["notes"])
    check("it is not registered as a found SANRAL file",
          "c111_form_of_offer" not in decision["sanral_files"])
    os.remove(path)


def test_valid_docx_is_classified_with_no_spurious_note(tmp_path="/tmp"):
    print("classify_documents: a real, readable C.1.1.1 Form of Offer .docx")
    path = os.path.join(tmp_path, "C.1.1.1 Form of Offer.docx")
    d = docx.Document()
    d.add_paragraph("C1.1.1 FORM OF OFFER (Incorporating SBD7)")
    d.save(path)
    decision = ct.classify_documents([_doc_entry("C.1.1.1 Form of Offer.docx", path)],
                                      buyer="SANRAL", all_text="SANRAL tender")
    check("no note is logged for a readable docx that matched", decision["notes"] == [], decision["notes"])
    os.remove(path)


def test_period_after_the_form_letter_is_also_recognized(tmp_path="/tmp"):
    print("classify_documents: a real C.1.1.1 Form of Offer whose IN-DOCUMENT heading has a "
          "period after the letter (\"C.1.1.1\", not \"C1.1.1\")")
    # QA audit 2026-10-20 ("enterprise readiness" pass): the two tests above
    # both use "C.1.1.1" in the FILENAME but "C1.1.1" (no period) in the
    # in-document heading text - the exact mismatch that let the period-in-
    # heading case go untested while SANRAL_DOCX's regexes silently failed
    # on it. This test exercises the heading text itself, not just the name.
    path = os.path.join(tmp_path, "Form of Offer - period heading.docx")
    d = docx.Document()
    d.add_paragraph("C.1.1.1 FORM OF OFFER (Incorporating SBD7)")
    d.save(path)
    entries = [_doc_entry("Form of Offer - period heading.docx", path),
               _doc_entry("Technical Submission.xlsm", TECH_WORKBOOK, fmt="xlsm")]
    decision = ct.classify_documents(entries, buyer="SANRAL", all_text="SANRAL tender")
    check("route is sanral (the workbook alone already guarantees this)", decision["route"] == "sanral", decision)
    check("a period after the form letter in the heading text is still recognized",
          "c111_form_of_offer" in decision["sanral_files"], decision)
    # Other notes are expected here (C.1.2.3/A-Forms/Financial workbook
    # genuinely weren't supplied in this fixture) - the only thing this
    # check cares about is that the period-heading document ITSELF isn't
    # flagged as unread/ambiguous/corrupt.
    check("the period-heading document itself is never named in a note (a clean match, not flagged)",
          not any("period heading" in n or "could not be opened" in n or "AMBIGUOUS" in n for n in decision["notes"]),
          decision["notes"])
    os.remove(path)


def test_duplicate_candidates_prefer_the_more_recently_modified_file(tmp_path="/tmp"):
    print("classify_documents: two files both match c111_form_of_offer - the more recently "
          "modified one should win, not whichever sorts first alphabetically")
    import time
    old_path = os.path.join(tmp_path, "Old_Draft_Form_Offer.docx")
    new_path = os.path.join(tmp_path, "Zfinal_Form_Offer.docx")
    d1 = docx.Document(); d1.add_paragraph("C1.1.1 FORM OF OFFER (Incorporating SBD7) - DRAFT")
    d1.save(old_path)
    time.sleep(1.1)
    d2 = docx.Document(); d2.add_paragraph("C1.1.1 FORM OF OFFER (Incorporating SBD7) - FINAL")
    d2.save(new_path)
    entries = [_doc_entry("Old_Draft_Form_Offer.docx", old_path), _doc_entry("Zfinal_Form_Offer.docx", new_path),
               _doc_entry("Technical Submission.xlsm", TECH_WORKBOOK, fmt="xlsm")]
    decision = ct.classify_documents(entries, buyer="SANRAL", all_text="SANRAL tender")
    check("the more recently modified file (Zfinal, alphabetically LAST) is the one actually used",
          decision["sanral_files"].get("c111_form_of_offer") == new_path, decision)
    check("the note names it as AMBIGUOUS and tells a person to check, not a soft 'remove the others if wrong'",
          any(n.startswith("AMBIGUOUS") and "must check" in n for n in decision["notes"]), decision["notes"])
    os.remove(old_path)
    os.remove(new_path)


def test_duplicate_technical_workbooks_signal_names_the_same_file_it_actually_used(tmp_path="/tmp"):
    print("classify_documents: two Technical Submission .xlsm candidates - the top-line 'signal' "
          "text must name the SAME file the tie-break actually picked (files/the AMBIGUOUS note), "
          "not whichever sorted first before the tie-break ran")
    # QA pass, 2026-10-26: family/signal were read from `tech[0]` - candidates["technical_xlsm"]
    # in its ORIGINAL, pre-tie-break (alphabetical) order - while files["technical_xlsm"] and the
    # AMBIGUOUS note already correctly used the tie-broken (most-recently-modified) pick. Whenever
    # the two orders disagreed, the signal told a person the WRONG file was "routed to the SANRAL
    # submission-pack flow", directly contradicting the AMBIGUOUS note two lines below it in the
    # very same decision. "Draft" < "ZZfinal" alphabetically but is modified FIRST (older mtime),
    # so this reproduces the divergence deterministically.
    import shutil
    import time
    old_path = os.path.join(tmp_path, "Draft_Technical_Submission.xlsm")
    new_path = os.path.join(tmp_path, "ZZfinal_Technical_Submission.xlsm")
    shutil.copy(TECH_WORKBOOK, old_path)
    time.sleep(1.1)
    shutil.copy(TECH_WORKBOOK, new_path)
    entries = [_doc_entry("Draft_Technical_Submission.xlsm", old_path, fmt="xlsm"),
               _doc_entry("ZZfinal_Technical_Submission.xlsm", new_path, fmt="xlsm")]
    decision = ct.classify_documents(entries, buyer="SANRAL", all_text="SANRAL tender")
    winner = os.path.basename(decision["sanral_files"].get("technical_xlsm") or "")
    check("the more recently modified file (ZZfinal) is the one actually used",
          winner == "ZZfinal_Technical_Submission.xlsm", decision)
    check("the signal text names that SAME file, not the alphabetically-first one",
          winner in decision["signal"] and "Draft_Technical_Submission.xlsm" not in decision["signal"],
          decision["signal"])
    os.remove(old_path)
    os.remove(new_path)


def test_real_a_forms_heading_deep_in_the_document_is_still_found():
    print("classify_documents: a real A-Forms.docx whose own heading sits at paragraph "
          "521 of 2104, well past the old 60-paragraph head cap")
    entries = [_doc_entry("A_Forms.docx", REAL_A_FORMS), _doc_entry("Technical Submission.xlsm", TECH_WORKBOOK, fmt="xlsm")]
    decision = ct.classify_documents(entries, buyer="SANRAL", all_text="SANRAL tender")
    check("the real A-Forms document is recognized, not silently missed",
          decision["sanral_files"].get("a_forms") == REAL_A_FORMS, decision)
    check("it's not reported as a missing SANRAL document",
          not any("A-Forms / returnable schedules" in n and "Not found" in n for n in decision["notes"]),
          decision["notes"])


def test_docx_already_marked_unreadable_is_skipped_without_a_note():
    print("classify_documents: a doc_entry already flagged unreadable upstream")
    entry = _doc_entry("whatever.docx", "/tmp/does-not-matter.docx")
    entry["unreadable"] = "already flagged by requirements.py"
    entry["_local"] = None  # not even checked since _local is falsy anyway, but exercises the real shape
    decision = ct.classify_documents([entry], buyer="SANRAL", all_text="SANRAL tender")
    check("no note is added for a doc_entry with no _local path at all", decision["notes"] == [])


FIN_WORKBOOK = os.path.join(os.path.dirname(HERE), "..", "templates", "samples", "financial",
                            "R.334-030-2025.1F_Financial_Submission_ran.xlsm")


def test_classify_workbook_real_technical_financial_and_corrupt(tmp_path="/tmp"):
    print("classify_workbook: real technical workbook, real financial workbook, a corrupt file")
    # None of classify_workbook()/read_setup()/find_nra_reference()/
    # find_delivery_address()/find_closing_time()/discrepancies() had any
    # dedicated test anywhere in the repo before this pass (test_manual_intake.py
    # needs Postgres and never calls any of them directly either) - despite
    # find_nra_reference/find_delivery_address/find_closing_time each
    # carrying their own documented "ambiguous" edge-case fix (QA passes,
    # 2026-10-23/24) with nothing regression-testing that it still holds.
    tech = ct.classify_workbook(TECH_WORKBOOK)
    check("the real Technical Submission workbook is classified as technical, with its family",
          tech["kind"] == "technical" and tech["family"] == "Conventional Consultant Tender", tech)
    fin = ct.classify_workbook(FIN_WORKBOOK)
    check("the real Financial Submission workbook is classified as financial, no family",
          fin["kind"] == "financial" and fin["family"] is None, fin)
    corrupt_path = os.path.join(tmp_path, "corrupt.xlsm")
    with open(corrupt_path, "wb") as f:
        f.write(b"not a real xlsm, just bytes")
    bad = ct.classify_workbook(corrupt_path)
    check("a corrupt/garbage file is reported unreadable, not raised or silently 'other'",
          bad["kind"] == "unreadable" and bad["sheets"] == [] and bad["error"], bad)
    os.remove(corrupt_path)


def test_read_setup_against_the_real_fixture():
    print("read_setup: the real Technical Submission workbook's own SETUP sheet")
    setup = ct.read_setup(TECH_WORKBOOK)
    check("contract_no read from SETUP!C2", setup["contract_no"] == "R.352-020-2025/1F", setup)
    check("closing_date read from SETUP!C3 as a real ISO date (Excel serial decoded)",
          setup["closing_date"] == "2026-10-22", setup)
    check("the two real key positions are read, each with its own managerial flag and alternate",
          {p["code"] for p in setup["positions"]} == {"PL", "DS1"}
          and all(p["managerial"] is False for p in setup["positions"]), setup)


def test_find_nra_reference_single_ambiguous_and_none():
    print("find_nra_reference: single match, most-repeated-wins when ambiguous, no match at all")
    check("a single NRA reference is read in the seed's own format",
          ct.find_nra_reference("Book 1 cover page. NRA: No. 2026/0071. Footer repeats NRA 2026/0071.")
          == {"value": "NRA 2026/0071", "ambiguous": False})
    check("no NRA reference anywhere returns None, not a guess",
          ct.find_nra_reference("This document mentions no NRA number at all.") is None)
    check("an empty/None input returns None without raising", ct.find_nra_reference(None) is None)
    # The real shape this ambiguity exists for: a stale number repeated on
    # every page footer (counted many times) vs. a single correcting
    # addendum mention - repetition count alone must not silently win.
    text = "NRA: 2026/0071\n" * 5 + "Addendum 1: the correct reference is NRA 2026/9999."
    result = ct.find_nra_reference(text)
    check("more than one distinct NRA number sets ambiguous=True even though one is repeated far more",
          result["ambiguous"] is True, result)
    check("the most-repeated value is still the default 'value' (unchanged normal-case behaviour)",
          result["value"] == "NRA 2026/0071", result)


def test_find_closing_time_single_ambiguous_and_formats():
    print("find_closing_time: single match, multiple SEPARATOR formats of the same time are not "
          "ambiguous, genuinely different times are")
    check("a single closing time is read as HH:MM",
          ct.find_closing_time("The closing time for the receipt of tenders is at 12:00 on 22 October 2026.")
          == {"value": "12:00", "ambiguous": False})
    check("'h' and '.' separators and a single-digit hour all normalize to the same HH:MM, not ambiguous",
          ct.find_closing_time(
              "The closing time and date for receipt of tenders is 9h00. "
              "The closing date and time for receipt of tenders is 09.00."
          ) == {"value": "09:00", "ambiguous": False})
    check("no closing-time sentence anywhere returns None", ct.find_closing_time("Nothing about a closing time here.") is None)
    # The real shape: an addendum extending the deadline from 12:00 to 14:00.
    text = ("The closing time for the receipt of tenders is at 12:00 on 22 October 2026. "
            "Addendum 2: the closing time for the receipt of tenders is at 14:00 on 29 October 2026.")
    result = ct.find_closing_time(text)
    check("a genuinely different closing time elsewhere in the text sets ambiguous=True", result["ambiguous"] is True, result)
    check("'value' stays the FIRST match found (unchanged default), not the later/extended one",
          result["value"] == "12:00", result)


def test_find_delivery_address_single_ambiguous_and_stops_correctly():
    print("find_delivery_address: single address, a genuinely different second address is ambiguous, "
          "line collection stops at a footer/blank/sentence")
    text = ("The tender box situated at:\n"
            "The Reception Area\n"
            "20 Shoreward Drive\n"
            "Bay West\n"
            "Gqeberha\n"
            "6025\n"
            "\n"
            "T-6\n")
    result = ct.find_delivery_address(text)
    check("the real multi-line address under the heading is read in full",
          result["lines"] == ["The Reception Area", "20 Shoreward Drive", "Bay West", "Gqeberha", "6025"], result)
    check("not ambiguous when only one address appears", result["ambiguous"] is False, result)
    check("collection stops at the blank line, never pulling in the page-footer code below it",
          "T-6" not in result["lines"], result)
    check("no delivery-address heading anywhere returns None", ct.find_delivery_address("Nothing about a tender box here.") is None)
    # The real shape: an addendum moving the delivery venue to a different address.
    moved = text + "\nAddendum 3: tenders must be deposited at the following address:\nUnit 4, Freedom Square\nPort Elizabeth\n7001\n"
    result = ct.find_delivery_address(moved)
    check("a genuinely different second address sets ambiguous=True", result["ambiguous"] is True, result)
    check("'lines' stays the FIRST address found (unchanged default), not the moved/later one",
          result["lines"][0] == "The Reception Area", result)


def test_discrepancies_contract_no_and_closing_date():
    print("discrepancies: contract number normalization, closing date mismatch, matching values, "
          "and a missing side on either end skips the check entirely")
    setup = {"contract_no": "R.352-020-2025/1F", "closing_date": "2026-10-22"}
    # Punctuation-only difference must NOT be reported - both sides normalize
    # to the same alphanumeric string.
    same = ct.discrepancies({"contract_no": "r352 020 2025 1f", "closing_date": "2026-10-22"}, setup)
    check("a punctuation/case-only difference in contract_no is not a discrepancy", same == [], same)
    mismatch = ct.discrepancies({"contract_no": "R.352-020-2025/2F", "closing_date": "2026-10-22"}, setup)
    check("a genuinely different contract number IS reported", len(mismatch) == 1 and "R.352-020-2025/2F" in mismatch[0], mismatch)
    date_mismatch = ct.discrepancies({"contract_no": "R.352-020-2025/1F", "closing_date": "2026-11-05"}, setup)
    check("a genuinely different closing date IS reported", len(date_mismatch) == 1 and "2026-11-05" in date_mismatch[0], date_mismatch)
    missing_user_side = ct.discrepancies({"contract_no": None, "closing_date": None}, setup)
    check("nothing the user typed yet (both None) is never a discrepancy, not a crash", missing_user_side == [], missing_user_side)
    missing_setup_side = ct.discrepancies({"contract_no": "anything", "closing_date": "2026-10-22"}, {"contract_no": None, "closing_date": None})
    check("no SETUP value to compare against is never a discrepancy either", missing_setup_side == [], missing_setup_side)


def test_assess_fit_from_text_categorized_tender_no_disciplines_not_duplicated():
    print("assess_fit_from_text: categorized tender, company has no disciplines on file - "
          "the actionable message appears once, not duplicated")
    # Sarel's own live feedback, 2026-10-23: the live Discovery feed showed a
    # generic, unhelpful version of this message. Fixed at the root
    # (assess_tender.assess_fit()) so assess_fit_from_text()'s own call site
    # - which already had a correct, duplicate-ANY-other-reason workaround -
    # doesn't double it up now that the root supplies the same wording.
    title = "APPOINTMENT OF EXTERNAL SECURITY SERVICES AT THE EILAND HOLIDAY RESORT, ROSE PARK & REITZPARK"
    result = ct.assess_fit_from_text(title, "", "", {"disciplines": [], "has_cidb_grading": False})
    check("verdict is needs_review, not rejected", result["verdict"] == "needs_review", result["verdict"])
    hits = [r for r in result["reasons"] if "Experience Library tags" in r]
    check("the actionable message appears exactly once (assess_fit() supplied it; "
          "assess_fit_from_text() must not add a second copy)", len(hits) == 1, result["reasons"])


def test_assess_fit_from_text_uncategorized_tender_no_disciplines_still_appends():
    print("assess_fit_from_text: uncategorized tender AND no disciplines on file - both facts "
          "are true and both are worth telling the user, so the no-disciplines note is still appended")
    # Here assess_fit() itself takes the "uncategorized" branch first (a
    # different primary reason), which never mentions disciplines at all -
    # this is the one case assess_fit_from_text()'s own extra check exists
    # for, and it must keep firing.
    title = "Appointment of a service provider ref ABC123-XYZ"
    result = ct.assess_fit_from_text(title, "", "", {"disciplines": [], "has_cidb_grading": False})
    check("verdict is needs_review", result["verdict"] == "needs_review", result["verdict"])
    check("the uncategorized reason is present",
          any("Could not determine discipline" in r for r in result["reasons"]), result["reasons"])
    hits = [r for r in result["reasons"] if "Experience Library tags" in r]
    check("the no-disciplines note is still appended exactly once, even though a different "
          "reason fired first", len(hits) == 1, result["reasons"])


def main():
    test_corrupt_docx_is_noted_not_silently_dropped()
    test_valid_docx_is_classified_with_no_spurious_note()
    test_period_after_the_form_letter_is_also_recognized()
    test_duplicate_candidates_prefer_the_more_recently_modified_file()
    test_duplicate_technical_workbooks_signal_names_the_same_file_it_actually_used()
    test_real_a_forms_heading_deep_in_the_document_is_still_found()
    test_docx_already_marked_unreadable_is_skipped_without_a_note()
    test_classify_workbook_real_technical_financial_and_corrupt()
    test_read_setup_against_the_real_fixture()
    test_find_nra_reference_single_ambiguous_and_none()
    test_find_closing_time_single_ambiguous_and_formats()
    test_find_delivery_address_single_ambiguous_and_stops_correctly()
    test_discrepancies_contract_no_and_closing_date()
    test_assess_fit_from_text_categorized_tender_no_disciplines_not_duplicated()
    test_assess_fit_from_text_uncategorized_tender_no_disciplines_still_appends()
    print(f"\n{len(FAILS)} failure(s)")
    sys.exit(1 if FAILS else 0)


if __name__ == "__main__":
    main()
