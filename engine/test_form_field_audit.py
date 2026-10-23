"""
Regression test for docs/specs/Form-Field-Audit.md (2026-10-18): every field
the audit fixed, run on the REAL SANRAL templates (templates/forms/,
templates/samples/) with a synthetic company built to be UNLIKE Klipcorp on
every axis the audit found a Klipcorp-shaped default for:

  - a close corporation (not "(Pty) Ltd"), two people with a signatory
    capacity, a separate PL and DS;
  - no B-BBEE level recorded (then: an EME affidavit; then: a valid
    certificate; then: an expired one);
  - Form B3 counts never entered (the 0001 DEFAULT 0 pattern), no head office,
    no JV share, no director;
  - no CSD report note.

Each check names the field and the old behaviour it guards against. No
network, no LibreOffice, no database. Run: python3 engine/test_form_field_audit.py
"""
import copy
import json
import os
import sys
import tempfile

import docx
import yaml

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, ".."))
for p in (HERE, os.path.join(HERE, "pack_generator"), os.path.join(HERE, "general"), os.path.join(HERE, "intake")):
    if p not in sys.path:
        sys.path.insert(0, p)

import fill_a_forms  # noqa: E402
import fill_docx_forms  # noqa: E402
import fill_conventional  # noqa: E402
import fill_financial  # noqa: E402
import make_envelope_covers  # noqa: E402
import confirm_pack  # noqa: E402
from provenance import ProvenanceLog  # noqa: E402
from sanral_fill import Workbook  # noqa: E402
from suggest_pack import read_required_positions  # noqa: E402

FORMS = os.path.join(REPO, "templates", "forms", "R.396-010-2025.1F")
TECH = os.path.join(REPO, "templates", "samples", "R.396-010-2025.1F_ran.xlsm")
FIN_DIR = os.path.join(REPO, "templates", "samples", "financial")
KLIPCORP_PHRASES = ("Klipcorp", "Henning", "sworn 2026-09-09", "June 2026 CSD", "2021 CSD", "Professional engineering services",
                    "PO Box 415", "CHECK PROFILE")

failures = []


def check(cond, what):
    print(("  ok   " if cond else "  FAIL ") + what)
    if not cond:
        failures.append(what)


PERSON_QUAL = {
    "ecsa_level": "Professional Engineer", "ecsa_number": "20009901", "ecsa_date": "2014-03-01",
    "sacpcmp_number": None, "highest_level": "Masters Degree", "highest_description": "MSc (Engineering Geology)",
    "highest_institution": "Test University", "highest_graduated": "2012-01-01",
    "initial_description": "BSc (Geology)", "initial_institution": "Test University", "initial_graduated": "2007-01-01",
}


def person(name, capacity=None, status=None):
    return {"name": name, "id_number": "0101015000081", "description": "", "role": None, "signatory_capacity": capacity,
            "phone": None, "email": None, "employment_status": status, "qualifications": dict(PERSON_QUAL), "projects": []}


UNLIKE = {
    "company": {
        "legal_name": "Ubuntu Geo-Surveys CC", "cipc_registration_number": "", "registered_address": "",
        "postal_address": None, "vat_number": None, "csd_supplier_number": "MAAA0999999",
        "bee_level": None, "black_ownership_pct": None, "status": None, "bee_proof_type": None, "bee_proof_expiry": None,
        "sars_tcs_pin": None, "professional_indemnity_insurance": None, "bank": {"bank_name": None, "branch_code": None},
        "director": {"name": ""}, "telephone": None, "email": None,
        # the untouched 0001 pattern: DEFAULT 0 x3, total employees never had a default
        # no head-office key at all (a CLI profile; the bridge passes None)
        "other_offices": 0, "ecsa_registered_staff_count": 0,
        "sacpcmp_registered_staff_count": 0, "total_employees": None, "typical_jv_share_pct": None,
        "coid_letter_of_good_standing_number": None, "coid_letter_of_good_standing_expiry": None,
    },
    "key_people": [person("Thandi Signatory-Member", "Member"), person("Lerato Project-Leader", "Managing Member", "full_time_employee"),
                   person("Riaan Design-Specialist", None, "not_full_time_employee")],
}


def write_json(obj, d, name):
    p = os.path.join(d, name)
    json.dump(obj, open(p, "w"), indent=1, default=str)
    return p


def entries(log):
    return {e["field_key"]: e for e in log.entries}


def a_forms(profile, d):
    _, log = fill_a_forms.main(os.path.join(FORMS, "A_Forms.docx"), write_json(profile, d, "p.json"), d, "R.396", "NRA 2026/0071")
    return entries(log), docx.Document(os.path.join(d, "A_Forms_FILLED.docx"))


def all_reasons(es):
    return " ".join(f"{e.get('value') or ''} {e.get('reason') or ''}" for e in es.values())


def test_a_forms():
    print("A-Forms (fill_a_forms.py)")
    d = tempfile.mkdtemp()
    es, doc = a_forms(UNLIKE, d)
    check(es["a_forms.A3.joint_venture"]["status"] == "assumed_default",
          "A3 'not a Joint Venture' is assumed_default (was confirmed - 'a known fact about the company')")
    check(es["a_forms.A7.supplier_commodity"]["status"] == "missing" and "a_forms.A7.commodity_code" not in es,
          "A7 commodity is missing, not Klipcorp's 'Professional engineering services'")
    check(es["a_forms.A12.bee_certificate"]["status"] == "missing" and "No B-BBEE level" in es["a_forms.A12.bee_certificate"]["reason"],
          "A12 with no level on file: missing, says why")
    check("a_forms.A16.addenda_received" in es and es["a_forms.A16.addenda_received"]["status"] == "missing",
          "A16 with no update record: logged missing (was never logged at all - the statement is in a table)")
    check(not any(p in all_reasons(es) for p in KLIPCORP_PHRASES), "no Klipcorp-specific wording in any A-Forms provenance line")

    prof = copy.deepcopy(UNLIKE)
    prof["company"].update(bee_level=1, bee_proof_type="certificate", bee_proof_expiry="2099-01-31", black_ownership_pct=51, status="EME")
    prof["tender_updates"] = [{"date": "2026-10-01", "title": "Addendum 1 - revised Form C2.4"}, {"date": "2026-10-08", "title": "Clarification 2"}]
    es, doc = a_forms(prof, tempfile.mkdtemp())
    check(es["a_forms.A12.bee_certificate"]["status"] == "confirmed" and "Level 1" in es["a_forms.A12.bee_certificate"]["value"]
          and "51% black owned" in es["a_forms.A12.bee_certificate"]["value"],
          "A12 with a valid certificate: confirmed from THIS company's level/ownership (was missing for everyone but Klipcorp)")
    t = fill_a_forms._a16_table(doc)
    check(t.rows[2].cells[2].text == "Addendum 1 - revised Form C2.4" and t.rows[3].cells[1].text == "2026-10-08",
          "A16 lists the addenda the update watch recorded, in its own table")
    check(es["a_forms.A16.addenda_received"]["status"] == "assumed_default", "A16 listing is assumed_default (check the portal)")

    prof["company"]["bee_proof_expiry"] = "2020-01-01"
    prof["tender_updates"] = []
    es, doc = a_forms(prof, tempfile.mkdtemp())
    check(es["a_forms.A12.bee_certificate"]["status"] == "missing" and "expired 2020-01-01" in es["a_forms.A12.bee_certificate"]["reason"],
          "A12 with an expired certificate: missing, says it expired")
    check("None issued" in fill_a_forms._a16_table(doc).rows[2].cells[2].text, "A16 with a watch that saw nothing: 'None issued'")


def test_docx_forms():
    print("C.1.1.1 / C.1.2.3 (fill_docx_forms.py)")
    d = tempfile.mkdtemp()
    log = fill_docx_forms.main(FORMS, write_json(UNLIKE, d, "p.json"), d)
    es = entries(log)
    offer = "\n".join(p.text for p in docx.Document(os.path.join(d, "C.1.1.1_Form_of_Offer_FILLED.docx")).paragraphs)
    check(" None " not in offer and es["c111.bee_level"]["status"] == "missing",
          "Form of Offer: no level on file -> claim left blank + missing (was ' None ' typed in, confirmed)")
    check("Ubuntu Geo-Surveys CC," not in offer and es["c111.organisation_address"]["status"] == "missing",
          "Form of Offer: no address -> no dangling comma, address missing")
    check(es["c123.director_name"]["status"] == "missing", "C.1.2.3 directors table: no director -> missing (was confirmed '')")
    prof = copy.deepcopy(UNLIKE)
    prof["company"]["director"] = {"name": "Thandi Signatory-Member"}
    d = tempfile.mkdtemp()
    es = entries(fill_docx_forms.main(FORMS, write_json(prof, d, "p.json"), d))
    t = docx.Document(os.path.join(d, "C.1.2.3_Contract_Data_FILLED.docx")).tables[2]
    check(t.rows[1].cells[2].text == "Member", "C.1.2.3: a close corporation's director row says Member, not Director")
    check(es["c123.directors_table_complete"]["status"] == "missing", "C.1.2.3: says only one director/member is on file")
    # the company's own directors register (company_principals, 0050) when it has rows
    prof["company"]["principals"] = [
        {"name": "Thandi Signatory-Member", "appointed_on": "2011-04-01", "designation": "Member", "registration_number": None},
        {"name": "Lerato Project-Leader", "appointed_on": None, "designation": None, "registration_number": "20009901"}]
    d = tempfile.mkdtemp()
    es = entries(fill_docx_forms.main(FORMS, write_json(prof, d, "p.json"), d))
    t = docx.Document(os.path.join(d, "C.1.2.3_Contract_Data_FILLED.docx")).tables[2]
    check([c.text for c in t.rows[2].cells] == ["Lerato Project-Leader", "", "Member", "20009901"]
          and t.rows[1].cells[1].text == "2011-04-01" and es["c123.directors_table"]["status"] == "confirmed",
          "C.1.2.3: every director/member from the register, with appointment date and registration number")
    check(es["c123.directors_appointment_dates"]["status"] == "missing" and "c123.directors_table_complete" not in es,
          "C.1.2.3: an undated register row is flagged; no 'only one director' line once the register is used")


def confirmed_pack(profile, d):
    """pack.yaml through confirm_pack.confirm(), PL = person_1, every other SETUP position = person_2."""
    primary, _ = read_required_positions(Workbook(TECH))
    sel = {"confirmed_by": "test", "confirmed_date": "2026-10-18",
           "selections": {p["code"]: {"confirmed_person": "person_1" if p["code"] == "PL" else "person_2"} for p in primary}}
    yaml.safe_dump(sel, open(os.path.join(d, "sel.yaml"), "w"))
    write_json({"contract_no": "R.396-010-2025/1F", "closing_date": "2026-10-22", "positions": {}}, d, "sugg.json")
    out = os.path.join(d, "pack.yaml")
    confirm_pack.confirm(os.path.join(d, "sugg.json"), os.path.join(d, "sel.yaml"), write_json(profile, d, "prof.json"), out)
    return out


def tech(profile):
    d = tempfile.mkdtemp()
    pack = confirmed_pack(profile, d)
    log = ProvenanceLog("fill_conventional.py")
    out = os.path.join(d, "filled.xlsm")
    notes = fill_conventional.main(TECH, pack, out, log)
    return Workbook(out), notes, entries(log), yaml.safe_load(open(pack))


def test_technical_workbook():
    print("TEND 1X front page + Form B3 (fill_conventional.py via confirm_pack.py)")
    wb, notes, es, pack = tech(UNLIKE)
    T = "TEND 1X"
    check(pack["company"]["head_office"] is None, "pack.yaml: head office not on file is None (was 'CHECK PROFILE - head_office not set')")
    check(wb.cached(T, "K33") == 9 and es["tend1x.bee_level"]["status"] == "missing",
          "K33: no level on file keeps SANRAL's 9 + missing (was an empty cell, unlogged)")
    check(wb.cached(T, "H34") == "GENERIC", "H34: no status on file keeps SANRAL's 'GENERIC' (was emptied)")
    check(wb.cached(T, "C33") == 2 and es["tend1x.verification_certificate"]["value"] == "No Valid Certificate",
          "C33: no B-BBEE proof -> 'No Valid Certificate' (SANRAL's shipped 'Verification Certificate' was never touched)")
    check(all(wb.cached(T, c) in (None, "") for c in ("F3339", "F3340", "F3341", "F3342", "F3344", "F3345")),
          "Form B3: untouched DEFAULT-0 counts, head office and JV share left blank (were 0 / placeholder text)")
    check(all(es[f"form_b3.{k}"]["status"] == "missing" for k in ("head_office", "other_offices", "ecsa_registered",
                                                                  "sacpcmp_registered", "total_employees", "jv_share_pct")),
          "Form B3: each of them logged missing")
    check(any("Form B3: REGISTERED PROFESSIONALS: ECSA is not on file" in n for n in notes), "Form B3 gaps are pack notes")
    check(not any(p in " ".join(notes) for p in KLIPCORP_PHRASES), "no Klipcorp-specific wording in the workbook notes")

    prof = copy.deepcopy(UNLIKE)
    prof["company"].update(bee_level=1, status="EME", black_ownership_pct=100, bee_proof_type="eme_affidavit",
                           bee_proof_expiry="2099-01-01", head_office_per_form_b3="Polokwane", other_offices=0,
                           ecsa_registered_staff_count=0, sacpcmp_registered_staff_count=2, total_employees=14,
                           typical_jv_share_pct=60)
    wb, notes, es, _ = tech(prof)
    check(wb.cached(T, "C33") == 1 and es["tend1x.verification_certificate"]["status"] == "assumed_default",
          "C33: an EME affidavit keeps SANRAL's answer but is assumed_default with the A12 caveat")
    check(wb.cached(T, "F3339") == "Polokwane" and wb.cached(T, "F3344") == 14 and wb.cached(T, "F3340") == 0,
          "Form B3: real values written (a real 0 other offices is kept once the record has been filled in)")
    check(es["form_b3.ecsa_registered"]["status"] == "assumed_default", "Form B3: ECSA count 0 on a filled-in record is flagged to confirm")
    check(wb.cached(T, "F3345") == 60 and es["form_b3.jv_share_pct"]["status"] == "assumed_default"
          and any("% SHARE OF JV AGREEMENT = 60" in n for n in notes),
          "Form B3: JV share is the company's standing value -> assumed_default + note (was written as fact)")

    prof["company"].update(bee_proof_type="certificate", bee_proof_expiry="2026-10-01")
    wb, notes, es, _ = tech(prof)
    check(wb.cached(T, "C33") == 2 and "before this tender closes" in es["tend1x.verification_certificate"]["reason"],
          "C33: a certificate that lapses before closing -> 'No Valid Certificate', says why")
    prof["company"]["bee_proof_expiry"] = "2027-06-30"
    wb, notes, es, _ = tech(prof)
    check(wb.cached(T, "C33") == 1 and es["tend1x.verification_certificate"]["status"] == "confirmed",
          "C33: a certificate valid at closing -> 'Verification Certificate', confirmed")


def test_covers():
    print("Envelope covers (make_envelope_covers.py)")
    d = tempfile.mkdtemp()
    args = ({"company": UNLIKE["company"]}, None, d, "SANRAL R.396-010-2025/1F", "NRA 2026/0071", "Test", "2026-10-22", "12:00")
    text = "\n".join(p.text for p in docx.Document(make_envelope_covers.make_outer(*args)).paragraphs)
    check("PO Box 415" not in text and "PRETORIA" not in text and "Chief Executive Officer" not in text,
          "no fixed SANRAL head-office address when the tender's own isn't on file")
    check("CIPC Registration No:" not in text, "no empty 'CIPC Registration No:' line for a company without one")
    lines = ["The Reception Area", "20 Shoreward Drive", "Bay West", "Gqeberha", "6025"]
    text = "\n".join(p.text for p in docx.Document(make_envelope_covers.make_inner(*args, 1, "TECHNICAL", "X", delivery_address=lines)).paragraphs)
    check("20 Shoreward Drive" in text and "Gqeberha" in text, "the tender's own tender-box address is printed when on file")


def test_intake_readers():
    print("Intake readers (classify_tender.py)")
    import classify_tender as c
    notice = ("CLOSING DATE\n\nThe closing time and date for receipt of tenders is 12:00 on 22 October 2026\n\n"
              "Tenders must be placed inside the tender box at:\n\nThe Reception Area\n20 Shoreward Drive\nBay West\n"
              "Gqeberha\n6025\n                    T-6\n")
    a = c.find_delivery_address(notice)
    check(a and a["lines"] == ["The Reception Area", "20 Shoreward Drive", "Bay West", "Gqeberha", "6025"]
          and a["ambiguous"] is False,
          "tender-box address read from R.396's own Tender Notice wording, not flagged ambiguous (only one found)")
    t = c.find_closing_time(notice)
    check(t == {"value": "12:00", "ambiguous": False}
          and c.find_closing_time("The closing date and time is 11h00 on 3 June") == {"value": "11:00", "ambiguous": False},
          "closing time read from the notice (12:00 and 11h00 forms), not flagged ambiguous (only one found)")
    check(c.find_delivery_address("No address here.") is None, "no address -> None, never a default")

    # Finding 61, QA pass 2026-10-23: a real, common shape - an addendum
    # extending the deadline and/or moving the tender box venue, on top of
    # Book 1's own original notice. collect_documents() feeds documents in
    # alphabetical filename order, not chronological order, so whichever one
    # happens to sort first in all_text is NOT a reliable "the original" or
    # "the correcting addendum" signal either way - confirmed directly below
    # by swapping which section comes first and watching the answer flip.
    addendum = ("ADDENDUM No 1\n\nThe closing time and date for receipt of tenders is 10:00 on 29 October 2026\n\n"
                "Tenders must be placed inside the tender box situated at:\n\nSANRAL Southern Region Office\n"
                "Corner Plein and Govan Mbeki\nHumewood\nGqeberha\n6001\n")
    book1_first = c.find_closing_time(notice + "\n" + addendum)
    addendum_first = c.find_closing_time(addendum + "\n" + notice)
    check(book1_first == {"value": "12:00", "ambiguous": True} and addendum_first == {"value": "10:00", "ambiguous": True},
          f"two different closing times across the documents: which one wins depends purely on filename sort order "
          f"({book1_first} vs {addendum_first}), so both are correctly flagged ambiguous=True rather than trusted silently")
    addr_book1_first = c.find_delivery_address(notice + "\n" + addendum)
    addr_addendum_first = c.find_delivery_address(addendum + "\n" + notice)
    check(addr_book1_first["lines"][0] == "The Reception Area" and addr_book1_first["ambiguous"] is True
          and addr_addendum_first["lines"][0] == "SANRAL Southern Region Office" and addr_addendum_first["ambiguous"] is True,
          "two different delivery addresses: which one wins again depends purely on sort order, both flagged ambiguous")


def test_financial_rows():
    print("Financial workbook rows (fill_financial.py)")
    ok = True
    for f in sorted(os.listdir(FIN_DIR)):
        wb = Workbook(os.path.join(FIN_DIR, f))
        ok &= fill_financial.row_for_position(wb, "PL") == 12 and fill_financial.row_for_position(wb, "DS1") == 13
        ok &= fill_financial.row_for_position(wb, "DS2") is None
    check(ok, "TENDERER INPUT DATA rows found by the sheet's own position codes on all 6 samples (PL 12, DS1 13)")


def test_general_fillers():
    print("General forms (fill_sbd_docx.py / fill_flat_pdf.py)")
    import fill_sbd_docx
    import fill_flat_pdf
    comp = {"bee_proof_type": "qse_affidavit", "bee_proof_expiry": "2026-10-01", "legal_name": "X CC"}
    check(fill_sbd_docx.bee_proof_lapse(comp, "2026-10-22") is not None, "SBD 1: proof lapsing before closing is detected")
    v = fill_flat_pdf.values_for({"company": comp, "tender": {"closing_date": "2026-10-22"}})
    check(v["bee_affidavit"] is None and v["bee_certificate"] is None, "flat SBD 1: no B-BBEE box ticked for lapsed proof")
    comp["bee_proof_expiry"] = "2027-10-01"
    v = fill_flat_pdf.values_for({"company": comp, "tender": {"closing_date": "2026-10-22"}})
    check(v["bee_affidavit"] == "yes", "flat SBD 1: valid affidavit still ticked")
    contact = [f for f in fill_flat_pdf.FORMS["1"]["fields"] + fill_flat_pdf.FORMS["ceq"]["fields"] if f["key"].endswith("contact_person")]
    check(len(contact) == 2 and all(f.get("assumed_reason") for f in contact),
          "SBD 1 / CEQ contact person = signatory is an assumed_default, not confirmed")
    import build_general_pack
    rows = dict((r[0], r) for r in build_general_pack.guide_rows("4", {"company": {}, "answers": {}, "signatory": {}}))
    check("Interest in any other related enterprise, bidding or not?" in rows,
          "completion guide carries SBD 4's real wording (was 'Interest in a related enterprise bidding?')")


def test_today_iso_is_sast_not_naive_utc():
    print("_today_iso() (fill_a_forms.py / fill_sbd_docx.py) uses SAST, not a naive UTC date")
    # Finding (2026-10-20): both files computed "today" via naive
    # datetime.date.today() - a bridge worker container has no TZ set and
    # runs in UTC by default, so this is one calendar day behind the real
    # SAST date for the whole 00:00-01:59 SAST window every day. In
    # fill_a_forms.py this silently flipped A12/A20's "is this B-BBEE proof /
    # COID letter still current" verdict from expired to current for a
    # document expiring on the real SAST today. Simulates the system clock
    # at UTC 2026-10-20 23:30 (= SAST 2026-10-21 01:30, inside the bad
    # window) by monkeypatching datetime.datetime.now, restored immediately
    # after - never leaves the real clock patched.
    import fill_sbd_docx
    import datetime as real_datetime

    class FixedDateTime(real_datetime.datetime):
        @classmethod
        def now(cls, tz=None):
            base = real_datetime.datetime(2026, 10, 20, 23, 30, tzinfo=real_datetime.timezone.utc)
            return base.astimezone(tz) if tz else base

    for mod, label in ((fill_a_forms, "fill_a_forms"), (fill_sbd_docx, "fill_sbd_docx")):
        real_dt = mod.datetime.datetime
        mod.datetime.datetime = FixedDateTime
        try:
            today = mod._today_iso()
        finally:
            mod.datetime.datetime = real_dt
        check(today == "2026-10-21", f"{label}._today_iso() reads the real SAST date (2026-10-21), "
              f"not the naive UTC one (2026-10-20) - got {today}")

    # End-to-end, through the REAL fill_a_forms.main() against the real
    # template: a B-BBEE proof expiring exactly on the real SAST today must
    # be reported as expired, not current, even while the clock sits in the
    # bad window.
    prof = copy.deepcopy(UNLIKE)
    prof["company"].update(bee_level=1, bee_proof_type="certificate", bee_proof_expiry="2026-10-21",
                           black_ownership_pct=51, status="EME")
    real_dt = fill_a_forms.datetime.datetime
    fill_a_forms.datetime.datetime = FixedDateTime
    try:
        es, _ = a_forms(prof, tempfile.mkdtemp())
    finally:
        fill_a_forms.datetime.datetime = real_dt
    check(es["a_forms.A12.bee_certificate"]["status"] == "missing"
          and "expired 2026-10-21" in es["a_forms.A12.bee_certificate"]["reason"],
          f"a B-BBEE proof expiring on the real SAST today (2026-10-21) is reported expired, "
          f"not confirmed-current, even from inside the bad UTC-midnight window ({es.get('a_forms.A12.bee_certificate')})")


if __name__ == "__main__":
    for fn in (test_a_forms, test_docx_forms, test_technical_workbook, test_covers, test_intake_readers,
               test_financial_rows, test_general_fillers, test_today_iso_is_sast_not_naive_utc):
        fn()
    print(f"\n{len(failures)} failure(s)")
    sys.exit(1 if failures else 0)
