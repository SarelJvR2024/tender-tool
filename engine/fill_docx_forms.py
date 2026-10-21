"""
Fills the mechanical, non-judgment fields of the C.1.1.1 Form of Offer and the
C.1.2.3 Contract Data form. The A-Forms bundle (A2-A23, including A18 bank
details) is handled separately by fill_a_forms.py.

Every field filled here is filled with a real, final value - never a bracketed
"[DRAFT]"/"[PENDING]" placeholder typed into the document itself (that reads as
the tenderer not having finished their own paperwork, in a document SANRAL will
actually receive). A field we genuinely don't have data for, or that needs a
real human act, is left blank exactly as the template already has it, and
listed instead on the ACTIONS_TO_SIGN.md checklist this pipeline writes
alongside the filled documents - see fill_a_forms.py's module docstring for why
that split exists.

Deliberately does NOT fill:
  - Any price field (PRICE OFFERED on C.1.1.1) - that's the BOQ, a business decision.
  - Any signature, date-of-signing, or witness field - these need a real human signature
    event, never text. Witness fields specifically are left blank rather than filled
    with the signatory's own name (confirmed by testing: a naive "fill every NAME
    (IN CAPITALS) field" approach wrongly duplicated the signatory into the WITNESS
    field too - a witness must be a different real person, never fabricated).
  - Form A2 (Authority of Signatory) - the template itself states in its own notes that
    submitting this example page completed "shall not be accepted as authorisation"; it
    needs a real board resolution on the company's letterhead, not this form filled in.
  - Form A12 (B-BBEE) - needs the tenderer's actual verification certificate/affidavit
    attached, not text typed into this template.

Usage:
    python3 fill_docx_forms.py <templates_dir> <profile.json> <output_dir>

<templates_dir> must contain C.1.1.1_Form_of_Offer.docx and C.1.2.3_Contract_Data.docx,
fetched from the real tender's Drive folder (not reconstructed from memory - these
forms vary per contract, e.g. different NRA reference numbers in the headers).
"""
import re
import sys
import json
import docx
import yaml
from provenance import ProvenanceLog


# Book 1's key-person titles (R.396-010-2025/1F T-24 abbreviation list, plus
# the C3.1.11 wording "Contract Engineer" / "Engineer's Representative (RE)"),
# used to find each position's row in C.1.2.3's key-person table BY ITS LABEL.
# Until 2026-10-17 the table was filled by fixed row index - row 1 = PL and
# row 3 = DS, both with the signatory - which is only right for a tender whose
# positions are exactly PL + DS1 and a company whose signatory holds both
# (all six samples, and Klipcorp). Any other SANRAL tender lists its own
# positions (C3.1.11 a: "The key persons required for this project are listed
# in Form C1.2.2"), and any other company confirms its own people per position.
POSITION_TITLES = {
    "PL": ["project leader"], "APL": ["assistant project leader"],
    "DS": ["design specialist"], "DA": ["design assistant"],
    "CE": ["contracts engineer", "contract engineer"], "ACE": ["assistant contracts engineer", "assistant contract engineer"],
    "RE": ["resident engineer", "engineer's representative", "engineer’s representative"],
    "ARE": ["assistant resident engineer", "assistant re"],
    "RM": ["route manager"], "ARM": ["assistant route manager"],
    "SMT": ["senior materials technician"], "CM": ["contracts manager"], "SA": ["site agent"],
    "MT": ["materials technician"], "ETS": ["electronics systems engineer", "electronic systems engineer"],
    "ES": ["electrical systems engineer"], "MS": ["mechanical systems engineer"],
}


def _base_code(code):
    return re.sub(r"\d+$", "", str(code or "").strip().upper())


def _row_code(label):
    """Book 1 position code a C.1.2.3 row label names, or None. Longest title
    wins, so "Assistant Project Leader" is APL, not PL. "Alternate to ..."
    rows return None (the alternate is optional, "if applicable")."""
    text = re.sub(r"\s+", " ", label.strip().lower())
    if text.startswith("alternate"):
        return None
    best = None
    for code, titles in POSITION_TITLES.items():
        for t in titles:
            if text.startswith(t) and (best is None or len(t) > best[1]):
                best = (code, len(t))
    return best[0] if best else None


def key_persons_by_row(table, assignments):
    """{table row index: (person, position code)} for C.1.2.3's key-person
    table. `assignments` = [(position code, person dict)] in SETUP order, as
    the confirmed pack lists them. Several positions sharing a base code (DS1,
    DS2) fill that code's rows in order."""
    queues = {}
    for code, person in assignments:
        queues.setdefault(_base_code(code), []).append((code, person))
    out = {}
    for i, row in enumerate(table.rows):
        if i == 0:
            continue
        code = _row_code(row.cells[0].text)
        if code and queues.get(code):
            pos, person = queues[code].pop(0)
            out[i] = (person, pos)
    return out


def fill_contract_data(src_dir, out_dir, comp, signatory, key_persons, log):
    d = docx.Document(f"{src_dir}/C.1.2.3_Contract_Data.docx")
    cell = d.tables[0].rows[0].cells[0]

    fills = {
        "The Service Provider is \t": comp["legal_name"],
        "Address: \t": comp["registered_address"],
        "Name: \t": signatory["name"] + f" ({signatory['signatory_capacity']})",
    }
    telephone_seen = 0
    for p in cell.paragraphs:
        if p.text.startswith("……………"):
            # this dotted line IS the postal-address blank, not a label to append after.
            # Found by the admin harness (spec item 30) running a brand-new test
            # company with no postal address on file - the blank was left untouched
            # (not even an empty string write) rather than crashing python-docx's
            # run-text setter on None, which is what happened before this fix.
            if comp.get("postal_address"):
                p.runs[0].text = comp["postal_address"]
                for r in p.runs[1:]:
                    r.text = ""
                log.confirmed("c123.postal_address", comp["postal_address"])
            else:
                log.missing("c123.postal_address", "No postal address on file.")
            continue
        if p.text == "Telephone: \t":
            telephone_seen += 1
            if telephone_seen == 2 and signatory.get("phone"):
                p.add_run(signatory["phone"])
                log.confirmed("c123.representative_telephone", signatory["phone"])
            elif telephone_seen == 2:
                log.missing("c123.representative_telephone", f"{signatory['name']}'s telephone number is not on file.")
            elif comp.get("telephone"):
                p.add_run(comp["telephone"])
                log.confirmed("c123.company_telephone", comp["telephone"])
            else:
                log.missing("c123.company_telephone", "The company's telephone number is not on file.")
            continue
        if p.text == "E-mail Address: \t":
            if signatory.get("email"):
                p.add_run(signatory["email"])
                log.confirmed("c123.representative_email", signatory["email"])
            else:
                log.missing("c123.representative_email", f"{signatory['name']}'s email address is not on file.")
            continue
        if p.text in fills:
            label = p.text
            value = fills[label]
            if value:
                p.add_run(value)
                log.confirmed(f"c123.{_slug(label)}", value)
            else:
                log.missing(f"c123.{_slug(label)}", f"No value on file for '{label.strip()}'.")

    kp_table = d.tables[1]
    rows = key_persons_by_row(kp_table, key_persons) if key_persons is not None else {}
    for row_idx, row in enumerate(kp_table.rows):
        if row_idx == 0:
            continue
        label = row.cells[0].text.strip()
        code = _row_code(label)
        if code is None:
            # _row_code() returns None for two different reasons that must
            # not be treated the same: a genuinely optional "Alternate to
            # ... (if applicable)" row (left blank, as before - no trace
            # needed, this is the documented, intended shape of the form),
            # versus a real key-position row whose title just doesn't match
            # anything in POSITION_TITLES (a future SANRAL tender introducing
            # a title this dict doesn't recognise yet). The second case used
            # to be silently skipped too - identical to the first - so the
            # row would ship blank in the real submission with NOTHING in
            # the provenance report telling anyone to fill it by hand. Found
            # during this session's enterprise-readiness audit: a brand-new
            # company's very first tender would hit this with no warning at
            # all, since admin_harness.py's own baseline-diff treats a first
            # run as "clean, baseline stored" unconditionally.
            if re.sub(r"\s+", " ", label.lower()).startswith("alternate"):
                continue
            log.missing(f"c123.key_person_row_{row_idx}",
                        f"'{label}' doesn't match any known SANRAL key-position title - "
                        "this row was left blank. Check the position list in fill_docx_forms.py's "
                        "POSITION_TITLES and add it if this is a real, currently-used title.")
            continue
        if row_idx not in rows:
            log.missing(f"c123.key_person_row_{row_idx}",
                        f"No key person confirmed for '{label}' in this pack - fill the name by hand."
                        if key_persons is not None else
                        f"No confirmed key-person assignment (pack.yaml) was given - fill '{label}' by hand.")
            continue
        person, pos = rows[row_idx]
        row.cells[1].text = person["name"]
        status = person.get("employment_status")
        if status == "not_full_time_employee":
            row.cells[2].text = "Not permanent"
            log.confirmed(f"c123.key_person_row_{row_idx}", person["name"])
        else:
            row.cells[2].text = "Permanent"
            if status == "full_time_employee":
                log.confirmed(f"c123.key_person_row_{row_idx}", person["name"])
            else:
                log.assumed_default(f"c123.key_person_row_{row_idx}", person["name"],
                                    f"{person['name']} ({pos}) written as Permanent - whether they are a full-time "
                                    "employee is not recorded on their People page; confirm before signing.")

    # "Name of Partners, Members or Directors | Appointment Date | Designation |
    # Professional Registration Number". The company record holds ONE director
    # (companies.director_name). Until 2026-10-18 that one row was written and
    # logged confirmed even when the name was blank, and nothing said the rest
    # of the table (every other director/member/partner, the dates, the
    # registration numbers) was left empty - complete only for a one-director
    # company such as Klipcorp.
    dir_table = d.tables[2]
    director = ((comp.get("director") or {}).get("name") or "").strip()
    principals = [p for p in (comp.get("principals") or []) if (p.get("name") or "").strip()]
    is_cc = bool(re.search(r"\bCC\b|close\s+corporation", comp.get("legal_name") or "", re.IGNORECASE))
    if principals:
        # The company's own directors / members register (company_principals,
        # 0050), every row the table has.
        rows = dir_table.rows[1:]
        for row, pr in zip(rows, principals):
            designation = pr.get("designation") or ("Member" if is_cc else "Director")
            row.cells[0].text = pr["name"]
            row.cells[1].text = str(pr.get("appointed_on") or "")
            row.cells[2].text = designation
            row.cells[3].text = pr.get("registration_number") or ""
        log.confirmed("c123.directors_table", "; ".join(p["name"] for p in principals[:len(rows)]))
        undated = [p["name"] for p in principals[:len(rows)] if not p.get("appointed_on")]
        if undated:
            log.missing("c123.directors_appointment_dates",
                        f"No appointment date on the directors register for {', '.join(undated)} - add it there or by hand.")
        if len(principals) > len(rows):
            log.missing("c123.directors_table_complete",
                        f"{len(principals)} directors/members on the register, the form has {len(rows)} rows - list "
                        f"{', '.join(p['name'] for p in principals[len(rows):])} on an attached page.")
    elif director:
        dir_table.rows[1].cells[0].text = director
        dir_table.rows[1].cells[2].text = "Member" if is_cc else "Director"
        log.confirmed("c123.director_name", director)
        log.computed("c123.director_designation", "Member" if is_cc else "Director",
                     "From the company record's director field" + (" (a close corporation's are members)" if is_cc else "") + ".")
        log.missing("c123.directors_table_complete",
                    f"Only one director/member is on file ({director}) and the directors register (Company page) is "
                    "empty - add every other partner, member or director, and each one's appointment date and "
                    "professional registration number, by hand or on the register.")
    else:
        log.missing("c123.director_name", "No director / member / partner on file - complete the table by hand.")

    d.save(f"{out_dir}/C.1.2.3_Contract_Data_FILLED.docx")


def _slug(label):
    return label.strip().rstrip(":\t").lower().replace(" ", "_")


def fill_form_of_offer(src_dir, out_dir, comp, signatory, log):
    d = docx.Document(f"{src_dir}/C.1.1.1_Form_of_Offer.docx")
    name_in_caps_count = 0
    for p in d.paragraphs:
        if p.text.startswith("I/we claim the following B-BBEE contributor status level"):
            if not comp.get("bee_level"):
                # Before 2026-10-18 a company with no level got " None " typed
                # into the claim, logged confirmed.
                log.missing("c111.bee_level", "No B-BBEE level on file - the claimed contributor status level is "
                                              "left blank; fill it by hand from your certificate/affidavit (or leave it "
                                              "if you claim none).")
                continue
            for r in p.runs:
                if ".........." in r.text:
                    r.text = r.text.replace("..........", f" {comp['bee_level']} ", 1)
                    log.confirmed("c111.bee_level", comp["bee_level"])
        if p.text.startswith("NAME AND ADDRESS OF ORGANISATION:"):
            value = ", ".join(x for x in (comp["legal_name"], comp.get("registered_address")) if x)
            p.add_run(f"  {value}")
            log.confirmed("c111.organisation_name_address", value)
            if not comp.get("registered_address"):
                log.missing("c111.organisation_address", "No registered address on file - add it after the name by hand.")
        if p.text.startswith("NAME (IN CAPITALS): "):
            name_in_caps_count += 1
            if name_in_caps_count == 1:
                p.add_run(f"  {signatory['name'].upper()}")
                log.confirmed("c111.signatory_name", signatory["name"].upper())
            # witness field (2nd+ occurrence) stays blank - a witness must be a real,
            # different person; see module docstring.
        if p.text.startswith("CAPACITY: "):
            p.add_run(f"  {signatory['signatory_capacity']}")
            log.confirmed("c111.signatory_capacity", signatory["signatory_capacity"])
    d.save(f"{out_dir}/C.1.1.1_Form_of_Offer_FILLED.docx")
    log.missing("c111.witness_name_signature",
                "Needs a witness to sign and print their name - must be a different real person "
                "from the signatory, never fabricated.")


def assignments_from_pack(pack_path):
    """[(position code, person dict)] from a confirmed pack.yaml, in the order
    the pack lists its positions (SETUP order); None when no pack is given."""
    if not pack_path:
        return None
    pack = yaml.safe_load(open(pack_path))
    people = pack.get("people", {})
    return [(a["position"], people[a["person"]]) for a in pack.get("key_persons", []) if a.get("person") in people]


def main(src_dir, profile_path, out_dir, provenance_out=None, pack_path=None):
    profile = json.load(open(profile_path))
    comp = profile["company"]
    signatory = profile["key_people"][0]  # the signatory (same convention as fill_a_forms.py)
    # Key persons come from the CONFIRMED pack (who the company put on each of
    # this tender's positions), never assumed to be the signatory. Without a
    # pack, process_pack_jobs.py's profile["role_people"] ({base code:
    # person}) is used when present; with neither, the rows are left blank and
    # logged missing.
    key_persons = assignments_from_pack(pack_path)
    if key_persons is None and profile.get("role_people"):
        key_persons = list(profile["role_people"].items())
    log = ProvenanceLog(source="fill_docx_forms.py")

    fill_contract_data(src_dir, out_dir, comp, signatory, key_persons, log)
    fill_form_of_offer(src_dir, out_dir, comp, signatory, log)
    print(f"Wrote 2 filled docx forms to {out_dir}")
    if provenance_out:
        log.save(provenance_out)
    return log


if __name__ == "__main__":
    main(*sys.argv[1:5])
