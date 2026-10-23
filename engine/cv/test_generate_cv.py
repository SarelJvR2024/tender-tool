"""
Focused tests for generate_cv.py's cover page (add_cover_page()) - the
tender-specific "this is why this person fits this position" page built on
top of suggest_pack.py's own stored scoring (never a new matcher). No
dedicated test file existed for generate_cv.py before this (the standard,
non-tender-specific CV path is already smoke-tested end to end by
process_cv_jobs.py --self-test; this file covers the --tender/--position
cover-page path specifically, which that self-test doesn't exercise).

Finding (2026-10-19, thirtieth run, docs/specs/QA-Edge-Case-Findings.md): when
a person wasn't among suggest_pack.py's own scored candidates for a position,
add_cover_page()'s own "missing" provenance reason interpolated
`tender_ctx['candidate']!r` - which is exactly the `None` this branch is
guarded on, so the stored reason always rendered as the nonsensical
"<name> not present in None candidates list for position <code>..." on
literally every occurrence. Fixed to describe the real situation instead.

Run: python3 engine/cv/test_generate_cv.py
"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.dirname(HERE))
import generate_cv as gc  # noqa: E402
from provenance import ProvenanceLog  # noqa: E402
import docx  # noqa: E402

FAILS = []


def check(name, cond, detail=""):
    print(("  ok   " if cond else "  FAIL ") + name + (f"  ({detail})" if detail and not cond else ""))
    if not cond:
        FAILS.append(name)


def _tender_ctx(candidate):
    return {
        "contract_no": "X.001-001-2026/1F", "closing_date": "2026-12-01",
        "position": {"code": "DS1", "description": "Structures specialist"},
        "candidate": candidate,
    }


def test_unmatched_candidate_reason_is_not_nonsensical():
    print("cover page: person not among suggest_pack.py's scored candidates for this position")
    d = docx.Document()
    log = ProvenanceLog(source="test")
    person = {"name": "Jane Doe"}
    gc.add_cover_page(d, person, "person_3", _tender_ctx(None), log)
    entry = next(e for e in log.entries if e["field_key"] == "cv_cover.position_match")
    check("the real bug: reason no longer says 'not present in None candidates list'",
          "None candidates" not in entry["reason"], entry["reason"])
    check("the fixed reason names the actual person_id and position code instead",
          "person_3" in entry["reason"] and "DS1" in entry["reason"], entry["reason"])
    check("status is still 'missing' (unchanged - only the wording was wrong)",
          entry["status"] == "missing")


def test_matched_candidate_logs_score_not_rationale_on_page():
    print("cover page: person IS a scored candidate for this position")
    d = docx.Document()
    log = ProvenanceLog(source="test")
    person = {"name": "Jane Doe"}
    candidate = {"specialism_warning": False, "rationale": "Matches DS1 keyword overlap", "score": 0.82}
    gc.add_cover_page(d, person, "person_3", _tender_ctx(candidate), log)
    check("match score is confirmed in provenance",
          any(e["field_key"] == "cv_cover.match_score" and e["value"] == "0.82" for e in log.entries))
    check("rationale is kept in provenance only (never printed on the buyer-facing page)",
          not any("Matches DS1 keyword overlap" in (p.text or "") for p in d.paragraphs))
    check("no specialism-warning paragraph for a candidate with no warning",
          not any("does not show direct evidence" in (p.text or "") for p in d.paragraphs))


def test_specialism_warning_shown_and_logged():
    print("cover page: a candidate flagged with no specialism evidence")
    d = docx.Document()
    log = ProvenanceLog(source="test")
    person = {"name": "Jane Doe"}
    candidate = {"specialism_warning": True, "rationale": "No keyword overlap found", "score": 0.3}
    gc.add_cover_page(d, person, "person_3", _tender_ctx(candidate), log)
    check("the warning paragraph is actually printed on the page",
          any("does not show direct evidence" in (p.text or "") for p in d.paragraphs))
    check("the warning is logged as assumed_default, not silently applied",
          any(e["field_key"] == "cv_cover.specialism_warning" and e["status"] == "assumed_default"
              for e in log.entries))


def test_person_name_cannot_escape_out_dir_via_path_traversal():
    print("output filename: a person name containing path-traversal characters")
    import tempfile
    out_dir = tempfile.mkdtemp(prefix="test_cv_out_")
    base_name = gc._safe_name("../../../../tmp/pwned_poc_cv")
    docx_path = os.path.join(out_dir, base_name + ".docx")
    real_out, real_docx = os.path.realpath(out_dir), os.path.realpath(docx_path)
    check("the resolved output path stays inside out_dir (no CWE-22 path traversal)",
          real_docx.startswith(real_out + os.sep), real_docx)


def test_null_projects_and_qualifications_degrade_instead_of_crashing():
    print("a person row with explicit null projects/qualifications (not just missing)")
    person = {"name": "Jane Doe", "projects": None, "qualifications": None}
    profile = {"company": {"legal_name": "Test Co"}}
    log = ProvenanceLog(source="test")
    try:
        gc.build_world_bank(profile, person, "person_1", "summary", None, None, log)
        check("build_world_bank does not crash on projects=None/qualifications=None", True)
    except Exception as e:
        check("build_world_bank does not crash on projects=None/qualifications=None", False, f"{type(e).__name__}: {e}")


def test_control_characters_are_stripped_not_crashed_on():
    print("a control character (realistic OCR artifact) in name/project text")
    person = {"name": "Jane\x0bDoe", "qualifications": {},
              "projects": [{"client_and_number": "Proj\x01A", "value_rm": 5, "position_held": "Engineer",
                             "start": "2020", "end": "2021", "contact": "Bob", "contact_phone": "012"}]}
    profile = {"company": {"legal_name": "Test Co"}}
    log = ProvenanceLog(source="test")
    import tempfile
    try:
        doc = gc.build_world_bank(profile, person, "person_1", "summary", None, None, log)
        with tempfile.TemporaryDirectory() as d:
            doc.save(os.path.join(d, "test_control_char.docx"))
        check("a control character in profile text does not crash docx generation", True)
    except Exception as e:
        check("a control character in profile text does not crash docx generation", False, f"{type(e).__name__}: {e}")


def test_years_experience_uses_sast_not_naive_utc_today():
    print("Form B2.1: years-of-experience fallback 'as of today' uses SAST, not a naive UTC date")
    # Finding (2026-10-21): the standard (non-tender) CV path's fallback
    # "as of today" date was a naive datetime.date.today() - a bridge worker
    # container has no TZ set and runs in UTC by default, so this is one
    # calendar day behind real SAST for the 00:00-01:59 SAST window every
    # day (same bug class as Finding 37, engine/fill_a_forms.py). Fed
    # straight into suggest_pack.years_experience_at()'s rounded year count
    # (round((closing_date - grad).days / 365.25)) - a one-day shift flips
    # the rounded figure whenever a person's initial_graduated date sits
    # near the rounding boundary. "2026-04-21" is confirmed (by direct
    # search) to be exactly such a date for this SAST/UTC pair. Simulates
    # the system clock at UTC 2026-10-20 23:30 (= SAST 2026-10-21 01:30,
    # inside the bad window) by monkeypatching datetime.datetime.now,
    # restored immediately after.
    import datetime as real_datetime

    class FixedDateTime(real_datetime.datetime):
        @classmethod
        def now(cls, tz=None):
            base = real_datetime.datetime(2026, 10, 20, 23, 30, tzinfo=real_datetime.timezone.utc)
            return base.astimezone(tz) if tz else base

    class FixedDate(real_datetime.date):
        @classmethod
        def today(cls):
            # What a naive datetime.date.today() would read under the same
            # fake UTC clock as FixedDateTime above - patched too so this
            # test still catches a regression to the OLD naive form, not
            # just a hypothetical one that happens to keep using datetime.now.
            return real_datetime.date(2026, 10, 20)

    profile = {"company": {"legal_name": "Test Co", "cipc_registration_number": "", "registered_address": ""}}
    person = {"name": "Test Person", "role": "Engineer",
              "qualifications": {"initial_graduated": "2026-04-21"}, "projects": []}
    real_dt, real_date = gc.datetime.datetime, gc.datetime.date
    gc.datetime.datetime, gc.datetime.date = FixedDateTime, FixedDate
    try:
        log = ProvenanceLog(source="test")
        gc.build_sanral_keyperson(profile, person, "p1", "summary", None, None, log)
    finally:
        gc.datetime.datetime, gc.datetime.date = real_dt, real_date
    entry = next(e for e in log.entries if e["field_key"] == "cv.b21.years_of_relevant_experience")
    check("years-of-experience is computed from the real SAST date (1), not the naive UTC one (0)",
          entry["value"] == "1", entry)


def test_world_bank_project_with_null_fields_does_not_print_the_word_none():
    print("build_world_bank: an ongoing project (end_date=None) and a missing referee/value "
          "must not render the literal word 'None' into a real, client-facing CV")
    # Finding (2026-10-21): dict.get(key, default) only substitutes `default`
    # when the key is ABSENT, not when it's present with value None -
    # start_date/end_date/value_rm/contact_name/contact_phone are all
    # nullable columns (0001_init_schema.sql's projects table), so a real
    # ongoing project (no end date yet - the single most common real shape,
    # true for almost any currently-employed person's latest project) or a
    # referee not yet captured sailed straight through as the bare string
    # "None" in the generated document. build_sanral_keyperson (same file,
    # same data) already guards the identical fields with `or` - this
    # sibling builder didn't.
    person = {"name": "Jane Doe", "role": "Engineer", "qualifications": {},
              "projects": [{"client_and_number": "Ongoing Project A/2022", "position_held": "Site Engineer",
                             "start": "2022-01-01", "end": None, "value_rm": None,
                             "contact": None, "contact_phone": None}]}
    profile = {"company": {"legal_name": "Test Co"}}
    log = ProvenanceLog(source="test")
    doc = gc.build_world_bank(profile, person, "person_1", "full", None, None, log)
    text = next(p.text for p in doc.paragraphs if "Ongoing Project" in p.text)
    check("the open-ended project's end date reads '?', not the word 'None'",
          "None" not in text, text)
    check("a missing referee still reads 'MISSING', not 'None'", "MISSING" in text, text)
    check("the start date is still shown correctly", "2022-01-01" in text, text)


def test_sanral_keyperson_project_with_null_position_held_does_not_crash():
    print("build_sanral_keyperson: a real project row with no captured position title "
          "(position_held=None) must not crash CV generation for that person")
    # Finding (2026-10-23): dict.get(key, default) only substitutes `default`
    # when the key is ABSENT, not when it's present with value None -
    # position_held is a nullable column (0001_init_schema.sql's projects
    # table), and _safe_text(None) itself crashes (re.sub needs a string,
    # not None) - every OTHER nullable project field in this same function
    # already guards with `or ""` (the "or" pattern that fixed the sibling
    # build_world_bank() bug, Finding 49), position_held was the one left
    # behind. Worse than Finding 49's "prints the word None": this one
    # failed the whole CV generation job, not just rendered wrong text.
    person = {"name": "Jane Doe", "role": "Engineer", "qualifications": {},
              "projects": [{"client_and_number": "Real Project A/2022", "position_held": None,
                             "start": "2022-01-01", "end": "2023-01-01", "value_rm": 10,
                             "contact": "A Buyer", "contact_phone": "012 345 6789"}]}
    profile = {"company": {"legal_name": "Test Co"}}
    log = ProvenanceLog(source="test")
    doc = gc.build_sanral_keyperson(profile, person, "person_1", "full", None, None, log)
    row = next(r for t in doc.tables for r in t.rows if any("Real Project" in c.text for c in r.cells))
    cells = [c.text for c in row.cells]
    check("the project renders at all (no crash) and the position cell is blank, not the word 'None'",
          cells[2] == "", cells)
    check("the rest of the row is unaffected (client/value/start/end/contact all still correct)",
          cells[0] == "Real Project A/2022" and cells[1] == "10" and cells[3] == "2022-01-01"
          and cells[4] == "2023-01-01" and cells[5] == "A Buyer", cells)


def main():
    test_unmatched_candidate_reason_is_not_nonsensical()
    test_matched_candidate_logs_score_not_rationale_on_page()
    test_specialism_warning_shown_and_logged()
    test_person_name_cannot_escape_out_dir_via_path_traversal()
    test_null_projects_and_qualifications_degrade_instead_of_crashing()
    test_control_characters_are_stripped_not_crashed_on()
    test_years_experience_uses_sast_not_naive_utc_today()
    test_world_bank_project_with_null_fields_does_not_print_the_word_none()
    test_sanral_keyperson_project_with_null_position_held_does_not_crash()
    print(f"\n{len(FAILS)} failure(s)")
    sys.exit(1 if FAILS else 0)


if __name__ == "__main__":
    main()
