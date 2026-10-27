"""
Tests for engine/cv/extract_cv_fields.py. No live Supabase and no Anthropic
API key needed - and therefore, stated plainly, this does NOT test how well
Claude reads a CV. What it does test, for real:

  1. Text extraction on synthetic CVs built here (python-docx -> .docx,
     LibreOffice -> .pdf and legacy .doc) - the exact readers the worker uses.
  2. The failure modes: an image upload, an image-only (scanned) PDF, a
     garbled-encoding text, a document far too long to be a CV - each refused
     with a reason, never guessed at.
  3. SA ID numbers are found locally and are not in the text sent out.
  4. The checks applied to a model answer (validate_proposal): quotes not in
     the CV, digits not in the CV, year-only dates, values outside SANRAL's
     dropdown lists, "not a CV", nothing found, mostly-unverifiable answers.
     The answers are hand-written TEST DOUBLES standing in for the model.
  5. The request call_llm() sends (model, schema, fallbacks) and how it
     handles refusal / max_tokens / bad JSON - against a stub client.
  6. The queue plumbing (claim, success, put-back on missing key, failure
     reason, stale reclaim) against an in-memory stand-in for supabase_request.
  7. EXPERIMENTAL OCR (needs tesseract): a scan made from the synthetic CV is
     read by OCR, flagged, every value left unticked; an illegible one refused.

Real-sample text extraction (the Klipcorp SANRAL Form B1/B2 PDF and the
scanned certificates) is in CV-Extraction-Findings.md; those files are
gitignored PII so they can't be committed as fixtures.

Usage: python3 engine/cv/test_extract_cv_fields.py   (exit 0 = all passed)
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
import types

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import extract_cv_fields as x  # noqa: E402

FAILS = []


def check(cond, msg):
    print(("  ok   " if cond else "  FAIL ") + msg)
    if not cond:
        FAILS.append(msg)


def sa_id(prefix12):
    """Completes a 12-digit prefix with the Luhn check digit."""
    for d in "0123456789":
        if x._luhn_ok(prefix12 + d):
            return prefix12 + d


SYNTH_ID = sa_id("900215012308")  # synthetic: 1990-02-15, citizen, made-up sequence


def make_docx(path):
    import docx
    d = docx.Document()
    d.add_heading("CURRICULUM VITAE - Naledi Test-Mokoena", 0)
    d.add_paragraph("Senior Civil Engineer | naledi.test@example.test | +27 82 555 0101")
    d.add_paragraph(f"ID number: {SYNTH_ID}")
    d.add_heading("Professional registration", 1)
    d.add_paragraph("Pr Eng, Engineering Council of South Africa (ECSA), registration no. 20169999, registered 14 June 2016")
    d.add_paragraph("Member, South African Institution of Civil Engineering (SAICE)")
    d.add_heading("Qualifications", 1)
    d.add_paragraph("MSc (Eng) Transportation, University of Cape Town, 2014")
    d.add_paragraph("BSc (Eng) Civil, University of the Witwatersrand, graduated 3 December 2009")
    d.add_heading("Project experience", 1)
    t = d.add_table(rows=1, cols=5)
    for c, h in zip(t.rows[0].cells, ["Client / contract", "Role", "Value", "Period", "Referee"]):
        c.text = h
    for row in [["Test Metro Roads Agency - TMRA/2019/044 Rehabilitation of Main Road", "Resident Engineer",
                 "R120 million", "Mar 2019 - Oct 2021", "J. Referee-Test, 011 555 0199"],
                ["Test Province DoT - TPD/2022/7 Bridge widening", "Design Engineer", "R45 million",
                 "2022 - present", ""]]:
        cells = t.add_row().cells
        for c, v in zip(cells, row):
            c.text = v
    d.save(path)


def soffice_convert(src, ext, out_dir):
    subprocess.run(["soffice", "--headless", "--convert-to", ext, "--outdir", out_dir, src],
                   capture_output=True, timeout=180)
    out = os.path.join(out_dir, os.path.splitext(os.path.basename(src))[0] + "." + ext)
    return out if os.path.exists(out) else None


def expect_fail(fn, needle, msg):
    try:
        fn()
    except x.ExtractionFailed as e:
        needles = needle if isinstance(needle, tuple) else (needle,)
        check(any(n.lower() in str(e).lower() for n in needles), f"{msg} -> refused: {str(e)[:90]}...")
        return
    check(False, f"{msg} -> expected ExtractionFailed")


def test_text_extraction(tmp):
    print("1. text extraction on synthetic CVs")
    docx_path = os.path.join(tmp, "naledi_cv.docx")
    make_docx(docx_path)
    targets = [docx_path]
    have_soffice = shutil.which("soffice") is not None
    if have_soffice:
        targets += [soffice_convert(docx_path, "pdf", tmp), soffice_convert(docx_path, "doc", tmp)]
    else:
        print("  (soffice not installed: .pdf/.doc variants skipped)")
    for p in targets:
        check(p is not None and os.path.exists(p), f"fixture exists: {os.path.basename(str(p))}")
        r = x.read_cv_text(p, tmp)
        t = r["text"]
        # LibreOffice's PDF wraps the narrow table column mid-token ("TMRA/2019/04" / "4") - real
        # behaviour of PDF CVs, which is why quotes are also matched with whitespace removed. That
        # claim used to be aspirational only - this check did a bare `"TMRA/2019/04" in t`, with no
        # normalisation at all, so it was exactly as brittle to a LibreOffice point release wrapping
        # the same table cell a few pixels differently as the comment above it said it wasn't. This
        # is the real, previously-undiagnosed cause of bridge-worker-ci.yml's "GitHub-runner-only"
        # failure on this exact check (QA-Edge-Case-Findings.md, 2026-10-19) - reusing
        # evidence_in_text(), the same production-tested, wrap-tolerant matcher real CV extraction
        # already verifies every quote against, rather than a second, less tolerant check invented
        # just for this test.
        norm_t = x._norm(t)
        check(x.evidence_in_text("20169999", norm_t) and x.evidence_in_text("University of Cape Town", norm_t)
              and x.evidence_in_text("TMRA/2019/04", norm_t),
              f"{os.path.basename(p)} via {r['method']}: registration, institution and table cell all read")
        redacted, ids = x.find_and_redact_sa_ids(t)
        check(ids == [SYNTH_ID] and SYNTH_ID not in redacted, f"{os.path.basename(p)}: SA ID found locally and removed")
        check("20169999" in redacted and "0101" in redacted, f"{os.path.basename(p)}: ECSA/phone numbers NOT redacted")
    return docx_path


def test_failure_modes(tmp):
    print("2. unreadable files are refused with a reason")
    from PIL import Image, ImageDraw
    img = Image.new("RGB", (1240, 1754), "white")
    ImageDraw.Draw(img).text((100, 100), "CURRICULUM VITAE  Test Person  Pr Eng 20169999", fill="black")
    png = os.path.join(tmp, "photo_of_cv.png")
    img.save(png)
    scan_pdf = os.path.join(tmp, "scanned_cv.pdf")
    img.save(scan_pdf, "PDF", resolution=150)
    # With OCR installed this one line IS read (38 characters) and then refused
    # as far too little text for a CV; without OCR, refused as a scan.
    expect_fail(lambda: x.read_cv_text(png, tmp), ("photo or scanned image", "characters of text could be read"), "PNG upload")
    expect_fail(lambda: x.read_cv_text(scan_pdf, tmp), ("no text layer", "characters of text could be read"),
                "image-only PDF (a scan)")
    expect_fail(lambda: x.assess_text_quality("Name: X\n"), "only", "almost no text")
    expect_fail(lambda: x.assess_text_quality("□■#$%&*+=" * 60 + "ab"), "garbled", "garbled font encoding")
    expect_fail(lambda: x.assess_text_quality("word " * 40000), "far more than one person", "a whole tender, not a CV")
    expect_fail(lambda: x.read_cv_text(os.path.join(tmp, "cv.pages"), tmp), "can't be read", "unsupported format")
    check(not x._plausible_sa_id("2016999912345"), "a 13-digit non-ID number is not treated as an ID")

    # Finding (2026-10-27): the "docx" branch above already guards docx_text() against a
    # corrupt file - the "doc" branch's own call to the SAME function on LibreOffice's own
    # conversion output didn't. convert_with_soffice() only checks the output path EXISTS,
    # never that LibreOffice actually wrote a readable docx, so a broken conversion raised a
    # bare PackageNotFoundError straight out of read_cv_text() instead of the clean,
    # user-facing ExtractionFailed every other failure mode here gets - same bug class
    # independently found in requirements.py's read_document() and classify_documents.py's
    # read_document_text(). Simulated (no real .doc sample known to trigger this in
    # LibreOffice itself) by monkeypatching convert_with_soffice() to return a path to a
    # genuinely broken "conversion".
    broken_docx = os.path.join(tmp, "broken_conversion.docx")
    with open(broken_docx, "wb") as f:
        f.write(b"not a real docx file, just garbage bytes")
    real_convert = x.convert_with_soffice
    x.convert_with_soffice = lambda path, target_ext, out_dir: broken_docx
    try:
        expect_fail(lambda: x.read_cv_text(os.path.join(tmp, "cv.doc"), tmp),
                    "could not be read", "a .doc whose LibreOffice conversion produced a broken file")
    finally:
        x.convert_with_soffice = real_convert


GOOD_ANSWER = {  # TEST DOUBLE: what a correct model answer for make_docx()'s CV looks like
    "document_kind": "cv", "document_kind_reason": "One engineer's CV.",
    "person": {
        "name": {"value": "Naledi Test-Mokoena", "evidence": "CURRICULUM VITAE - Naledi Test-Mokoena"},
        "role": {"value": "Senior Civil Engineer", "evidence": "Senior Civil Engineer"},
        "phone": {"value": "+27 82 555 0101", "evidence": "+27 82 555 0101"},
        "email": {"value": "naledi.test@example.test", "evidence": "naledi.test@example.test"},
        "ecsa_level": {"value": "Professional Engineer", "evidence": "Pr Eng, Engineering Council of South Africa (ECSA)"},
        "ecsa_number": {"value": "20169999", "evidence": "registration no. 20169999"},
        "ecsa_date": {"value": "2016-06-14", "evidence": "registered 14 June 2016"},
        "sacpcmp_number": {"value": None, "evidence": None}, "sacpcmp_date": {"value": None, "evidence": None},
        "highest_level": {"value": "Masters Degree", "evidence": "MSc (Eng) Transportation"},
        "highest_description": {"value": "MSc (Eng) Transportation", "evidence": "MSc (Eng) Transportation"},
        "highest_institution": {"value": "University of Cape Town", "evidence": "University of Cape Town, 2014"},
        "highest_graduated": {"value": "2014", "evidence": "University of Cape Town, 2014"},
        "initial_description": {"value": "BSc (Eng) Civil", "evidence": "BSc (Eng) Civil"},
        "initial_institution": {"value": "University of the Witwatersrand", "evidence": "University of the Witwatersrand"},
        "initial_graduated": {"value": "2009-12-03", "evidence": "graduated 3 December 2009"},
    },
    "projects": [
        {"client_and_number": "Test Metro Roads Agency - TMRA/2019/044 Rehabilitation of Main Road",
         "position_held": "Resident Engineer", "value_rm": 120, "start_date": "2019-03", "end_date": "2021-10",
         "contact_name": "J. Referee-Test", "contact_phone": "011 555 0199", "contact_email": None,
         "evidence": "Test Metro Roads Agency - TMRA/2019/044 Rehabilitation of Main Road"},
        {"client_and_number": "Test Province DoT - TPD/2022/7 Bridge widening", "position_held": "Design Engineer",
         "value_rm": 45, "start_date": "2022", "end_date": None, "contact_name": None, "contact_phone": None,
         "contact_email": None, "evidence": "Test Province DoT - TPD/2022/7 Bridge widening"},
    ],
    "other_registrations": [{"body": "SAICE", "level": "Member", "number": None,
                             "evidence": "Member, South African Institution of Civil Engineering (SAICE)"}],
    "parse_problems": [],
}


def _person_as_wire_entries(person_dict):
    """GOOD_ANSWER["person"] is a dict (used throughout this file to test
    validate_proposal()'s own logic) - the real wire schema (output_schema())
    is an array of {field, value, evidence} entries (see extract_cv_fields.py's
    own 2026-10-23 comment on PERSON_FIELD_ENTRY_SCHEMA). Converts one to the
    other for the one test that checks a fixture against the real schema."""
    return [{"field": k, "value": v["value"], "evidence": v["evidence"]} for k, v in person_dict.items()]


def _count_union_params(schema):
    """What the API's own error actually counts (2026-10-23 re-review,
    QA-cv.pdf): 'N parameters with type arrays or anyOf' - only
    union/nullable-typed properties, not every property. A first pass at
    this check (`_count_schema_params`, counting every property regardless
    of union-ness) overstated the real number and understated how much
    `person`'s own fix already helped - it reported this schema at ~22
    "total params" and left the real, API-relevant number (12, see below)
    unverified. Counting only union-typed properties, recursively through
    `properties`, array `items` and `anyOf` branches, is what actually
    matches the provider's limit."""
    if not isinstance(schema, dict):
        return 0
    n = 0
    for v in (schema.get("properties") or {}).values():
        if isinstance(v, dict) and ("anyOf" in v or isinstance(v.get("type"), list)):
            n += 1
        n += _count_union_params(v)
    items = schema.get("items")
    if isinstance(items, dict):
        n += _count_union_params(items)
    for branch in schema.get("anyOf") or []:
        n += _count_union_params(branch)
    return n


def test_validation(docx_path, tmp):
    print("3. checks applied to a model answer (hand-written test doubles)")
    import jsonschema
    jsonschema.validate(dict(GOOD_ANSWER, person=_person_as_wire_entries(GOOD_ANSWER["person"])), x.output_schema())
    check(True, "array-shaped test-double answer is valid against the real structured-output schema sent to the API")
    n = _count_union_params(x.output_schema())
    # `person`'s own union count dropped from 32 (one value+evidence pair
    # per field, the same fixed-object-per-field bug confirmed live in
    # classify_documents.py's `fields`) to 2, by moving to an array of
    # {field, value, evidence} entries. `projects`/`other_registrations`
    # were deliberately left as typed per-field records rather than
    # flattened the same way (a shared per-project evidence quote,
    # `validate_proposal()`'s own verification logic built around that
    # shape) - but counted correctly (union-typed properties only, matching
    # the API's own "N parameters with type arrays or anyOf" wording, not
    # every property regardless of union-ness), their 8 + 2 union params are
    # well within the limit on their own. 12 total, confirmed under the
    # real 16-param limit the API enforces - not just "smaller than before".
    check(n < 16, f"schema has {n} union-typed parameters (API limit 16), down from 42 before `person` was fixed")
    text = x.read_cv_text(docx_path, tmp)["text"]
    _, ids = x.find_and_redact_sa_ids(text)
    prop, warnings = x.validate_proposal(json.loads(json.dumps(GOOD_ANSWER)), text, ids)
    p = prop["person"]
    check(p["ecsa_number"]["verified"] and p["ecsa_date"]["value"] == "2016-06-14", "full date + quoted number verified")
    check(p["highest_graduated"]["value"] == "2014-01-01" and not p["highest_graduated"]["verified"]
          and "only gives the year" in p["highest_graduated"]["warnings"][0], "year-only date: placeholder, flagged, not pre-ticked")
    check(p["id_number"]["value"] == SYNTH_ID, "ID number proposed from the local match")
    check(len(prop["projects"]) == 2 and prop["projects"][0]["start_date"] == "2019-03-01"
          and prop["projects"][0]["verified"], "projects kept, month-only start date normalised with a warning")
    check(prop["other_registrations"][0]["verified"], "other registration recorded (no column to save it yet)")
    pdf = os.path.join(tmp, "naledi_cv.pdf")
    if os.path.exists(pdf):
        ptext = x.read_cv_text(pdf, tmp)["verify_text"]
        pprop, _ = x.validate_proposal(json.loads(json.dumps(GOOD_ANSWER)), ptext, ids)
        ok_fields = sum(v["verified"] for v in pprop["person"].values())
        check(all(pr["verified"] for pr in pprop["projects"]) and pprop["person"]["ecsa_number"]["verified"],
              f"same answer against the PDF's wrapped table text: projects still traced ({ok_fields}/{len(pprop['person'])} person fields verified)")

    bad = json.loads(json.dumps(GOOD_ANSWER))
    bad["person"]["ecsa_number"] = {"value": "20161111", "evidence": "registration no. 20161111"}  # invented
    bad["person"]["highest_level"] = {"value": "MSc", "evidence": "MSc (Eng) Transportation"}     # not a SANRAL option
    bad["person"]["sacpcmp_number"] = {"value": "PrCPM 123", "evidence": "SACPCMP PrCPM 123"}     # not in CV
    prop, _ = x.validate_proposal(bad, text, ids)
    check(not prop["person"]["ecsa_number"]["verified"] and any("digits" in w for w in prop["person"]["ecsa_number"]["warnings"]),
          "invented registration number flagged (quote and digits not in CV)")
    hl = prop["person"]["highest_level"]
    check(hl["value"] is None and not hl["verified"] and "SANRAL's options" in hl["warnings"][0],
          "value outside SANRAL's dropdown list blanked, reviewer asked to choose")
    check(not prop["person"]["sacpcmp_number"]["verified"], "value quoted from text that isn't in the CV flagged")

    not_cv = dict(GOOD_ANSWER, document_kind="not_a_cv", document_kind_reason="It is a tender's returnable schedule.")
    expect_fail(lambda: x.validate_proposal(not_cv, text, []), "doesn't look like one person's cv", "model says not a CV")
    empty = dict(GOOD_ANSWER, person={k: {"value": None, "evidence": None} for k in GOOD_ANSWER["person"]}, projects=[])
    expect_fail(lambda: x.validate_proposal(empty, text, []), "none of the cv fields", "nothing found")
    made_up = json.loads(json.dumps(GOOD_ANSWER))
    for k, v in made_up["person"].items():
        if v["value"]:
            v["evidence"] = "zzqx " + k
    for pr in made_up["projects"]:
        pr["client_and_number"], pr["evidence"] = "Imaginary Client " + pr["position_held"], "nowhere"
    expect_fail(lambda: x.validate_proposal(made_up, text, []), "couldn't be traced back", "mostly unverifiable answer")
    two_ids = x.validate_proposal(json.loads(json.dumps(GOOD_ANSWER)), text, [SYNTH_ID, sa_id("850101500108")])
    check("id_number" not in two_ids[0]["person"] and any("2 different SA ID" in w for w in two_ids[1]),
          "two different ID numbers in one CV: none proposed, warning raised")


class StubResponse:
    def __init__(self, text=None, stop_reason="end_turn"):
        self.content = [types.SimpleNamespace(type="thinking", thinking="")] + (
            [types.SimpleNamespace(type="text", text=text)] if text is not None else [])
        self.stop_reason, self.model = stop_reason, "stub-model"
        self.usage = types.SimpleNamespace(input_tokens=5000, output_tokens=1800)


class StubClient:
    """Stands in for anthropic.Anthropic(): records the request, returns a canned response."""
    def __init__(self, response):
        self.calls = []
        outer = self
        self.beta = types.SimpleNamespace(messages=types.SimpleNamespace(
            create=lambda **kw: (outer.calls.append(kw), response)[1]))


def test_llm_call():
    print("4. call_llm request shape and stop-reason handling (stub client)")
    c = StubClient(StubResponse(json.dumps(GOOD_ANSWER)))
    data, meta = x.call_llm("cv text", client=c)
    kw = c.calls[0]
    check(kw["model"] == "claude-opus-5-5" and kw["fallbacks"] == "default"
          and kw["betas"] == ["server-side-fallback-2026-07-01"], "model, server-side fallback and beta header")
    check(kw["output_config"]["format"]["type"] == "json_schema" and kw["output_config"]["effort"] == "medium"
          and "thinking" not in kw and "temperature" not in kw, "structured output + effort, no disabled-thinking/sampling params")
    check("<cv_text>\ncv text\n</cv_text>" in kw["messages"][0]["content"], "CV text sent inside <cv_text> tags")
    check(data["document_kind"] == "cv" and meta["usage"]["output_tokens"] == 1800, "JSON parsed from the text block, usage kept")
    for resp, needle, label in [(StubResponse(None, "refusal"), "declined", "refusal"),
                                (StubResponse("{}", "max_tokens"), "more content", "max_tokens"),
                                (StubResponse("not json"), "valid structured", "bad JSON")]:
        expect_fail(lambda: x.call_llm("t", client=StubClient(resp)), needle, f"stop/answer: {label}")
    saved = os.environ.pop("ANTHROPIC_API_KEY", None)
    try:
        x.default_llm_client()
        check(False, "missing key should raise SetupError")
    except x.SetupError as e:
        check("ANTHROPIC_API_KEY is not set" in str(e), "missing ANTHROPIC_API_KEY: clear SetupError")
    finally:
        if saved:
            os.environ["ANTHROPIC_API_KEY"] = saved

    # Re-review 2026-10-23: a live upload exposed the raw provider error and
    # request_id straight to the customer (ExtractionFailed used to embed
    # e.message verbatim). A BadRequestError - the exact shape of the real
    # "too many union parameters" rejection - must now come out as a clean,
    # fixed customer message, with the technical detail only on `.detail`.
    import anthropic
    # Finding 65, 2026-10-23: this is httpx2, not httpx - anthropic's own
    # _base_client.py imports httpx2 throughout ("`httpx` objects don't work
    # with `httpx2`, and would otherwise fail deep inside a request") and
    # anthropic==1.15.0's real dependency is httpx2<3,>=2.0.0, confirmed from
    # a real CI run's own pip install log, not assumed. A bare `import httpx`
    # here was wrong in two ways at once: ModuleNotFoundError since nothing
    # in this dependency tree installs plain httpx, and even if it were
    # installed, passing an httpx.Response into anthropic.BadRequestError
    # would silently misbehave per that same docstring.
    import httpx2 as httpx
    raw = ("Schemas contains too many parameters with union types (42 parameters with type arrays or anyOf). "
           "This causes exponential compilation cost. Reduce the number of nullable or union-typed parameters "
           "(limit: 16 parameters with unions).")
    resp = httpx.Response(400, request=httpx.Request("POST", "https://api.anthropic.com/v1/messages"),
                          json={"type": "error", "error": {"type": "invalid_request_error", "message": raw},
                                "request_id": "req_test123"})

    def boom(**kw):
        raise anthropic.BadRequestError(raw, response=resp, body=resp.json())
    bad_client = types.SimpleNamespace(beta=types.SimpleNamespace(messages=types.SimpleNamespace(create=boom)))
    try:
        x.call_llm("cv text", client=bad_client)
        check(False, "BadRequestError should raise ExtractionFailed")
    except x.ExtractionFailed as e:
        check("Error code" not in str(e) and "request_id" not in str(e) and "42 parameters" not in str(e),
              "customer-facing message has no raw provider error or request id")
        check(e.detail and "42 parameters" in e.detail, "raw provider detail kept on .detail for system_alerts, not lost")


class MemoryDb:
    """In-memory stand-in for supabase_request + storage_download (only the URL
    shapes extract_cv_fields.py sends)."""
    def __init__(self, files):
        self.tables = {"cv_uploads": [], "cv_extraction_results": []}
        self.files = files

    def _parse(self, path):
        import urllib.parse
        table, _, q = path.partition("?")
        filters = {}
        for k, v in urllib.parse.parse_qsl(q):
            if v.startswith("eq."):
                filters[k] = v[3:]
        return table, filters

    def request(self, url, key, method, path, body=None):
        table, f = self._parse(path)
        rows = self.tables.setdefault(table, [])
        match = [r for r in rows if all(str(r.get(k)) == v for k, v in f.items())]
        if method == "GET":
            return [dict(r) for r in match]
        if method == "PATCH":
            for r in match:
                r.update(body)
            return [dict(r) for r in match]
        if method == "POST":
            existing = [r for r in rows if r.get("cv_upload_id") == body.get("cv_upload_id")]
            if existing:
                existing[0].update(body)
            else:
                rows.append(dict(body))
            return [body]
        raise AssertionError(method)

    def download(self, url, key, bucket, object_path, dest):
        shutil.copy(self.files[object_path], dest)


def test_queue(docx_path, tmp):
    print("5. queue plumbing (in-memory Supabase stand-in, stub LLM client)")
    from PIL import Image
    scan = os.path.join(tmp, "scan2.pdf")
    Image.new("RGB", (600, 800), "white").save(scan, "PDF")
    db = MemoryDb({"c1/people/p1/uploads/a-naledi_cv.docx": docx_path, "c1/people/p1/uploads/b-scan.pdf": scan})
    x.supabase_request, x.storage_download = db.request, db.download
    base = {"company_id": "c1", "person_id": "p1", "mode": "ai_extract", "status": "extraction_requested"}
    db.tables["cv_uploads"] = [dict(base, id="u1", storage_path="c1/people/p1/uploads/a-naledi_cv.docx", created_at="1"),
                               dict(base, id="u2", storage_path="c1/people/p1/uploads/b-scan.pdf", created_at="2"),
                               dict(base, id="u3", mode="as_is", status="uploaded", storage_path="x", created_at="3")]

    def no_key():
        raise x.SetupError("ANTHROPIC_API_KEY is not set")
    try:
        x.poll_once("u", "k", client_factory=no_key)
        check(False, "missing key should stop the pass")
    except x.SetupError:
        st = {r["id"]: r["status"] for r in db.tables["cv_uploads"]}
        check(st["u1"] == "extraction_requested" and not db.tables["cv_extraction_results"],
              "missing key: pass stops, upload put back in the queue (not failed), nothing written")

    stub = StubClient(StubResponse(json.dumps(GOOD_ANSWER)))
    handled = x.poll_once("u", "k", client_factory=lambda: stub)
    st = {r["id"]: r for r in db.tables["cv_uploads"]}
    res = db.tables["cv_extraction_results"]
    check(handled == 2 and st["u1"]["status"] == "extracted" and len(res) == 1
          and res[0]["review_status"] == "pending_review" and res[0]["cv_upload_id"] == "u1",
          "docx: result stored for review, upload marked extracted")
    check(SYNTH_ID not in stub.calls[0]["messages"][0]["content"], "the ID number never reached the API request")
    check(st["u2"]["status"] == "extraction_failed" and "no text layer" in st["u2"]["extraction_error"]
          and len(stub.calls) == 1, "scan: failed with reason, and no API call was made for it")
    check(st["u3"]["status"] == "uploaded", "as-is upload untouched")

    db.tables["cv_uploads"].append(dict(base, id="u4", status="extracting", storage_path="x",
                                        extraction_claimed_at="2020-01-01T00:00:00+00:00", created_at="4"))
    x.poll_once("u", "k", quiet_when_empty=True, client_factory=lambda: stub)
    u4 = [r for r in db.tables["cv_uploads"] if r["id"] == "u4"][0]
    check(u4["status"] == "extraction_failed" and "stopped mid-job" in u4["extraction_error"], "stale 'extracting' reclaimed as failed")

    # A write that REPORTS an outcome (fail_upload) must never itself crash
    # the pass if IT fails (a transient Supabase blip right as a different
    # error is being reported) - found missing, UX audit 2026-10-23: a prior
    # contributor to a row staying 'extracting' past the stale-reclaim window.
    db.tables["cv_uploads"].append(dict(base, id="u5", storage_path="c1/people/p1/uploads/missing.docx", created_at="5"))
    real_request = db.request

    def flaky_request(url, key, method, path, body=None):
        if method == "PATCH" and path == "cv_uploads?id=eq.u5":  # fail_upload()'s own write, not claim_upload()'s
            raise RuntimeError("simulated transient Supabase failure while marking u5 failed")
        return real_request(url, key, method, path, body)
    x.supabase_request = flaky_request
    try:
        handled5 = x.poll_once("u", "k", quiet_when_empty=True, client_factory=lambda: stub)
        check(handled5 >= 1, "poll_once() itself doesn't crash when fail_upload()'s own write fails")
    finally:
        x.supabase_request = real_request


def test_ocr(docx_path, tmp):
    """EXPERIMENTAL OCR path (engine/ocr/scan_ocr.py). Needs tesseract; skipped
    (stated, not silently) where it isn't installed."""
    print("6. OCR of a scanned CV (experimental)")
    import scan_ocr
    ok, why = scan_ocr.available()
    redacted, ids = x.find_and_redact_sa_ids("ID 9002150123081 and 1234567890123", redact_all_13_digit=True)
    check("1234567890123" not in redacted and "9002150123081" not in redacted,
          "OCR'd text: every 13-digit run is removed before sending, even one whose checksum fails (a misread ID)")
    if not ok or shutil.which("soffice") is None:
        print(f"  (skipped: {why if not ok else 'soffice not installed'})")
        return
    from PIL import Image, ImageFilter
    pdf = soffice_convert(docx_path, "pdf", tmp)
    subprocess.run(["pdftoppm", "-r", "200", "-gray", "-png", pdf, os.path.join(tmp, "scanpage")], check=True)
    pages = sorted(f for f in os.listdir(tmp) if f.startswith("scanpage"))
    imgs = [Image.open(os.path.join(tmp, f)).convert("L") for f in pages]
    scan_pdf = os.path.join(tmp, "scanned_cv_ocr.pdf")
    imgs[0].save(scan_pdf, "PDF", resolution=200, save_all=True, append_images=imgs[1:])
    r = x.read_cv_text(scan_pdf, tmp)
    check(r["method"].startswith("OCR (") and r["ocr"] and r["ocr"]["pages_ocr"], f"scanned PDF read by OCR: {r['method']}")
    check("20169999" in r["text"] and "University of Cape Town" in r["text"], "OCR text has the ECSA number and institution")
    stub = StubClient(StubResponse(json.dumps(GOOD_ANSWER)))
    proposal, warnings, info = x.extract_from_file(scan_pdf, tmp, client=stub)
    check(warnings and warnings[0].startswith("OCR:"), "review step is told the text came from OCR")
    check(all(not v["verified"] and x.OCR_FIELD_WARNING in v["warnings"] for v in proposal["person"].values())
          and all(not p["verified"] for p in proposal["projects"]),
          "every value read from the scan starts unticked, with the OCR warning")
    check(SYNTH_ID not in stub.calls[0]["messages"][0]["content"], "OCR'd ID number never reached the API request")
    # an illegible scan: the same page blurred until OCR can't read it -> refused, never guessed
    blurred = os.path.join(tmp, "blurred_cv.pdf")
    imgs[0].resize((imgs[0].width // 4, imgs[0].height // 4)).filter(ImageFilter.GaussianBlur(2)).save(blurred, "PDF", resolution=50)
    expect_fail(lambda: x.read_cv_text(blurred, tmp), "OCR couldn't read it clearly", "illegible scan")
    # a photo upload of the same page is OCR'd too
    png = os.path.join(tmp, "cv_photo.png")
    imgs[0].save(png)
    r = x.read_cv_text(png, tmp)
    check(r["method"] == f"OCR ({r['ocr']['engine']}) - image upload" and "20169999" in r["text"], "image upload read by OCR")


def test_sa_id_redaction_line_break():
    """Finding (2026-10-21): the separator between an ID's 6/4/3 digit groups
    only matched a literal space or hyphen - but test_text_extraction()'s own
    comment above already confirms real PDF CVs get mid-token line-wraps
    from a narrow table column (the same artifact this test exercises for an
    ID specifically). Neither half of a line-split ID is 13 digits on its
    own, so the whole number passed through completely unredacted and would
    have reached the Anthropic API verbatim - the exact failure this
    module's "found locally and redacted BEFORE any text leaves this
    machine" promise exists to prevent."""
    print("7. SA ID redaction survives a mid-number line break (narrow table-cell PDF wrap)")
    line_split = f"ID Number: {SYNTH_ID[:6]}\n{SYNTH_ID[6:10]}\n{SYNTH_ID[10:]}\nName: Jane Doe"
    redacted, ids = x.find_and_redact_sa_ids(line_split)
    check(ids == [SYNTH_ID] and SYNTH_ID not in redacted.replace("\n", ""),
          "an ID split across line breaks is still found and redacted, not leaked")
    # Mirror case: two genuinely UNRELATED digit runs that happen to sit on
    # either side of a line break must NOT be fused into a fake ID - proves
    # the fix isn't over-broad.
    unrelated = "Phone: 011555\n1234 ext 99"
    redacted2, ids2 = x.find_and_redact_sa_ids(unrelated)
    check(redacted2 == unrelated and ids2 == [], "unrelated digits split by a line break are never fused into a false ID")


def main():
    tmp = tempfile.mkdtemp(prefix="cv_extract_test_")
    try:
        docx_path = test_text_extraction(tmp)
        test_failure_modes(tmp)
        test_validation(docx_path, tmp)
        test_llm_call()
        test_queue(docx_path, tmp)
        test_ocr(docx_path, tmp)
        test_sa_id_redaction_line_break()
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print(f"\n{'ALL PASSED' if not FAILS else f'{len(FAILS)} FAILED'}")
    return 1 if FAILS else 0


if __name__ == "__main__":
    sys.exit(main())
