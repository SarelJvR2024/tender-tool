"""
Generates a DRAFT "Authority for Signatory" letter on company letterhead, for use with
Form A2 (Certificate of Authority for Signatory).

This does NOT fill Form A2 itself - the template's own notes say the example page,
completed, "shall not be accepted as authorisation." What SANRAL actually wants is a
real letter on the tenderer's letterhead, signed by whoever has authority to appoint a
signatory (for Klipcorp: the director, per the company's own past practice). This script
reproduces Klipcorp's own real historical letter format (found in Drive:
"Letter of Authority to Sign_Klipcorp .docx", dated 2023-07-20), substituting this
tender's own details - it does not invent a new format or wording.

The output is explicitly a DRAFT: the director's signature block is left blank for a real
signature, never auto-signed. Someone with actual authority (Krisun Naidoo, per the
precedent letters) must review and sign a printed copy.

Usage:
    python3 fill_authority_letter.py <profile.json> <logo.png> <output.docx> \\
        --signatory-name "Henning Jansen van Rensburg" \\
        --signatory-capacity "Senior Manager" \\
        --rfp-number "SANRAL NRA 2026/0073 / Contract R.334-030-2025/1F" \\
        --rfp-description "Consulting engineering services for the routine ad-hoc \\
maintenance of R334 Section 3 from Great Fish River (km 0.0) to Mpeko (km 13.2)"
"""
import sys
import json
import argparse
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def build_letter(profile_path, logo_path, out_path, signatory_name, signatory_capacity,
                  rfp_number, rfp_description):
    profile = json.load(open(profile_path))
    comp = profile["company"]
    director = comp["director"]

    d = docx.Document()

    # Letterhead: company logo at the top of the page.
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(logo_path, width=Inches(1.8))

    d.add_paragraph()

    p = d.add_paragraph("TO WHOM IT MAY CONCERN")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(13)

    p = d.add_paragraph("AUTHORITY FOR SIGNATORY")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(13)

    d.add_paragraph()

    body = (
        f"It is hereby certified that Mr {signatory_name} in his capacity as "
        f"{signatory_capacity} is duly authorised, on behalf of the Board of Directors, "
        f"to sign all documents in connection with the tender on behalf of "
        f"{comp['legal_name']},\n\n"
        f"RFP NUMBER: {rfp_number}\n"
        f"{rfp_description}"
    )
    d.add_paragraph(body)

    d.add_paragraph()
    d.add_paragraph()
    d.add_paragraph("Your's Sincerely")
    d.add_paragraph()
    d.add_paragraph()
    d.add_paragraph("_____________________________")
    d.add_paragraph(director["name"])
    d.add_paragraph("DIRECTOR")

    d.save(out_path)
    print(f"Wrote authority letter to {out_path}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("profile_path")
    ap.add_argument("logo_path")
    ap.add_argument("out_path")
    ap.add_argument("--signatory-name", required=True)
    ap.add_argument("--signatory-capacity", required=True)
    ap.add_argument("--rfp-number", required=True)
    ap.add_argument("--rfp-description", required=True)
    args = ap.parse_args()
    build_letter(args.profile_path, args.logo_path, args.out_path,
                 args.signatory_name, args.signatory_capacity,
                 args.rfp_number, args.rfp_description)


if __name__ == "__main__":
    main()
