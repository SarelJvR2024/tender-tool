"""
Checks engine/company_documents.py (Hardcoded-Assumptions-Audit Q7): every
company's letter of authority to sign is produced the way THAT company chose -
its own upload as-is, or its own guideline - and a company that chose nothing
gets "not set up", never Klipcorp's historical wording.

Each case writes a real .docx and reads it back with python-docx.

    python3 engine/test_company_documents.py [out_dir]
    COMPANY_DOCS_PSQL="psql -h /tmp/sock -p 54411 -U postgres -d tpltest" python3 engine/test_company_documents.py
        (also part 7: the real table (0037) on a scratch Postgres with every
        migration replayed - RLS, and the pack bridge's
        process_pack_jobs.load_document_templates() feeding produce())
"""
import json
import os
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import docx  # noqa: E402
import company_documents as cd  # noqa: E402
import fill_authority_letter  # noqa: E402
from provenance import ProvenanceLog  # noqa: E402

FAILS = []

# The wording fill_authority_letter.py used to write for every company
# (Klipcorp's 2023-07-20 letter). None of it may appear unless a company's own
# upload/guideline contains it.
KLIPCORP_PHRASES = ["It is hereby certified that", "duly authorised, on behalf of the Board of Directors",
                    "Your's Sincerely", "TO WHOM IT MAY CONCERN", "AUTHORITY FOR SIGNATORY"]


def check(cond, msg):
    print(("  PASS " if cond else "  FAIL ") + msg)
    if not cond:
        FAILS.append(msg)


def text_of(path):
    return "\n".join(p.text for p in docx.Document(path).paragraphs)


def profile(name, director=None):
    return {"company": {"legal_name": name, **({"director": {"name": director}} if director else {})},
            "key_people": [{"name": "Signatory Person", "signatory_capacity": "Managing Member"}]}


TENDER = {"contract_no": "R.396-010-2025/1F", "nra_reference": "NRA 2026/0071", "closing_date": "2026-11-05",
          "description": "Consulting engineering services for the routine ad-hoc maintenance of R396 Section 1"}
RFP = "SANRAL NRA 2026/0071 / Contract R.396-010-2025/1F"


def run(out):
    os.makedirs(out, exist_ok=True)

    print("1. Uploaded - the company's own letter, used as-is")
    own = os.path.join(out, "own_letter_src.docx")
    d = docx.Document()
    for line in ["BOARD RESOLUTION 14/2026 OF UPLOADED OWN-LETTER ENGINEERS (PTY) LTD",
                 "Resolved that the person named below may sign tender documents for the company.",
                 "Chairperson: ____________"]:
        d.add_paragraph(line)
    d.save(own)
    p = profile("Uploaded Own-Letter Engineers (Pty) Ltd")
    p["document_templates"] = {"authority_to_sign": {"mode": "uploaded", "path": own,
                                                     "original_filename": "Resolution 14-2026.docx"}}
    log = ProvenanceLog("test")
    r = fill_authority_letter.produce_letter(p, os.path.join(out, "uploaded"), p["key_people"][0], TENDER, RFP, log)
    check(r["status"] == "uploaded" and r["method"] == "uploaded", f"status uploaded ({r['status']})")
    check(open(r["path"], "rb").read() == open(own, "rb").read(), "output file is byte-identical to the upload")
    check(text_of(r["path"]) == text_of(own), "read back: the company's own text, unchanged")
    check(not any(k in text_of(r["path"]) for k in KLIPCORP_PHRASES), "no Klipcorp wording")

    print("2. Guideline - the company's own wording with {{fields}}")
    guideline = ("# SIGNING AUTHORITY - {{company_legal_name}}\n\n"
                 "The directors of {{company_legal_name}} confirm by written resolution that {{signatory_name}}, "
                 "our {{signatory_capacity}}, may sign and submit every document for tender {{tender_reference}} "
                 "({{tender_description}}), closing {{closing_date}}.\n\n"
                 "Signed for the directors on {{date}}:\n\n__________________\n{{director_name}}, Director")
    p = profile("Guideline Wording Consultants CC", director="Thandi Director")
    p["document_templates"] = {"authority_to_sign": {"mode": "ai_guideline", "guideline": guideline}}
    log = ProvenanceLog("test")
    r = fill_authority_letter.produce_letter(p, os.path.join(out, "guideline"), p["key_people"][0], TENDER, RFP, log)
    got = text_of(r["path"])
    facts = cd.authority_facts(p, p["key_people"][0], TENDER, RFP)
    expected, _, _ = cd.fill_fields(guideline, facts, "authority_to_sign")
    expected = expected.replace("# ", "", 1)
    print("    --- read back ---\n    " + got.replace("\n", "\n    "))
    check(r["status"] == "generated" and r["method"] == "fields", f"generated from fields ({r['method']})")
    check(got.split() == expected.split(), "read back == the guideline with only its {{fields}} filled")
    check(cd.guard_only_company_wording(got, guideline, facts) == [],
          "every word is the company's guideline or a tender/company fact")
    check(not any(k.lower() in got.lower() for k in KLIPCORP_PHRASES), "no Klipcorp wording")
    check("Thandi Director" in got and RFP in got and "Managing Member" in got, "this tender's and company's facts placed")
    check(all(e["status"] == "confirmed" for e in log.entries), "provenance: every field confirmed")

    print("3. Guideline asking for a field that isn't on file")
    p = profile("No Director On File (Pty) Ltd")
    p["document_templates"] = {"authority_to_sign": {"mode": "ai_guideline",
                                                     "guideline": "{{signatory_name}} signs. Approved: {{director_name}}. {{bogus_field}}"}}
    log = ProvenanceLog("test")
    r = fill_authority_letter.produce_letter(p, os.path.join(out, "missing"), p["key_people"][0], TENDER, RFP, log)
    got = text_of(r["path"])
    check("[NOT ON FILE: Director's name]" in got, "missing value printed visibly as [NOT ON FILE: ...]")
    check("[UNKNOWN FIELD: bogus_field]" in got, "unknown field printed visibly")
    check(sum(e["status"] == "missing" for e in log.entries) == 2, "both logged missing (-> ACTIONS_TO_SIGN.md)")

    print("4. Nothing set up - a new company")
    p = profile("Brand New Startup (Pty) Ltd")
    log = ProvenanceLog("test")
    r = fill_authority_letter.produce_letter(p, os.path.join(out, "not_set_up"), p["key_people"][0], TENDER, RFP, log)
    check(r["status"] == "not_set_up" and r["path"] is None, "status not_set_up, no document produced")
    check("not set up" in r["reason"] and "Brand New Startup" in r["reason"], f"reason: {r['reason']}")
    check(not os.listdir(os.path.join(out, "not_set_up")), "nothing written (no fallback wording)")
    check(log.entries[0]["status"] == "missing", "logged missing (-> ACTIONS_TO_SIGN.md)")
    p["document_templates"] = {"authority_to_sign": {"mode": "ai_guideline", "guideline": "   "}}
    r = fill_authority_letter.produce_letter(p, os.path.join(out, "blank"), p["key_people"][0], TENDER, RFP, ProvenanceLog("t"))
    check(r["status"] == "not_set_up", "a blank guideline is 'not set up' too")

    print("5. AI placement into a guideline with no {{fields}} - output is checked, not trusted")
    plain = ("Our managing member Signatory Person holds signing authority for this tender. "
             "Resolution of the members, Lindiwe Holdings.")

    class Stub:
        def __init__(self, text):
            self.text = text
            self.messages = self

        def create(self, **kw):
            self.kw = kw
            return type("R", (), {"content": [type("B", (), {"text": self.text})()]})()

    good = Stub("Our managing member Signatory Person holds signing authority for this tender "
                f"{RFP}. Resolution of the members, Lindiwe Holdings.")
    p = profile("Lindiwe Holdings")
    p["document_templates"] = {"authority_to_sign": {"mode": "ai_guideline", "guideline": plain}}
    log = ProvenanceLog("test")
    r = fill_authority_letter.produce_letter(p, os.path.join(out, "ai_good"), p["key_people"][0], TENDER, RFP, log, llm=good)
    check(r["method"] == "ai_placed" and RFP in text_of(r["path"]), "well-behaved AI: facts placed, accepted")
    check("Output only the document text" in good.kw["system"], "AI told to use the company's words exactly")
    bad = Stub("We are committed to excellence and world-class service delivery. " + good.text)
    log = ProvenanceLog("test")
    r = fill_authority_letter.produce_letter(p, os.path.join(out, "ai_bad"), p["key_people"][0], TENDER, RFP, log, llm=bad)
    got = text_of(r["path"])
    check(r["method"] == "verbatim" and got == plain, "AI that adds boilerplate is rejected; guideline used verbatim")
    check("committed" not in got and "excellence" not in got, "the invented wording never reaches the document")
    rej = [e for e in log.entries if e["field_key"] == "authority_to_sign.ai_placement"]
    check(rej and "committed" in rej[0]["reason"], f"rejection logged with the invented words: {rej[0]['reason'][:110]}...")
    log = ProvenanceLog("test")
    r = fill_authority_letter.produce_letter(p, os.path.join(out, "no_llm"), p["key_people"][0], TENDER, RFP, log)
    check(r["method"] == "verbatim" and log.entries[0]["status"] == "assumed_default",
          "no AI configured: guideline verbatim, flagged 'no fields marked - check it'")

    print("6. Same tender, two companies -> two different letters")
    outs = {}
    for name, sub in (("uploaded", "uploaded"), ("guideline", "guideline")):
        d = os.path.join(out, sub)
        outs[name] = text_of(os.path.join(d, sorted(os.listdir(d))[0]))
    check(outs["uploaded"] != outs["guideline"], "uploaded and guideline companies' letters differ")

    print("6b. authority_facts()'s {{date}} field uses SAST, not a naive UTC date")
    # Finding (2026-10-27): authority_facts() computed "today" via naive
    # datetime.date.today() - a bridge worker container has no TZ set and
    # runs in UTC by default, so this is one calendar day behind the real
    # SAST date for the whole 00:00-01:59 SAST window every day, printing
    # yesterday's date on a real authority-to-sign letter. Simulates the
    # system clock at UTC 2026-10-20 23:30 (= SAST 2026-10-21 01:30, inside
    # the bad window) by monkeypatching this module's own datetime.date AND
    # datetime.datetime (the fixed code reads the aware one; the original
    # buggy code called datetime.date.today() directly, so both must be
    # patched for this test to actually exercise either version), restored
    # immediately after - never leaves the real clock patched. Deliberately
    # NOT today's real date, so a coincidental match can't hide a reverted
    # fix (confirmed: patching only datetime.datetime here first let the
    # revert pass by accident, since the then-still-live datetime.date.today()
    # fell through to the real sandbox clock, which happened to already read
    # the target date).
    import datetime as real_datetime

    class FixedDateTime(real_datetime.datetime):
        @classmethod
        def now(cls, tz=None):
            base = real_datetime.datetime(2026, 10, 20, 23, 30, tzinfo=real_datetime.timezone.utc)
            return base.astimezone(tz) if tz else base

    class FixedDate(real_datetime.date):
        @classmethod
        def today(cls):
            return cls(2026, 10, 20)  # the naive UTC date at that same simulated moment

    real_dt, real_date = cd.datetime.datetime, cd.datetime.date
    cd.datetime.datetime, cd.datetime.date = FixedDateTime, FixedDate
    try:
        facts = cd.authority_facts(profile("SAST Date Test (Pty) Ltd"), {"name": "Jane"}, {}, "REF1")
    finally:
        cd.datetime.datetime, cd.datetime.date = real_dt, real_date
    check(facts["date"] == "21 October 2026",
          f"authority_facts()'s date reads the real SAST date (21 October 2026), not the naive "
          f"UTC one (20 October 2026) - got {facts['date']!r}")


def run_db(psql, out):
    sys.path.insert(0, os.path.join(HERE, "intake"))
    sys.path.insert(0, os.path.join(HERE, "pack"))
    from test_manual_intake import Client
    from bridge_selftest import lit
    import process_pack_jobs
    print("7. The real table on a scratch Postgres (0001..latest replayed)")
    storage = os.path.join(out, "storage")
    c = Client(psql, storage)
    c.sql("delete from company_document_templates; delete from companies where legal_name like 'DocTpl %'")
    c.sql("grant usage on schema public to authenticated; grant usage on schema storage to authenticated; "
          "grant all on all tables in schema public to authenticated; grant all on storage.objects to authenticated; "
          "grant execute on all functions in schema public to authenticated")
    cos = {n: c.insert("companies", {"legal_name": n})["id"] for n in
           ("DocTpl Uploader (Pty) Ltd", "DocTpl Guideline CC", "DocTpl Nothing Yet (Pty) Ltd")}
    users = {}
    for n, cid in cos.items():
        users[n] = c.sql("insert into auth.users default values returning id").splitlines()[0]
        c.insert("company_members", {"company_id": cid, "user_id": users[n], "role": "owner"})
    up, gl, none = cos["DocTpl Uploader (Pty) Ltd"], cos["DocTpl Guideline CC"], cos["DocTpl Nothing Yet (Pty) Ltd"]
    own = os.path.join(out, "own_letter_src.docx")
    path = f"{up}/templates/authority_to_sign/abc-own.docx"
    ok, msg = c.as_user(users["DocTpl Uploader (Pty) Ltd"],
                        f"insert into storage.objects (bucket_id, name) values ('company-documents', {lit(path)}) returning name")
    check(ok, f"member uploads into its own company-documents prefix ({msg[:60]})")
    c.upload("company-documents", path, own, None)
    ok, msg = c.as_user(users["DocTpl Uploader (Pty) Ltd"],
                        "insert into company_document_templates (company_id, document_type, mode, storage_path, original_filename) "
                        f"values ({lit(up)}, 'authority_to_sign', 'uploaded', {lit(path)}, 'Resolution 14-2026.docx') returning id")
    check(ok, "member saves its own 'uploaded' choice under RLS")
    ok, msg = c.as_user(users["DocTpl Guideline CC"],
                        "insert into company_document_templates (company_id, document_type, mode, guideline) values "
                        f"({lit(gl)}, 'authority_to_sign', 'ai_guideline', "
                        f"{lit('Our members authorise {{signatory_name}} to sign tender {{tender_reference}} for {{company_legal_name}}.')}) returning id")
    check(ok, "member saves its own guideline under RLS")
    ok, msg = c.as_user(users["DocTpl Guideline CC"],
                        f"insert into company_document_templates (company_id, document_type, mode, guideline) values "
                        f"({lit(none)}, 'authority_to_sign', 'ai_guideline', 'hijack') returning id")
    check(not ok, "a member can't set another company's letter (RLS refuses)")
    ok, n = c.as_user(users["DocTpl Guideline CC"], "select count(*) from company_document_templates")
    check(ok and n == "1", f"a member sees only its own company's row ({n})")
    ok, msg = c.as_user(users["DocTpl Guideline CC"],
                        f"insert into storage.objects (bucket_id, name) values ('company-documents', {lit(up + '/templates/x/evil.docx')}) returning name")
    check(not ok, "a member can't upload into another company's template prefix")

    texts = {}
    for name, cid in cos.items():
        tpls = process_pack_jobs.load_document_templates(c, cid, os.path.join(out, "db", cid))
        prof = profile(name)
        prof["document_templates"] = tpls
        r = fill_authority_letter.produce_letter(prof, os.path.join(out, "db_out", cid), prof["key_people"][0],
                                                 TENDER, RFP, ProvenanceLog("t"))
        texts[name] = (r["status"], text_of(r["path"]) if r["path"] else r["reason"])
        print(f"    {name}: {r['status']} -> {texts[name][1][:110]}")
    check(texts["DocTpl Uploader (Pty) Ltd"] == ("uploaded", text_of(own)), "uploader: its own file, unchanged")
    check(texts["DocTpl Guideline CC"] == ("generated", f"Our members authorise Signatory Person to sign tender {RFP} for DocTpl Guideline CC."),
          "guideline company: its own sentence with the fields filled, nothing else")
    check(texts["DocTpl Nothing Yet (Pty) Ltd"][0] == "not_set_up", "company with no row: not set up")


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else tempfile.mkdtemp(prefix="company_documents_")
    run(out)
    if os.environ.get("COMPANY_DOCS_PSQL"):
        run_db(os.environ["COMPANY_DOCS_PSQL"], out)
    else:
        print("(part 7 skipped - set COMPANY_DOCS_PSQL to a scratch Postgres with the migrations replayed)")
    print(f"\nOutput: {out}")
    print(f"{len(FAILS)} failure(s)")
    sys.exit(1 if FAILS else 0)
