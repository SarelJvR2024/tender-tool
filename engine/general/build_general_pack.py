"""
Builds a complete non-SANRAL submission pack from a tender's requirement
manifest (engine/general/requirements.py), the company's data, the answers
captured on the Tender Workspace (tender_answers) and the company's files.

One ordered PDF (engine/pack/pdf_merge.py's merge_with_contents - the same
merger, contents page, page stamps and read-back check the SANRAL pack uses):

  Read first       Completion guide: every SBD/MBD field with the value to
                   write, for the forms filled here AND the ones that must be
                   completed by hand (flat PDF editions, forms inside a
                   bundled tender document), plus everything still to sign.
  Returnable forms The forms filled by fill_sbd_docx.fill_sbd_document() from
                   each editable .docx edition (a legacy .doc is converted by
                   LibreOffice first); the blank published edition of each
                   form that has to be done by hand.
  Buyer's own forms The buyer-specific forms (ESKOM Form A / declarations,
                   Sentech consent, municipal Form of Quotation...) as
                   published, to complete by hand - Tender Tool has no
                   template for them and says so.
  Company documents Vault copies + documents uploaded for this tender, in the
                   order the tender names them.
  Experience        The Experience Library items this tender's own evidence
                   rule accepts.
  Key personnel     CVs of the people tagged for the roles the tender names.

Anything that should be in the pack but isn't is listed on the contents page
with the reason (pdf_merge's "Not included" table) - never silently dropped.

This module is pure: it takes local file paths and dicts and never talks to
Supabase, so engine/general/general_harness.py can run it on real tender
documents with no live project. engine/general/process_general_jobs.py is the
Supabase adapter around it.
"""
import datetime
import json
import os
import re
import shutil
import subprocess
import sys

import docx
from docx.shared import Pt, Inches

HERE = os.path.dirname(os.path.abspath(__file__))
ENGINE_DIR = os.path.abspath(os.path.join(HERE, ".."))
for p in (HERE, ENGINE_DIR, os.path.join(ENGINE_DIR, "pack"), os.path.join(ENGINE_DIR, "cv")):
    if p not in sys.path:
        sys.path.insert(0, p)

import fill_sbd_docx  # noqa: E402
import fill_flat_pdf  # noqa: E402
import requirements as req  # noqa: E402
from provenance import ProvenanceLog  # noqa: E402
from pdf_merge import merge_with_contents, verify_pack, expiry_warning  # noqa: E402

FORMS_ENV = "Returnable forms"
BUYER_ENV = "Buyer's own forms (complete by hand)"
DOCS_ENV = "Company documents"
EXP_ENV = "Experience evidence"
CV_ENV = "Key personnel CVs"
GUIDE_ENV = "Read first"

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".gif"}
OFFICE_EXT = {".doc", ".docx", ".odt", ".rtf"}

FILLABLE_CODES = {"1", "4", "6.1", "8", "9"}

# Vault file per returnable (companies.*_path, 0014) - the rest are uploads.
VAULT_FILE_FOR = {
    "tax_compliance": "sars_tcs_pin_document_path",
    "bbbee_proof": "bee_proof_document_path",
    "coid": "coid_letter_document_path",
    "bank_letter": "bank_confirmation_letter_path",
    "pi_insurance": "pi_insurance_document_path",
}
VAULT_EXPIRY_FOR = {
    "tax_compliance": "sars_tcs_pin_expiry",
    "bbbee_proof": "bee_proof_expiry",
    "coid": "coid_letter_of_good_standing_expiry",
    "bank_letter": "bank_confirmation_letter_valid_until",
    "pi_insurance": "pi_insurance_expiry",
}

# Real incident, 2026-10-28 (Reviewer Civils & Consulting): a file named
# "...BBBEE_Affidavit.docx" was uploaded into the COID vault slot, and a real
# pack-generation run merged it into the pack labelled as the COID Letter of
# Good Standing. Same filename-only heuristic as lib/vaultDocuments.ts's own
# vaultDocumentMismatchWarning() on the app side (that file's own comment
# explains why filename, not real text extraction: this engine is the only
# place that reads real PDF/DOCX content, and build_general_pack() only ever
# sees a document AFTER it's already in the pack - too late to ask the vault
# page's own upload-time check to run again). Keyed by the companies.*_path
# column (VAULT_FILE_FOR's own values), not the returnable key, since that's
# the one name both VAULT_FILE_FOR here and VAULT_DOCUMENTS there agree on.
VAULT_MISMATCH_KEYWORDS = {
    "coid_letter_document_path": {
        "reject": (["bbbee", "b-bbee", "bee", "affidavit", "broad-based", "broadbased"], "a B-BBEE Affidavit"),
        "expect": ["coid", "compensation", "good standing", "letter of good standing"],
    },
    "bee_proof_document_path": {
        "expect": ["bee", "bbbee", "b-bbee", "affidavit", "sanas", "broad-based", "broadbased"],
    },
    "sars_tcs_pin_document_path": {
        "expect": ["tax", "sars", "pin", "tcs", "compliance"],
    },
}


def vault_document_mismatch_warning(column, local_path):
    """Non-blocking content-sanity warning for one vault document already
    downloaded to `local_path` - filename only, same reasoning as this
    module's own VAULT_MISMATCH_KEYWORDS comment. Returns None for a column
    with no rule, or for a filename that passes its rule."""
    rule = VAULT_MISMATCH_KEYWORDS.get(column)
    if not rule or not local_path:
        return None
    name = os.path.basename(local_path).lower()
    reject = rule.get("reject")
    if reject:
        words, looks_like = reject
        if any(w in name for w in words):
            return f"This document appears to be {looks_like}, not {VAULT_LABEL_FOR.get(column, 'this compliance document')}."
    expect = rule.get("expect")
    if expect and not any(w in name for w in expect):
        return f"This file's name doesn't look like {VAULT_LABEL_FOR.get(column, 'this compliance document')} - check the right file was uploaded."
    return None


# Labels for vault_document_mismatch_warning()'s own messages - same wording
# lib/vaultDocuments.ts's VAULT_DOCUMENTS gives each column on the app side.
VAULT_LABEL_FOR = {
    "coid_letter_document_path": "COID Letter of Good Standing",
    "bee_proof_document_path": "a B-BBEE certificate or sworn affidavit",
    "sars_tcs_pin_document_path": "a SARS Tax Compliance Status PIN letter",
}


def to_pdf(path, work_dir):
    """PDF / image / Word file -> (pdf_path, None) or (None, reason).
    Same rules as the SANRAL assembler's _to_pdf (engine/pack/assemble_pack.py);
    kept here so this engine doesn't import the SANRAL assembler."""
    if not path or not os.path.exists(path):
        return None, "file not found"
    ext = os.path.splitext(path)[1].lower()
    os.makedirs(work_dir, exist_ok=True)
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
            return (path, None) if len(PdfReader(path).pages) else (None, "PDF has no pages")
        except Exception as e:
            return None, f"PDF could not be read ({type(e).__name__})"
    if ext in IMAGE_EXT:
        from PIL import Image
        out = os.path.join(work_dir, os.path.splitext(os.path.basename(path))[0] + ".pdf")
        try:
            img = Image.open(path)
            img.load()
            img.convert("RGB").save(out, "PDF", resolution=150)
            return out, None
        except Exception as e:
            return None, f"image could not be converted ({type(e).__name__})"
    if ext in OFFICE_EXT:
        local = os.path.join(work_dir, os.path.basename(path))
        if os.path.abspath(local) != os.path.abspath(path):
            shutil.copy(path, local)
        pdf = req.convert_with_soffice(local, "pdf", work_dir)
        return (pdf, None) if pdf else (None, "Word file could not be converted to PDF")
    return None, f"unsupported file type '{ext}'"


# ---------------------------------------------------------------------------
# Completion guide
# ---------------------------------------------------------------------------

def _yn(v):
    return {"yes": "Yes", "no": "No"}.get((v or "").lower(), v)


def guide_rows(code, ctx):
    """(field, value-or-None, note) for one form - the same facts
    fill_sbd_document() writes, for a person completing a flat PDF by hand."""
    c, a, s, t = ctx["company"], ctx["answers"], ctx.get("signatory") or {}, ctx.get("tender") or {}
    proof = {"certificate": "B-BBEE verification certificate", "qse_affidavit": "QSE sworn affidavit",
             "eme_affidavit": "EME sworn affidavit"}.get(c.get("bee_proof_type"))
    if code == "1":
        return [
            ("Name of bidder", c.get("legal_name"), None),
            ("Postal address", c.get("postal_address"), None),
            ("Street address", c.get("registered_address"), None),
            ("Telephone number", c.get("telephone"), None),
            ("Cellphone number", c.get("cellphone"), None),
            ("E-mail address", c.get("email"), None),
            ("VAT registration number", c.get("vat_number"), None),
            ("TCS PIN", c.get("sars_tcs_pin"), None),
            ("CSD number", c.get("csd_supplier_number"), None),
            ("B-BBEE proof attached", proof, "tick the matching box"),
            ("Accredited representative in SA?", _yn(a.get("sbd1.accredited_representative")), None),
            ("Foreign-based supplier?", _yn(a.get("sbd1.foreign_based")) or ("No" if c.get("cipc_registration_number") else None),
             None if a.get("sbd1.foreign_based") else "assumed No - CIPC-registered SA company"),
            ("Capacity under which the bid is signed", s.get("capacity"), "attach the board resolution / proof of authority"),
            ("Total bid price", a.get("total_bid_price"), "from the priced schedule"),
            ("Signature and date", None, "sign by hand at submission"),
        ]
    if code == "4":
        return [
            ("Employed by the state?", _yn(a.get("sbd4.employed_by_state")), None),
            ("Relationship with an employee of this buyer?", _yn(a.get("sbd4.relationship_with_employee")), None),
            # The form asks about ANY related enterprise "whether or not they are
            # bidding for this contract" (Regulatory review G4b; renamed
            # 2026-10-18 with the general-harness baseline reset).
            ("Interest in any other related enterprise, bidding or not?", _yn(a.get("sbd4.interest_in_related_enterprise")),
             None if a.get("sbd4.interest_in_related_enterprise") else
             "not answered - if Yes, list every CSD-registered company linked to the directors"),
            ("Declarant name (I, ..., the undersigned)", s.get("name"), None),
            ("Designation", s.get("capacity"), None),
            ("Name of bidder", c.get("legal_name"), None),
            ("Signature and date", None, "sign by hand"),
        ]
    if code == "6.1":
        return [
            ("Name of company/firm", c.get("legal_name"), None),
            ("Company registration number", c.get("cipc_registration_number"), None),
            ("VAT registration number", c.get("vat_number"), None),
            ("Type of company/firm", "(Pty) Limited" if re.search(r"\(Pty\)", c.get("legal_name") or "") else None, "tick it"),
            ("Specific-goal points claimed", a.get("sbd6_1.points_claimed"), "enter in the buyer's own table, attach the proof"),
            ("Bidder's representative", s.get("name"), None),
            ("Address", c.get("registered_address"), None),
            ("Signature and date", None, "sign by hand"),
        ]
    if code == "8":
        return [
            ("Listed on the Database of Restricted Suppliers?", _yn(a.get("sbd8.restricted_supplier")), None),
            ("Listed on the Register for Tender Defaulters?", _yn(a.get("sbd8.tender_defaulter")), None),
            ("Convicted of fraud/corruption in the past 5 years?", _yn(a.get("sbd8.convicted")), None),
            ("Contract with an organ of state terminated in the past 5 years?", _yn(a.get("sbd8.contract_terminated")), None),
            ("Signature, position, name of bidder, date", None, "complete by hand"),
        ]
    if code == "9":
        return [
            ("Bid number and description", " - ".join(x for x in [t.get("contract_no"), t.get("description")] if x) or None, None),
            ("Name of institution", t.get("buyer"), None),
            ("On behalf of (name of bidder)", c.get("legal_name"), None),
            ("Signature, position, date", None, "complete by hand"),
        ]
    return []


def _field_label(key):
    return key.split(".", 1)[-1].replace("_", " ")


def _log_overlay(log, reports):
    """Provenance for every field fill_flat_pdf looked for: typed (and read
    back) -> confirmed / assumed_default; anything else -> missing with the
    reason and the value to write by hand. The buyer's own pre-print is not
    logged (nothing for the company to do)."""
    for code, rep in reports.items():
        prefix = None
        for key, f in rep.get("fields", {}).items():
            prefix = key.split(".", 1)[0]
            if f["status"] == "placed":
                if f.get("reason"):
                    log.assumed_default(key, f.get("value"), f["reason"])
                else:
                    log.confirmed(key, f.get("value"))
            elif f["status"] in ("prefilled", "not_found"):
                continue  # the buyer printed it / this edition has no such field
            elif f.get("value"):
                log.missing(key, f"Typed-onto-PDF: {f.get('reason') or 'not placed'} - write by hand: {f['value']}")
            else:
                log.missing(key, f"Not on file - write it by hand on the printed form ({f.get('reason') or 'no value'}).")
        if prefix:
            log.missing(f"{prefix}.signature_and_date", "Needs a real signature and date - never typed in.")


def write_guide(path, manifest, ctx, form_results, by_hand):
    d = docx.Document()
    sec = d.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.8)
    t = ctx.get("tender") or {}
    h = d.add_paragraph("COMPLETION GUIDE")
    h.runs[0].bold = True
    h.runs[0].font.size = Pt(15)
    d.add_paragraph(f"{t.get('contract_no') or ''} - {t.get('buyer') or ''}".strip(" -"))
    if t.get("description"):
        d.add_paragraph(t["description"][:300])
    p = d.add_paragraph(
        "Generated by Tender Tool from your company record and the answers on the Tender Workspace. "
        "Forms marked 'filled' are already completed in this pack; for forms marked 'by hand', copy the values "
        "below onto the buyer's original form. Blank values were not on file - nothing has been guessed.")
    p.runs[0].italic = True
    p.runs[0].font.size = Pt(9)

    for f in manifest["forms"]:
        if f["code"] not in FILLABLE_CODES:
            continue
        status = form_results.get(f["code"], {}).get("status") or ("filled" if f["status"] == "auto" else "by hand")
        placed = form_results.get(f["code"], {}).get("placed") or []
        hd = d.add_paragraph(f"{f['label']} - {f['title']} ({status})")
        hd.runs[0].bold = True
        tbl = d.add_table(rows=1, cols=3)
        tbl.style = "Table Grid"
        for cell, label in zip(tbl.rows[0].cells, ("Field", "Value", "Note")):
            cell.text = label
        for field, value, note in guide_rows(f["code"], ctx):
            row = tbl.add_row().cells
            row[0].text = field
            row[1].text = value or ("(by hand)" if note and ("by hand" in note or "sign" in note) else "- not on file -")
            row[2].text = note or ""
        if placed:
            p = d.add_paragraph("Typed onto the printed form (check each before signing): " +
                                ", ".join(_field_label(k) for k in sorted(placed)) +
                                ". Everything else in the table above is still to be written by hand.")
            p.runs[0].font.size = Pt(8.5)
            p.runs[0].italic = True
        for row in tbl.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(8.5)
    if by_hand:
        hd = d.add_paragraph("STILL TO DO BY HAND")
        hd.runs[0].bold = True
        for b in by_hand:
            para = d.add_paragraph(f"{b['title']}: {b['reason']}", style="List Bullet")
            para.runs[0].font.size = Pt(9)
    d.save(path)
    return path


# ---------------------------------------------------------------------------
# The pack
# ---------------------------------------------------------------------------

def _doc_path(manifest_doc, local_docs):
    """manifest documents carry names; local_docs maps name -> local path
    (the downloaded copy, or the .docx LibreOffice converted from a .doc)."""
    return local_docs.get(manifest_doc)


def build_general_pack(*, manifest, local_docs, company, signatory, answers, tender, files, work_dir, out_pdf,
                       evidence_items=(), cvs=()):
    """
    manifest   requirements.extract_requirements() output (stripped or not)
    local_docs {manifest document name: local path} for the tender's own documents
    company    companies-row-shaped dict (0001 + 0014 + 0018 columns)
    signatory  {name, capacity} or None
    answers    {tender_answers.field_key: value}
    tender     {contract_no, description, buyer, closing_date}
    files      {"vault:<companies column>": local path, "upload:<returnable key>": local path}
    evidence_items  [{title, path, document_type}] already filtered to the tender's rule
    cvs        [{title, path}] for people tagged for the tender's roles
    Returns {"manifest": merged-pack manifest, "provenance": [entries], "outputs": {name: path}}.
    """
    os.makedirs(work_dir, exist_ok=True)
    conv_dir = os.path.join(work_dir, "pdf")
    filled_dir = os.path.join(work_dir, "filled_forms")
    os.makedirs(filled_dir, exist_ok=True)
    ctx = {"company": company, "signatory": signatory, "answers": answers, "tender": tender}
    log = ProvenanceLog(source="engine/general/build_general_pack.py")
    sections, by_hand, outputs = [], [], {}
    form_results = {}

    # --- 1. standard forms: fill every editable edition (bundles once) -----
    editable_docs = [d for d in manifest["documents"] if d.get("editable") and d.get("forms")]
    done_codes = set()
    for d in editable_docs:
        src = _doc_path(d["name"], local_docs)
        codes = [c for c in d["forms"] if c in FILLABLE_CODES and c not in done_codes]
        if not src or not codes:
            continue
        base = re.sub(r"[^A-Za-z0-9._-]+", "_", os.path.splitext(os.path.basename(d["name"]))[0])[:60]
        out = os.path.join(filled_dir, f"{base}_FILLED.docx")
        try:
            fill_sbd_docx.fill_sbd_file(src, out, ctx, log, set(codes))
        except Exception as e:  # a buyer edition python-docx chokes on
            for c in codes:
                form_results[c] = {"status": "by hand", "reason": f"could not fill {d['name']}: {type(e).__name__}"}
            continue
        outputs[os.path.basename(out)] = out
        pdf, why = to_pdf(out, conv_dir)
        labels = ", ".join(f"{(manifest.get('form_suite') or 'SBD')} {c}" for c in codes)
        sections.append({"title": f"{labels} - filled ({d['name']})", "envelope": FORMS_ENV, "pdf": pdf,
                         "omitted_reason": None if pdf else f"filled .docx could not be converted to PDF ({why}) - print the .docx"})
        for c in codes:
            form_results[c] = {"status": "filled"}
            done_codes.add(c)

    # --- 1b. flat PDF editions: type onto the printed form ----------------
    # (requirements.py marked these "overlay": the form's page is in the PDF's
    # text layer). One typed copy per PDF, every overlay form in it at once.
    overlay_jobs = {}
    for f in manifest["forms"]:
        if f["status"] == "overlay" and f["code"] not in done_codes and f.get("overlay"):
            overlay_jobs.setdefault(f["overlay"]["document"], []).append(f["code"])
    for b in manifest.get("buyer_forms", []):
        if b.get("overlay"):
            overlay_jobs.setdefault(b["overlay"]["document"], []).append(b["overlay"]["code"])
    overlaid, overlay_reports = {}, {}
    for doc_name, codes in overlay_jobs.items():
        src = _doc_path(doc_name, local_docs)
        if not src or not src.lower().endswith(".pdf"):
            continue
        base = re.sub(r"[^A-Za-z0-9._-]+", "_", os.path.splitext(os.path.basename(doc_name))[0])[:60]
        out = os.path.join(filled_dir, f"{base}_TYPED.pdf")
        try:
            rep = fill_flat_pdf.fill_flat_pdf(src, out, ctx, codes)
        except Exception as e:  # a PDF pdfplumber can't parse - stays by hand
            rep = {"forms": {c: {"pages": [], "fields": {}, "reason": f"could not read the PDF ({type(e).__name__})"} for c in codes},
                   "placed": 0, "written": False}
        for c in codes:
            overlay_reports[c] = rep["forms"].get(c, {"pages": [], "fields": {}})
        if rep["written"]:
            overlaid[doc_name] = out
            outputs[os.path.basename(out)] = out
    _log_overlay(log, overlay_reports)

    for f in manifest["forms"]:
        if f["code"] in done_codes:
            continue
        if f["status"] == "overlay" and f["code"] in overlay_reports:
            orep = overlay_reports[f["code"]]
            placed = [k for k, v in orep["fields"].items() if v["status"] == "placed"]
            left = [(k, v) for k, v in orep["fields"].items() if v["status"] in ("not_placed", "missing") and v.get("value")]
            doc_name = f["overlay"]["document"]
            if placed and doc_name in overlaid:
                form_results[f["code"]] = {"status": f"typed onto the PDF ({len(placed)} fields)", "placed": placed,
                                           "left": [k for k, _ in left]}
                standalone = next((d for d in manifest["documents"] if d["name"] == doc_name and d["kind"] == "standard_form"), None)
                if standalone:
                    pdf, why = to_pdf(overlaid[doc_name], conv_dir)
                    sections.append({"title": f"{f['label']} {f['title']} - details typed on, check and sign", "envelope": FORMS_ENV,
                                     "pdf": pdf, "omitted_reason": None if pdf else f"could not read the typed copy ({why})"})
                by_hand.append({"title": f"{f['label']} {f['title']}",
                                "reason": "check the typed details, then sign" + (
                                    "; still to write: " + ", ".join(_field_label(k) for k, _ in left[:8]) if left else "")})
                continue
            # nothing could be placed - falls through to the by-hand path below
        if f["status"] == "attach":
            form_results[f["code"]] = {"status": "from the vault"}
            continue
        reason = f["reason"]
        if f["code"] in ("7",) and not f.get("editions"):
            by_hand.append({"title": f"{f['label']} {f['title']}", "reason": reason})
            continue
        # a standalone (flat) edition of just this form -> include it blank
        standalone = next((e for e in f.get("editions", []) if
                           next((d for d in manifest["documents"] if d["name"] == e["document"] and d["kind"] == "standard_form"), None)), None)
        if standalone and _doc_path(standalone["document"], local_docs):
            pdf, why = to_pdf(_doc_path(standalone["document"], local_docs), conv_dir)
            sections.append({"title": f"{f['label']} {f['title']} - blank, complete by hand", "envelope": FORMS_ENV, "pdf": pdf,
                             "omitted_reason": None if pdf else f"could not convert the published form ({why})"})
        elif not f.get("editions"):
            # named but no copy found anywhere in the published documents
            sections.append({"title": f"{f['label']} {f['title']}", "envelope": FORMS_ENV, "pdf": None,
                             "omitted_reason": "no copy in the published documents - get it from the buyer and complete by hand"})
        # else: it sits inside a bundled tender document, which is included
        # whole below ("complete the forms inside by hand") - not "missing".
        by_hand.append({"title": f"{f['label']} {f['title']}", "reason": reason})
        form_results[f["code"]] = {"status": "by hand"}
        if f["code"] in FILLABLE_CODES and not overlay_reports.get(f["code"], {}).get("fields"):
            for field, value, note in guide_rows(f["code"], ctx):
                key = f"{f['label'].lower().replace(' ', '').replace('.', '_')}.{re.sub(r'[^a-z0-9]+', '_', field.lower()).strip('_')}"
                log.missing(key, f"{f['label']} is completed by hand ({reason[:120]})" + (f" - write: {value}" if value else " - not on file"))

    # --- 2. the tender document itself when its forms are done by hand ----
    # (municipal single-PDF bundles - Dawid Kruiper TN051, Overberg Q12 - are
    # returned whole, with the MBD forms inside completed by hand)
    for d in manifest["documents"]:
        if d["kind"] == "tender_document" and (d.get("forms") or d["name"] in overlaid) and not d.get("editable"):
            src = overlaid.get(d["name"]) or _doc_path(d["name"], local_docs)
            if src:
                pdf, why = to_pdf(src, conv_dir)
                title = (f"{d['name']} - your details typed onto its forms; check, complete the rest and sign"
                         if d["name"] in overlaid else f"{d['name']} - complete the forms inside by hand")
                sections.append({"title": title, "envelope": BUYER_ENV, "pdf": pdf,
                                 "omitted_reason": None if pdf else f"could not convert ({why})"})

    # --- 3. buyer's own forms (as published, by hand) --------------------
    for b in manifest.get("buyer_forms", []):
        src = _doc_path(b["document"], local_docs) if b["key"].startswith("doc:") else None
        if src:
            pdf, why = to_pdf(src, conv_dir)
            sections.append({"title": f"{b['label']} - complete by hand", "envelope": BUYER_ENV, "pdf": pdf,
                             "omitted_reason": None if pdf else f"could not convert ({why}) - print the original"})
        ov = b.get("overlay")
        orep = overlay_reports.get(ov["code"]) if ov else None
        placed = [k for k, v in (orep or {}).get("fields", {}).items() if v["status"] == "placed"]
        if placed and ov["document"] in overlaid:
            form_results[ov["code"]] = {"status": f"typed onto the PDF ({len(placed)} fields)", "placed": placed}
            by_hand.append({"title": b["label"], "reason": f"your company details are typed on (p.{ov['page']} of {ov['document']}) - "
                                                           f"check them; {ov['by_hand']} stay by hand"})
        else:
            by_hand.append({"title": b["label"], "reason": b["reason"]})

    # --- 4. company documents -------------------------------------------
    closing = tender.get("closing_date")
    bidding_jv = (answers.get("bidding_as_jv") or "").lower()
    for r in manifest.get("returnables", []):
        if r.get("conditional") == "jv" and bidding_jv != "yes":
            continue
        path = files.get("upload:" + r["key"])
        col = VAULT_FILE_FOR.get(r["key"])
        if not path and col:
            path = files.get("vault:" + col)
        title = r["label"]
        if not path:
            reason = ("not uploaded to the vault" if col else "not uploaded for this tender") + \
                     (" (the tender only mentions it - check whether it applies)" if r["strength"] == "mentioned" else "")
            sections.append({"title": title, "envelope": DOCS_ENV, "pdf": None, "omitted_reason": reason})
            continue
        pdf, why = to_pdf(path, conv_dir)
        exp_col = VAULT_EXPIRY_FOR.get(r["key"])
        # An expired vault copy is still included, with a checklist WARNING
        # (pdf_merge.expiry_warning, same rule as the SANRAL pack) - a
        # per-tender upload replaces the vault copy, so its date isn't known.
        warning = expiry_warning(company.get(exp_col), closing) if exp_col and not files.get("upload:" + r["key"]) else None
        # Content-sanity check (see VAULT_MISMATCH_KEYWORDS's own comment) -
        # only against the vault copy's own filename: a per-tender upload
        # (files.get("upload:" + r["key"])) was chosen fresh for THIS tender,
        # not pulled blind from a vault slot that might hold the wrong thing.
        if not warning and col and not files.get("upload:" + r["key"]):
            warning = vault_document_mismatch_warning(col, path)
        if warning:
            by_hand.append({"title": title, "reason": f"the vault copy {warning[0].lower() + warning[1:]}"})
        sections.append({"title": title, "envelope": DOCS_ENV, "pdf": pdf, "warning": warning if pdf else None,
                         "omitted_reason": None if pdf else f"could not convert ({why})"})

    # --- 5. experience evidence -----------------------------------------
    for item in evidence_items:
        pdf, why = to_pdf(item.get("path"), conv_dir)
        sections.append({"title": item["title"], "envelope": EXP_ENV, "pdf": pdf,
                         "omitted_reason": None if pdf else (item.get("missing_reason") or f"could not convert ({why})")})

    # --- 6. CVs ----------------------------------------------------------
    for cv in cvs:
        pdf, why = to_pdf(cv.get("path"), conv_dir)
        sections.append({"title": cv["title"], "envelope": CV_ENV, "pdf": pdf,
                         "omitted_reason": None if pdf else (cv.get("missing_reason") or f"could not convert ({why})")})

    for x in ("signature and date on every form", "initial every page if the buyer asks for it"):
        by_hand.append({"title": "Sign", "reason": x})

    # --- guide first ------------------------------------------------------
    guide_docx = write_guide(os.path.join(work_dir, "00b_Completion_guide.docx"), manifest, ctx, form_results, by_hand)
    outputs["Completion_guide.docx"] = guide_docx
    guide_pdf, why = to_pdf(guide_docx, conv_dir)
    sections.insert(0, {"title": "Completion guide", "envelope": GUIDE_ENV, "pdf": guide_pdf,
                        "omitted_reason": None if guide_pdf else f"could not render ({why})"})

    heading = [x for x in [tender.get("contract_no"), tender.get("buyer"), (tender.get("description") or "")[:120], company.get("legal_name")] if x]
    notes = [
        "Generated by Tender Tool. Signatures, dates and prices are never typed in for you.",
        f"Standard forms: {sum(1 for v in form_results.values() if v['status'] == 'filled')} filled from your data, "
        f"{sum(1 for v in form_results.values() if v['status'].startswith('typed'))} typed onto the buyer's printed PDF, "
        f"{sum(1 for v in form_results.values() if v['status'] == 'by hand')} to complete by hand (see the completion guide).",
    ] + [f"Gap: {g}" for g in manifest.get("gaps", [])[:4]]
    pack_manifest = merge_with_contents(sections, out_pdf, os.path.join(work_dir, "contents"), heading, notes=notes,
                                        front_count=0, stamp_label=(tender.get("contract_no") or "")[:40])
    problems = verify_pack(out_pdf, pack_manifest)
    pack_manifest["verify_problems"] = problems
    pack_manifest["by_hand"] = by_hand
    pack_manifest["form_results"] = form_results
    pack_manifest["provenance_summary"] = log.summary()
    prov_path = os.path.join(work_dir, "field_provenance.json")
    log.save(prov_path)
    outputs["field_provenance.json"] = prov_path
    return {"manifest": pack_manifest, "provenance": log.entries, "outputs": outputs}


def local_docs_for(manifest, doc_paths, work_dir):
    """Re-runs requirements.collect_documents()/read_document() over the
    downloaded files so manifest document names map back to local paths (a
    .doc maps to its LibreOffice .docx conversion, a zip member to its
    extracted file)."""
    out = {}
    for d in req.collect_documents(doc_paths, work_dir):
        if d["format"] == "doc":
            conv = req.convert_with_soffice(d["path"], "docx", work_dir)
            out[d["name"]] = conv or d["path"]
        else:
            out[d["name"]] = d["path"]
    return out


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser(description="Build a general submission pack from local files (no Supabase).")
    ap.add_argument("--docs", nargs="+", required=True, help="the tender's documents (files/dirs/zips)")
    ap.add_argument("--profile", required=True, help="a config/test_companies/*.json style profile")
    ap.add_argument("--answers", help="JSON file {field_key: value}")
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    work = os.path.abspath(args.out)
    os.makedirs(work, exist_ok=True)
    m = req.extract_requirements(args.docs, work_dir=os.path.join(work, "src"))
    prof = json.load(open(args.profile))
    comp = dict(prof["company"])
    person = prof["key_people"][0]
    res = build_general_pack(manifest=m, local_docs=local_docs_for(m, args.docs, os.path.join(work, "src2")), company=comp,
                             signatory={"name": person["name"], "capacity": person.get("signatory_capacity")},
                             answers=json.load(open(args.answers)) if args.answers else {},
                             tender={"contract_no": "LOCAL", "buyer": None, "description": None, "closing_date": None},
                             files={}, work_dir=work, out_pdf=os.path.join(work, "submission_pack.pdf"))
    print(json.dumps({k: v for k, v in res["manifest"].items() if k != "sections"}, indent=1)[:3000])
