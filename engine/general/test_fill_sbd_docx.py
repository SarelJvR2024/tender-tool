"""
Regression test for fill_sbd_docx.py's fill_sbd_document() - the actual SBD
1/4/6.1/8/9 field-filling logic used on every non-SANRAL general-tender
submission. Beyond bee_proof_lapse() and _today_iso() (already covered by
test_form_field_audit.py), no dedicated test file existed for this module at
all before this pass (confirmed by grep) - answer_yes_no() in particular, the
most complex piece (four different real-world Yes/No layouts, regex-matched
against real table/paragraph text), had zero coverage of its own.

Builds small, realistic docx fixtures with python-docx directly, matching
the four layouts answer_yes_no()'s own docstring/comments name:
  1. Separate Yes/No cells (Mthashana SBD 8 4.1-4.4)
  2. A single "YES / NO" slash cell (DSTI SBD 4 2.1-2.3)
  3. One cell containing both words inline (SBD 1 Part B)
  4. A plain paragraph with no table at all (older SBD 4 edition)

Run: python3 engine/general/test_fill_sbd_docx.py
"""
import os
import sys

import docx

HERE = os.path.dirname(os.path.abspath(__file__))
ENGINE = os.path.dirname(HERE)
sys.path.insert(0, ENGINE)
import fill_sbd_docx as sbd  # noqa: E402
from provenance import ProvenanceLog  # noqa: E402

FAILS = []


def check(name, cond, detail=""):
    print(("  ok   " if cond else "  FAIL ") + name + (f"  ({detail})" if detail and not cond else ""))
    if not cond:
        FAILS.append(name)


def _entry(log, key):
    return next((e for e in log.entries if e["field_key"] == key), None)


# ----------------------------------------------------------------------------
# answer_yes_no(): the four real layouts
# ----------------------------------------------------------------------------

def test_separate_yes_no_cells():
    print("answer_yes_no(): separate Yes/No cells (Mthashana SBD 8 layout)")
    d = docx.Document()
    t = d.add_table(rows=1, cols=3)
    t.rows[0].cells[0].text = "Was the bidder convicted for fraud?"
    t.rows[0].cells[1].text = "Yes"
    t.rows[0].cells[2].text = "No"
    n = sbd.answer_yes_no(d, r"convicted\s+for\s+fraud", "yes")
    check("exactly one place marked", n == 1, n)
    check("the Yes cell is ticked", "[X] Yes" in t.rows[0].cells[1].text, t.rows[0].cells[1].text)
    check("the No cell is left untouched", t.rows[0].cells[2].text == "No", t.rows[0].cells[2].text)

    d2 = docx.Document()
    t2 = d2.add_table(rows=1, cols=3)
    t2.rows[0].cells[0].text = "Was the bidder convicted for fraud?"
    t2.rows[0].cells[1].text = "Yes"
    t2.rows[0].cells[2].text = "No"
    sbd.answer_yes_no(d2, r"convicted\s+for\s+fraud", "no")
    check("a 'no' answer ticks the No cell, not the Yes cell",
          t2.rows[0].cells[2].text == "[X] No" and t2.rows[0].cells[1].text == "Yes", t2.rows[0].cells[2].text)


def test_yes_slash_no_cell():
    print("answer_yes_no(): a single 'YES / NO' slash cell (DSTI SBD 4 layout)")
    d = docx.Document()
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Are you presently employed by the state?"
    t.rows[0].cells[1].text = "YES / NO"
    n = sbd.answer_yes_no(d, r"employed\s+by\s+the\s+state", "yes")
    check("exactly one place marked", n == 1, n)
    check("the cell is replaced with the plain answer", t.rows[0].cells[1].text == "YES", t.rows[0].cells[1].text)


def test_inline_yes_no_cell():
    print("answer_yes_no(): one cell with both words inline (SBD 1 Part B layout)")
    d = docx.Document()
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Is the bidder an accredited representative?"
    t.rows[0].cells[1].text = "Yes   No"
    n = sbd.answer_yes_no(d, r"accredited\s+representative", "no")
    check("exactly one place marked", n == 1, n)
    check("the No in the inline cell is marked, the Yes left alone",
          "[X] No" in t.rows[0].cells[1].text and "[X] Yes" not in t.rows[0].cells[1].text,
          t.rows[0].cells[1].text)


def test_paragraph_only_layout():
    print("answer_yes_no(): a plain paragraph, no table at all (older SBD 4 edition)")
    d = docx.Document()
    d.add_paragraph("Are you presently in the service of the state? YES/NO")
    n = sbd.answer_yes_no(d, r"presently\s+in\s+the\s+service\s+of\s+the\s+state", "no")
    check("exactly one place marked", n == 1, n)
    check("the answer is appended to the paragraph", d.paragraphs[0].text.endswith("(No)"), d.paragraphs[0].text)


def test_paragraph_layout_only_marks_the_first_repeat():
    print("answer_yes_no(): a repeated prose bullet (Mthashana's older edition) only marks the FIRST match")
    d = docx.Document()
    d.add_paragraph("Bidders must declare whether they are employed by the state? YES/NO")
    d.add_paragraph("Note: being employed by the state in a different capacity does not apply.")
    n = sbd.answer_yes_no(d, r"employed\s+by\s+the\s+state", "yes")
    check("only the real question paragraph (with '?'/YES/NO) is marked, not the unrelated bullet repeating the same phrase",
          n == 1 and d.paragraphs[0].text.endswith("(Yes)") and "(Yes)" not in d.paragraphs[1].text,
          [p.text for p in d.paragraphs])


def test_no_match_returns_zero_not_a_crash():
    print("answer_yes_no(): a question this edition simply doesn't have - returns 0, never raises")
    d = docx.Document()
    d.add_paragraph("Some unrelated paragraph about delivery addresses.")
    n = sbd.answer_yes_no(d, r"interest\s+in\s+any\s+other\s+related\s+enterprise", "no")
    check("no match -> 0, no exception", n == 0)


# ----------------------------------------------------------------------------
# SBD4/SBD8: an unanswered declaration - assumed_default vs never-assumed
# ----------------------------------------------------------------------------

def _sbd4_doc():
    d = docx.Document()
    for q in ("Are you presently in the service of the state?",
              "Do you have any relationship with any person who is employed by the procuring institution?",
              "Do you have any interest in any other related enterprise or company?"):
        t = d.add_table(rows=1, cols=3)
        t.rows[0].cells[0].text = q
        t.rows[0].cells[1].text = "Yes"
        t.rows[0].cells[2].text = "No"
    return d


def test_sbd4_unanswered_question_is_assumed_default_no_never_silently_confirmed():
    print("SBD4: an unanswered declaration defaults to ticking No, but the provenance says "
          "assumed_default, never confirmed - the document and the audit trail must not disagree")
    d = _sbd4_doc()
    log = ProvenanceLog("test")
    sbd.fill_sbd_document(d, {"company": {}, "answers": {}, "tender": {}}, log, {"4"})
    entry = _entry(log, "sbd4.employed_by_state")
    check("the provenance entry exists and is assumed_default, not confirmed",
          entry is not None and entry["status"] == "assumed_default", entry)
    check("the assumed value is 'No', with a reason a reviewer can act on",
          entry["value"] == "No" and entry["reason"], entry)
    no_cell = d.tables[0].rows[0].cells[2]
    check("the document itself shows No ticked (matches what was logged as assumed, not left blank)",
          "[X] No" in no_cell.text, no_cell.text)


def test_sbd4_real_answer_is_confirmed_not_assumed():
    print("SBD4: a real answer on file is logged confirmed, with the actual value - not downgraded to assumed")
    d = _sbd4_doc()
    log = ProvenanceLog("test")
    sbd.fill_sbd_document(d, {"company": {}, "answers": {"sbd4.employed_by_state": "no"}, "tender": {}}, log, {"4"})
    entry = _entry(log, "sbd4.employed_by_state")
    check("confirmed, with the real answer capitalised", entry["status"] == "confirmed" and entry["value"] == "No", entry)


def _sbd8_doc():
    d = docx.Document()
    for q in ("Is the bidder listed on the Database of Restricted Suppliers?",
              "Is the bidder listed on the Register for Tender Defaulters?",
              "Was the bidder convicted for fraud?",
              "Has any contract between the bidder and any organ of state terminated due to non-performance?"):
        t = d.add_table(rows=1, cols=3)
        t.rows[0].cells[0].text = q
        t.rows[0].cells[1].text = "Yes"
        t.rows[0].cells[2].text = "No"
    return d


def test_sbd8_unanswered_question_is_never_assumed_unlike_sbd4():
    print("SBD8: unlike SBD4, an unanswered restricted-supplier declaration is NEVER assumed - "
          "left entirely blank on the document, flagged missing, not ticked at all")
    d = _sbd8_doc()
    log = ProvenanceLog("test")
    sbd.fill_sbd_document(d, {"company": {}, "answers": {}, "tender": {}}, log, {"8"})
    entry = _entry(log, "sbd8.restricted_supplier")
    check("the provenance entry is 'missing', not assumed or confirmed", entry["status"] == "missing", entry)
    yes_cell, no_cell = d.tables[0].rows[0].cells[1], d.tables[0].rows[0].cells[2]
    check("neither Yes nor No is ticked on the actual document - a real person must answer this one",
          yes_cell.text == "Yes" and no_cell.text == "No", (yes_cell.text, no_cell.text))


def test_sbd8_real_answer_is_confirmed():
    print("SBD8: a real 'no' answer on file is actually ticked and logged confirmed")
    d = _sbd8_doc()
    log = ProvenanceLog("test")
    sbd.fill_sbd_document(d, {"company": {}, "answers": {"sbd8.restricted_supplier": "no"}, "tender": {}}, log, {"8"})
    entry = _entry(log, "sbd8.restricted_supplier")
    check("confirmed", entry["status"] == "confirmed" and entry["value"] == "No", entry)
    no_cell = d.tables[0].rows[0].cells[2]
    check("the No cell is actually ticked", "[X] No" in no_cell.text, no_cell.text)


# ----------------------------------------------------------------------------
# fill_cells(): the CSD "MAAA" prefix-ok exception
# ----------------------------------------------------------------------------

def test_fill_cells_prefill_ok_lets_a_known_prefix_through():
    print("fill_cells(): a cell pre-printed with DSTI's own 'MAAA' CSD prefix is still filled, "
          "a cell with any OTHER existing text is left alone (never overwritten)")
    d = docx.Document()
    t1 = d.add_table(rows=1, cols=2)
    t1.rows[0].cells[0].text = "CSD No"
    t1.rows[0].cells[1].text = "MAAA"
    n = sbd.fill_cells(d, r"CSD\s+No", "MAAA0999999", "any", prefill_ok=r"MAAA")
    check("the pre-printed MAAA prefix is overwritten with the real CSD number", n == 1
          and t1.rows[0].cells[1].text == "MAAA0999999", t1.rows[0].cells[1].text)

    d2 = docx.Document()
    t2 = d2.add_table(rows=1, cols=2)
    t2.rows[0].cells[0].text = "CSD No"
    t2.rows[0].cells[1].text = "Some other pre-existing text"
    n2 = sbd.fill_cells(d2, r"CSD\s+No", "MAAA0999999", "any", prefill_ok=r"MAAA")
    check("a cell with unrelated existing text is never overwritten", n2 == 0
          and t2.rows[0].cells[1].text == "Some other pre-existing text", t2.rows[0].cells[1].text)


def test_whitespace_only_value_is_never_placed_or_confirmed():
    print("fill_sbd_document(): a whitespace-only value (a data-entry glitch) is never written "
          "as invisible spaces, and never logged 'confirmed' - the same real shape Finding 95 "
          "found independently in fill_flat_pdf.py")
    for bad_value, label in (("   ", "plain spaces"), ("\t\n ", "tab/newline only")):
        # Bidder table (fill_cells, scope='bidder') + a blank-paragraph layout
        # (fill_blank_paragraphs) in the same document - both of this file's
        # two main writer shapes, in one real end-to-end fill_sbd_document() pass.
        d = docx.Document()
        t = d.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text = "NAME OF BIDDER"
        t.rows[0].cells[1].text = ""
        d.add_paragraph("POSTAL ADDRESS: ................................")
        log = ProvenanceLog("test")
        ctx = {"company": {"legal_name": bad_value, "postal_address": bad_value}, "signatory": {},
              "answers": {}, "tender": {}}
        sbd.fill_sbd_document(d, ctx, log, {"1"})

        check(f"[{label}] the bidder table cell is left blank, not filled with invisible whitespace",
              t.rows[0].cells[1].text == "", repr(t.rows[0].cells[1].text))
        check(f"[{label}] the dotted paragraph is left untouched, not filled with invisible whitespace",
              d.paragraphs[0].text == "POSTAL ADDRESS: ................................",
              repr(d.paragraphs[0].text))
        name_entry = _entry(log, "sbd1.name_of_bidder")
        addr_entry = _entry(log, "sbd1.postal_address")
        check(f"[{label}] name_of_bidder is logged missing, never 'confirmed' with a whitespace value",
              name_entry is not None and name_entry["status"] == "missing", name_entry)
        check(f"[{label}] postal_address is logged missing, never 'confirmed' with a whitespace value",
              addr_entry is not None and addr_entry["status"] == "missing", addr_entry)

    # A real, non-whitespace value (including one with incidental surrounding
    # whitespace) must still be placed exactly as before - the fix must not
    # over-correct into refusing genuine values.
    d = docx.Document()
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "NAME OF BIDDER"
    t.rows[0].cells[1].text = ""
    log = ProvenanceLog("test")
    ctx = {"company": {"legal_name": "  Mzansi Roads & Civils (Pty) Ltd  "}, "signatory": {},
          "answers": {}, "tender": {}}
    sbd.fill_sbd_document(d, ctx, log, {"1"})
    check("a real value (even one with incidental surrounding whitespace) is still placed exactly as given",
          t.rows[0].cells[1].text == "  Mzansi Roads & Civils (Pty) Ltd  ", repr(t.rows[0].cells[1].text))
    check("...and logged confirmed, not refused",
          _entry(log, "sbd1.name_of_bidder")["status"] == "confirmed", _entry(log, "sbd1.name_of_bidder"))


def test_a_value_with_a_literal_backslash_is_placed_verbatim_not_corrupted_or_crashed():
    print("fill_dotted_paragraph_count(): a value containing a literal backslash (a pasted Windows "
          "path fragment, an OCR artefact, a typo) is placed exactly as given, not silently "
          "corrupted via re.sub()'s own backreference syntax or crashed with re.error")
    # '\1' is a real backreference in a replacement STRING (even though this
    # pattern has no capture group at all, so the real failure here is
    # re.error - "invalid group reference 1" - not silent corruption; see
    # engine/fill_a_forms.py's own sibling bug/fix for the silent-corruption
    # shape, where the pattern DOES have a group 1 to splice in).
    for bad_value, label in (("Smith\\1 Construction", "backslash-digit (backreference syntax)"),
                              ("Acme\\g<bad> Ltd", "backslash-g (named-group syntax)"),
                              ("O'Brien & Sons", "control: an ordinary value with no backslash")):
        d = docx.Document()
        d.add_paragraph("Name of bidder ....................")
        try:
            n = sbd.fill_dotted_paragraph_count(d, "Name of bidder", bad_value)
        except Exception as e:
            check(f"[{label}] fill_dotted_paragraph_count() does not raise", False, f"{type(e).__name__}: {e}")
            continue
        check(f"[{label}] fill_dotted_paragraph_count() filled exactly one paragraph", n == 1,
              f"expected n=1, got n={n}")
        check(f"[{label}] the value is placed verbatim, not corrupted",
              bad_value in d.paragraphs[0].text, repr(d.paragraphs[0].text))


def main():
    test_separate_yes_no_cells()
    test_yes_slash_no_cell()
    test_inline_yes_no_cell()
    test_paragraph_only_layout()
    test_paragraph_layout_only_marks_the_first_repeat()
    test_no_match_returns_zero_not_a_crash()
    test_sbd4_unanswered_question_is_assumed_default_no_never_silently_confirmed()
    test_sbd4_real_answer_is_confirmed_not_assumed()
    test_sbd8_unanswered_question_is_never_assumed_unlike_sbd4()
    test_sbd8_real_answer_is_confirmed()
    test_fill_cells_prefill_ok_lets_a_known_prefix_through()
    test_whitespace_only_value_is_never_placed_or_confirmed()
    test_a_value_with_a_literal_backslash_is_placed_verbatim_not_corrupted_or_crashed()
    print(f"\n{len(FAILS)} failure(s)")
    sys.exit(1 if FAILS else 0)


if __name__ == "__main__":
    main()
