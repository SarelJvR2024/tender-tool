"""
End-to-end test of engine/general/process_general_jobs.py with no live
Supabase project (none exists in this environment): an in-memory fake of the
PostgREST + Storage calls the bridge makes, seeded with rows shaped exactly
like migration 0018's tables, and REAL tender documents from
templates/general_samples/ served as the tender's "download links".

It proves, by reading the results back rather than trusting exit codes:
  1. an extract_requirements job is claimed, the tender's documents are read,
     and tenders.document_requirements / evaluation_rules are written;
  2. a generate_pack job builds a merged PDF that pdf_merge.verify_pack()
     accepts, uploads it + the filled forms to 'submission-packs' under
     <company>/tenders/<tender>/<job>/, writes field_provenance for the pack,
     and the PDF actually contains the company's data and the human's answers;
  3. a job already claimed by another worker is not re-processed;
  4. a pack requested before the documents were read fails with a reason;
  5. a buyer with no SBD/MBD forms (ESKOM) still gets a pack (ordered
     documents + its own forms listed by hand), not a crash or a fake fill;
  6. an EXPIRED vault document is still put in the pack, with a warning on
     its checklist row and on the contents page (Sarel, 2026-10-17);
  7. private documents on a SHARED tender (migration 0032): one company's
     own files are read with the published ones, the result goes to THAT
     company's tender_packs row only, the shared tender row and another
     company's reading are untouched, and a row pointing outside the
     company's own folder is refused before anything is downloaded.

Usage: python3 engine/general/test_general_bridge.py
"""
import io
import json
import os
import shutil
import sys
import tempfile
import threading
import urllib.parse
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import process_general_jobs as bridge  # noqa: E402

SAMPLES = os.path.join(HERE, "..", "..", "templates", "general_samples")


class Fake:
    """GET/PATCH with eq. filters, POST upsert on a table, Storage GET/POST."""

    def __init__(self):
        self.tables, self.storage, self.lock = {}, {}, threading.Lock()
        fake = self

        class H(BaseHTTPRequestHandler):
            def log_message(self, *a):
                pass

            def _send(self, code, payload, raw=None):
                body = raw if raw is not None else json.dumps(payload).encode()
                self.send_response(code)
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)

            def _parts(self):
                u = urllib.parse.urlparse(self.path)
                q = urllib.parse.parse_qs(u.query)
                filters = {k: v[0][3:] for k, v in q.items() if v[0].startswith("eq.")}
                return u.path, filters, q

            def _body(self):
                n = int(self.headers.get("Content-Length") or 0)
                return self.rfile.read(n) if n else b""

            def do_GET(self):
                path, filters, _ = self._parts()
                if path.startswith("/storage/v1/object/"):
                    key = urllib.parse.unquote(path[len("/storage/v1/object/"):])
                    data = fake.storage.get(key)
                    return self._send(200, None, data) if data is not None else self._send(404, {"message": "not found"})
                table = path.split("/rest/v1/", 1)[1]
                with fake.lock:
                    rows = [dict(r) for r in fake.tables.get(table, []) if all(str(r.get(k)) == v for k, v in filters.items())]
                self._send(200, rows)

            def do_PATCH(self):
                path, filters, _ = self._parts()
                table = path.split("/rest/v1/", 1)[1]
                patch = json.loads(self._body() or b"{}")
                out = []
                with fake.lock:
                    for r in fake.tables.get(table, []):
                        if all(str(r.get(k)) == v for k, v in filters.items()):
                            r.update(patch)
                            out.append(dict(r))
                self._send(200, out)

            def do_POST(self):
                path, _, q = self._parts()
                data = self._body()
                if path.startswith("/storage/v1/object/"):
                    with fake.lock:
                        fake.storage[urllib.parse.unquote(path[len("/storage/v1/object/"):])] = data
                    return self._send(200, {"Key": path})
                table = path.split("/rest/v1/", 1)[1]
                keys = [k for k in q.get("on_conflict", [""])[0].split(",") if k]
                rows = json.loads(data)
                if isinstance(rows, dict):  # a plain insert of one row (PostgREST accepts either)
                    rows = [rows]
                with fake.lock:
                    t = fake.tables.setdefault(table, [])
                    for row in rows:
                        hit = next((r for r in t if all(r.get(k) == row.get(k) for k in keys)), None) if keys else None
                        if hit:
                            hit.update(row)
                        else:
                            t.append(dict(row))
                self._send(201, rows)

        self.httpd = ThreadingHTTPServer(("127.0.0.1", 0), H)
        self.url = f"http://127.0.0.1:{self.httpd.server_address[1]}"
        threading.Thread(target=self.httpd.serve_forever, daemon=True).start()

    def add(self, table, **row):
        self.tables.setdefault(table, []).append(row)
        return row

    def one(self, table, **f):
        return next(r for r in self.tables[table] if all(r.get(k) == v for k, v in f.items()))


def tiny_pdf(text):
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(72, 760, text)
    c.showPage()
    c.save()
    return buf.getvalue()


def seed(fake, tender_id, sample_dir, with_answers=True):
    files = sorted(os.listdir(sample_dir))
    fake.add("tenders", id=tender_id, slug=tender_id, contract_no=f"T-{tender_id}", buyer="Test buyer", description="Test tender",
             closing_date="2026-11-06T11:00:00Z", document_urls=[{"url": f"sample://{sample_dir}/{f}", "title": f} for f in files],
             document_requirements=None, evaluation_rules=None)
    fake.add("tender_packs", id=f"pack-{tender_id}", company_id="co1", tender_id=tender_id, status="suggested")
    if with_answers:
        for k, v in {"signatory_person_id": "p1", "total_bid_price": "R 98 765.43", "sbd1.accredited_representative": "no",
                     "sbd4.employed_by_state": "no", "sbd4.relationship_with_employee": "yes", "sbd4.interest_in_related_enterprise": "no",
                     "sbd6_1.points_claimed": "Level 1: 20 points"}.items():
            fake.add("tender_answers", company_id="co1", tender_id=tender_id, field_key=k, value=v, storage_path=None)
        fake.add("tender_answers", company_id="co1", tender_id=tender_id, field_key="upload:csd_report", value=None,
                 storage_path=f"co1/tenders/{tender_id}/u1-csd.pdf", original_filename="csd.pdf")
        fake.storage[f"company-documents/co1/tenders/{tender_id}/u1-csd.pdf"] = tiny_pdf("SYNTHETIC CSD REPORT")


def fetch(url):
    assert url.startswith("sample://")
    with open(url[len("sample://"):], "rb") as f:
        return f.read()


def main():
    from pypdf import PdfReader
    fake = Fake()
    client = bridge.Supabase(fake.url, "service-key")
    work = tempfile.mkdtemp(prefix="tt_general_bridge_")
    failures = []

    def check(cond, msg):
        print(("  ok   " if cond else "  FAIL ") + msg)
        if not cond:
            failures.append(msg)

    fake.add("companies", id="co1", legal_name="Synthetic Test Bidder (Pty) Ltd", postal_address="PO Box 9, Testville",
             registered_address="9 Test Road, Testville, Gauteng", vat_number="4999999999", csd_supplier_number="MAAA0999999",
             sars_tcs_pin="TESTPIN99", sars_tcs_pin_expiry="2027-06-24", sars_tcs_pin_document_path="co1/vault/sars/pin.pdf",
             cipc_registration_number="2019/999999/07", telephone="012 345 6789", cellphone="082 000 0000", email="bids@synthetic.example",
             bee_level=1, bee_proof_type="qse_affidavit", bee_proof_expiry="2027-09-09", bee_proof_document_path="co1/vault/bee/aff.pdf")
    fake.storage["company-documents/co1/vault/sars/pin.pdf"] = tiny_pdf("SYNTHETIC TCS PIN LETTER")
    fake.storage["company-documents/co1/vault/bee/aff.pdf"] = tiny_pdf("SYNTHETIC B-BBEE AFFIDAVIT")
    fake.add("people", id="p1", company_id="co1", name="Pat Synthetic", signatory_capacity="Managing Director", role="Director")

    print("1. extraction + pack on a real editable SBD bundle (DSTI RFQ15/10/2026)")
    seed(fake, "t-dsti", os.path.join(SAMPLES, "dsti_rfq15"))
    fake.add("pack_generation_jobs", id="j1", engine="general", kind="extract_requirements", status="requested", company_id="co1",
             tender_id="t-dsti", tender_pack_id="pack-t-dsti", requested_at="2026-10-16T10:00:00Z")
    bridge.poll_once(client, work, fetch)
    job = fake.one("pack_generation_jobs", id="j1")
    t = fake.one("tenders", id="t-dsti")
    check(job["status"] == "ready", f"extract job ready (got {job['status']} {job.get('error_message')})")
    m = t.get("document_requirements") or {}
    check([f["code"] for f in m.get("forms", []) if f["status"] == "auto"] == ["1", "4", "6.1"], "SBD 1/4/6.1 found, all fillable from the .docx bundle")
    check(any(q["key"] == "sbd4.relationship_with_employee" for q in m.get("questions", [])), "SBD 4 questions asked")
    check(t.get("evaluation_rules") is not None, "evaluation_rules written (empty before)")

    fake.add("pack_generation_jobs", id="j2", engine="general", kind="generate_pack", status="requested", company_id="co1",
             tender_id="t-dsti", tender_pack_id="pack-t-dsti", requested_at="2026-10-16T10:05:00Z")
    bridge.poll_once(client, work, fetch)
    job = fake.one("pack_generation_jobs", id="j2")
    check(job["status"] == "ready", f"pack job ready (got {job['status']} {job.get('error_message')})")
    key = f"submission-packs/{job.get('storage_path_pdf')}"
    check(job.get("storage_path_pdf", "").startswith("co1/tenders/t-dsti/j2/") and key in fake.storage, "pack PDF uploaded under <company>/tenders/<tender>/<job>/")
    if key in fake.storage:
        rd = PdfReader(io.BytesIO(fake.storage[key]))
        txt = " ".join(p.extract_text() or "" for p in rd.pages)
        check(len(rd.pages) == job["page_count"], f"page count matches ({len(rd.pages)})")
        check("SYNTHETIC CSD REPORT" in txt and "SYNTHETIC TCS PIN LETTER" in txt, "tender upload + vault document merged in")
        check("COMPLETION GUIDE" in txt, "completion guide included")
    filled = next((v for k, v in job.get("output_files", {}).items() if k.endswith("_FILLED.docx")), None)
    if filled:
        import docx as _docx
        sys.path.insert(0, HERE)
        from requirements import docx_text
        tmp = os.path.join(work, "readback.docx")
        open(tmp, "wb").write(fake.storage[f"submission-packs/{filled}"])
        ftxt = docx_text(tmp)
        check("Synthetic Test Bidder (Pty) Ltd" in ftxt and "4999999999" in ftxt and "TESTPIN99" in ftxt and "MAAA0999999" in ftxt,
              "filled .docx read back: name, VAT, TCS PIN, CSD number in the form itself")
        check("Pat Synthetic" in ftxt and "Managing Director" in ftxt, "filled .docx: chosen signatory + capacity")
        check("98 765.43" not in ftxt, "DSTI's SBD 1 edition has no TOTAL BID PRICE field - nothing invented")
    prov = [r for r in fake.tables.get("field_provenance", []) if r["tender_pack_id"] == "pack-t-dsti"]
    by = {r["field_key"]: r for r in prov}
    check(by.get("sbd4.relationship_with_employee", {}).get("value") == "Yes" and by["sbd4.relationship_with_employee"]["status"] == "confirmed",
          "a 'Yes' answer is recorded as confirmed Yes, not the old assumed No")
    check(by.get("sbd1.total_bid_price", {}).get("status") == "missing", "bid price logged missing for this edition (label absent), not confirmed")
    check(by.get("sbd1.telephone_number", {}).get("value") == "012 345 6789", "company telephone (0018 column) filled")
    check(any(k.endswith("_FILLED.docx") for k in job.get("output_files", {})), "filled .docx uploaded alongside")

    print("2. a job another worker already claimed is left alone")
    fake.add("pack_generation_jobs", id="j3", engine="general", kind="generate_pack", status="generating", company_id="co1",
             tender_id="t-dsti", tender_pack_id="pack-t-dsti", requested_at="2026-10-16T10:06:00Z")
    bridge.poll_once(client, work, fetch)
    check(fake.one("pack_generation_jobs", id="j3")["status"] == "generating", "status untouched")

    print("3. pack before extraction fails with a reason")
    seed(fake, "t-early", os.path.join(SAMPLES, "tourism_nb143"), with_answers=False)
    fake.add("pack_generation_jobs", id="j4", engine="general", kind="generate_pack", status="requested", company_id="co1",
             tender_id="t-early", tender_pack_id="pack-t-early", requested_at="2026-10-16T10:07:00Z")
    bridge.poll_once(client, work, fetch)
    j4 = fake.one("pack_generation_jobs", id="j4")
    check(j4["status"] == "failed" and "haven't been read" in (j4.get("error_message") or ""), "failed: documents not read yet")

    print("4. a buyer with no SBD/MBD forms (ESKOM E3689, zip of its own documents)")
    seed(fake, "t-eskom", os.path.join(SAMPLES, "eskom_e3689"))
    for jid, kind, at in (("j5", "extract_requirements", "10:08"), ("j6", "generate_pack", "10:09")):
        fake.add("pack_generation_jobs", id=jid, engine="general", kind=kind, status="requested", company_id="co1",
                 tender_id="t-eskom", tender_pack_id="pack-t-eskom", requested_at=f"2026-10-16T{at}:00Z")
    bridge.poll_once(client, work, fetch)
    m = fake.one("tenders", id="t-eskom")["document_requirements"] or {}
    check(m.get("form_suite") is None and not m.get("forms"), "no standard forms claimed")
    check(any("no SBD/MBD standard forms" in g for g in m.get("gaps", [])), "the gap is stated, not hidden")
    j6 = fake.one("pack_generation_jobs", id="j6")
    check(j6["status"] == "ready", f"pack still built ({j6['status']} {j6.get('error_message')})")
    check(any("NEC3" in b["title"] or "Form of Offer" in b["title"] for b in (j6.get("manifest") or {}).get("by_hand", [])),
          "ESKOM's own forms listed as by-hand")

    print("6. an expired vault document goes in WITH a warning (co2, TCS PIN lapsed 2026-01-31)")
    fake.add("companies", id="co2", legal_name="Second Synthetic Bidder (Pty) Ltd", registered_address="1 Other Road, Testville, Gauteng",
             csd_supplier_number="MAAA0888888", sars_tcs_pin="TESTPIN22", sars_tcs_pin_expiry="2026-01-31",
             sars_tcs_pin_document_path="co2/vault/sars/pin.pdf", bee_level=2, bee_proof_type="certificate", bee_proof_expiry="2027-03-01")
    fake.storage["company-documents/co2/vault/sars/pin.pdf"] = tiny_pdf("SECOND BIDDER EXPIRED TCS PIN LETTER")
    fake.add("people", id="p2", company_id="co2", name="Sam Second", signatory_capacity="Director", role="Director")
    fake.add("tender_packs", id="pack-co2-t-dsti", company_id="co2", tender_id="t-dsti", status="suggested")
    fake.add("pack_generation_jobs", id="j7", engine="general", kind="generate_pack", status="requested", company_id="co2",
             tender_id="t-dsti", tender_pack_id="pack-co2-t-dsti", requested_at="2026-10-16T10:10:00Z")
    bridge.poll_once(client, work, fetch)
    j7 = fake.one("pack_generation_jobs", id="j7")
    check(j7["status"] == "ready", f"pack ready ({j7['status']} {j7.get('error_message')})")
    cl = (j7.get("manifest") or {}).get("checklist") or []
    tax = [c for c in cl if "SECOND" not in c["title"] and c["status"] == "included_warning"]
    check(len(tax) == 1 and "EXPIRED 2026-01-31" in (tax[0]["note"] or ""), f"checklist: the expired PIN letter is included_warning ({[(c['title'], c['note']) for c in tax]})")
    check(all(c["status"] in ("included", "included_warning", "not_included") for c in cl) and len(cl) >= 5, f"checklist lists every document ({len(cl)} rows)")
    key7 = f"submission-packs/{j7.get('storage_path_pdf')}"
    if key7 in fake.storage:
        rd = PdfReader(io.BytesIO(fake.storage[key7]))
        flat = " ".join(" ".join((p.extract_text() or "").split()) for p in rd.pages)
        check("SECOND BIDDER EXPIRED TCS PIN LETTER" in flat, "the expired PIN letter itself IS in the PDF (not excluded)")
        check("INCLUDED - WARNING: EXPIRED 2026-01-31" in flat, "contents page shows the warning on that row")
        check("Pack checklist:" in flat, "contents page carries the checklist summary")

    print("7. private documents on a SHARED tender (0032)")
    # Shared tender: Discovery found it, the listing only has an advert.
    fake.storage["advert.pdf"] = tiny_pdf("ADVERT ONLY - collect the bid document from the buyer")
    advert = os.path.join(work, "advert.pdf")
    open(advert, "wb").write(fake.storage["advert.pdf"])
    fake.add("tenders", id="t-shared", slug="ocds-shared-1", contract_no="RFQ-SHARED", buyer="Shared buyer", description="Shared tender",
             closing_date="2026-11-06T11:00:00Z", owner_company_id=None, document_urls=[{"url": f"sample://{advert}", "title": "advert.pdf"}],
             document_requirements=None, evaluation_rules=None, intake=None, family="General (non-SANRAL)")
    fake.add("tender_packs", id="pack-co1-shared", company_id="co1", tender_id="t-shared", status="suggested", document_requirements=None, intake=None)
    fake.add("tender_packs", id="pack-co2-shared", company_id="co2", tender_id="t-shared", status="suggested", document_requirements=None, intake=None)
    # co1's own copy of the bid document (the real DSTI SBD bundle), uploaded privately.
    dsti = os.path.join(SAMPLES, "dsti_rfq15")
    for i, f in enumerate(sorted(os.listdir(dsti))):
        sp = f"private/co1/t-shared/{i}-{f}"
        fake.storage[f"tender-documents/{sp}"] = open(os.path.join(dsti, f), "rb").read()
        fake.add("tender_private_documents", id=f"pd{i}", company_id="co1", tender_id="t-shared", title=f, storage_path=sp,
                 created_at=f"2026-10-17T09:0{i}:00Z")
    for jid, co, at in (("j8", "co1", "10:11"), ("j9", "co2", "10:12")):
        fake.add("pack_generation_jobs", id=jid, engine="general", kind="extract_requirements", status="requested", company_id=co,
                 tender_id="t-shared", tender_pack_id=f"pack-{co}-shared", requested_at=f"2026-10-16T{at}:00Z")
    bridge.poll_once(client, work, fetch)
    j8, j9 = fake.one("pack_generation_jobs", id="j8"), fake.one("pack_generation_jobs", id="j9")
    check(j8["status"] == "ready" and j9["status"] == "ready", f"both reads ready ({j8['status']} {j8.get('error_message')} / {j9['status']} {j9.get('error_message')})")
    p1, p2, ts = fake.one("tender_packs", id="pack-co1-shared"), fake.one("tender_packs", id="pack-co2-shared"), fake.one("tenders", id="t-shared")
    m1 = p1.get("document_requirements") or {}
    check([f["code"] for f in m1.get("forms", []) if f["status"] == "auto"] == ["1", "4", "6.1"],
          "co1's own reading (with its private bid document) found SBD 1/4/6.1")
    check(sorted((p1.get("intake") or {}).get("private_documents") or []) == sorted(f"pd{i}" for i in range(len(os.listdir(dsti)))),
          "co1's intake records which private documents it read")
    shared_m = ts.get("document_requirements") or {}
    check(not shared_m.get("forms"), "the SHARED tender's reading (co2's job, published advert only) has no SBD forms - nothing from co1's files")
    shared_names = " ".join(d["name"] for d in shared_m.get("documents", []))
    check(not any(f in shared_names for f in os.listdir(dsti)) and "SBD" not in json.dumps(ts.get("evaluation_rules") or {}),
          "no private file name or rule reached the shared tenders row")
    check(p2.get("document_requirements") is None, "co2's pack has no private reading")
    check(not any(r.get("tender_id") == "t-shared" for r in fake.tables.get("tender_documents", [])), "nothing registered in shared tender_documents")
    # co1's pack is built from its private reading: SBD 1 filled with co1's data.
    fake.add("tender_answers", company_id="co1", tender_id="t-shared", field_key="signatory_person_id", value="p1", storage_path=None)
    fake.add("pack_generation_jobs", id="j10", engine="general", kind="generate_pack", status="requested", company_id="co1",
             tender_id="t-shared", tender_pack_id="pack-co1-shared", requested_at="2026-10-16T10:13:00Z")
    bridge.poll_once(client, work, fetch)
    j10 = fake.one("pack_generation_jobs", id="j10")
    check(j10["status"] == "ready", f"co1's pack ready ({j10['status']} {j10.get('error_message')})")
    fr = (j10.get("manifest") or {}).get("form_results") or {}
    check(fr.get("1", {}).get("status") == "filled", f"co1's pack fills SBD 1 from the private bid document ({fr.get('1')})")
    # A row pointing at another company's folder is refused before any download.
    fake.add("tender_private_documents", id="bad", company_id="co2", tender_id="t-shared", title="x",
             storage_path="private/co1/t-shared/0-stolen.pdf", created_at="2026-10-17T10:00:00Z")
    fake.add("pack_generation_jobs", id="j11", engine="general", kind="extract_requirements", status="requested", company_id="co2",
             tender_id="t-shared", tender_pack_id="pack-co2-shared", requested_at="2026-10-16T10:14:00Z")
    bridge.poll_once(client, work, fetch)
    j11 = fake.one("pack_generation_jobs", id="j11")
    check(j11["status"] == "failed" and "not under this company's own folder" in (j11.get("error_message") or ""),
          f"a private-document row outside the company's own folder is refused ({j11.get('error_message')})")
    check(fake.one("tender_packs", id="pack-co2-shared").get("document_requirements") is None, "co2 got nothing from co1's folder")

    shutil.rmtree(work, ignore_errors=True)
    fake.httpd.shutdown()
    print(f"\n{'ALL PASSED' if not failures else f'{len(failures)} FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
