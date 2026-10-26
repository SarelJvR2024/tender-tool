"""
"What does THIS non-SANRAL tender require?" - reads a tender's own document
set (whatever the buyer published: one bundled PDF, a folder of SBD/MBD .docx
forms, a zip of buyer-specific templates) and returns a requirement manifest
the Tender Workspace renders and the pack generator follows.

Why this exists (docs/specs/General-Tender-Engine-Findings.md, plus the round
3 corpus in docs/specs/General-Tender-Preparation-Findings.md): non-SANRAL
buyers do NOT share one returnable list.
  - National/provincial departments, SOEs and smaller entities use National
    Treasury's SBD suite (SBD 1/4/6.1, sometimes 3.x/5/7/8/9).
  - Municipalities use the MFMA twin, the MBD suite (MBD 1/4/6.1/7.1/8/9) -
    same purpose and field set, different numbering prefix, usually inside
    one bundled PDF (Dawid Kruiper, Overberg, Cape Winelands, uMlalazi, BCMDA).
  - ESKOM runs its own forms (Form A, Declaration of Shareholding, Supplier
    Declaration of Interest, NEC3 contract) - no SBD/MBD form at all.
  - Each buyer's own wording decides which supporting documents, roles,
    experience evidence and locality rows apply.
So the manifest is read from the documents, never assumed from the buyer type.

Re-used, not re-implemented:
  - experience evidence + locality: engine/discovery/score_outlook.py
    extract_experience_evidence() / extract_locality() (built buyer-agnostic,
    confirmed on 12 non-SANRAL rules + 6 locality tables).
  - preference system / functionality / specific-goal bases:
    score_outlook.estimate_score_outlook().
  - CIDB grading / compulsory briefing: assess_tender.CIDB_RE / BRIEFING_RE.
  - role codes: the 17 non-SANRAL rows of migration 0016's tender_roles.

Every item carries a verbatim `quote` from the document it came from, so the
UI can show WHY something is listed. Nothing is added that the documents don't
mention. A form or document this engine can't fill is listed as "by hand" with
the reason - never silently dropped and never faked.

Usage (dry run, prints the manifest):
    python3 engine/general/requirements.py <file-or-dir> [<file-or-dir> ...] [--json out.json]
"""
import argparse
import datetime
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "discovery"))
if HERE not in sys.path:
    sys.path.insert(0, HERE)  # fill_flat_pdf (overlay check)
sys.path.insert(0, os.path.join(HERE, "..", "ocr"))
import scan_ocr  # noqa: E402  (EXPERIMENTAL - docs/specs/OCR-Evaluation.md)
from score_outlook import estimate_score_outlook, extract_experience_evidence, extract_locality  # noqa: E402
from key_person_minimums import extract_key_person_minimums  # noqa: E402
from assess_tender import CIDB_RE, BRIEFING_RE  # noqa: E402

MANIFEST_VERSION = 1

# ---------------------------------------------------------------------------
# Text extraction for every format seen in the real corpus
# (pdf 476 / docx 35 / zip 25 / xlsx 7 / doc 5 across 300 live releases,
# 2026-10-16 - see the findings doc)
# ---------------------------------------------------------------------------


def docx_text(path):
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            seen, cells = set(), []
            for c in row.cells:
                if id(c._tc) in seen:
                    continue
                seen.add(id(c._tc))
                cells.append(c.text.strip())
            parts.append(" | ".join(cells))
    return "\n".join(parts)


def convert_with_soffice(path, target_ext, out_dir):
    """LibreOffice headless conversion (.doc -> .docx, .docx -> .pdf). Returns
    the converted path or None. Same soffice dependency engine/cv/generate_cv.py
    and engine/print_reports.py already have."""
    try:
        subprocess.run(["soffice", "--headless", "--convert-to", target_ext, "--outdir", out_dir, path],
                       capture_output=True, timeout=180)
    except (subprocess.TimeoutExpired, FileNotFoundError):
        return None
    out = os.path.join(out_dir, os.path.splitext(os.path.basename(path))[0] + "." + target_ext.split(":")[0])
    return out if os.path.exists(out) else None


def pdf_text(path):
    return pdf_text_or_error(path)[0]


def pdf_text_or_error(path):
    """(text, error). A file pdftotext can't open at all (truncated download,
    an HTML error page saved as .pdf, a password-protected PDF) is a different
    problem from a scanned PDF with no text layer, and the user fixes it
    differently (re-download vs. ask for a text copy / OCR) - so the two are
    never reported with the same words. Before 2026-10-17 both came out as
    "no text layer (scanned image)"."""
    try:
        out = subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True, timeout=120)
    except subprocess.TimeoutExpired:
        return "", "pdftotext timed out reading it (very large or damaged PDF)"
    except FileNotFoundError:
        return "", "pdftotext (poppler-utils) is not installed on this worker"
    if out.returncode != 0:
        first = next((ln.strip() for ln in (out.stderr or "").splitlines() if ln.strip()), "")
        return "", "not a readable PDF - corrupt, truncated or not really a PDF" + (f" ({first[:120]})" if first else "") + " - re-download it"
    return out.stdout, None


def collect_documents(paths, work_dir):
    """Expands dirs and zips into a flat list of {name, path, format}. Zips are
    unpacked into work_dir (ESKOM publishes everything as zips)."""
    found = []

    def add(path, display):
        ext = os.path.splitext(path)[1].lower().lstrip(".")
        if ext == "zip":
            dest = os.path.join(work_dir, re.sub(r"[^A-Za-z0-9_.-]", "_", os.path.basename(path)) + "_unz")
            try:
                with zipfile.ZipFile(path) as z:
                    for member in z.namelist():
                        # zip-slip guard: never extract outside dest
                        target = os.path.realpath(os.path.join(dest, member))
                        if not target.startswith(os.path.realpath(dest) + os.sep):
                            continue
                        if member.endswith("/"):
                            continue
                        os.makedirs(os.path.dirname(target), exist_ok=True)
                        with z.open(member) as src, open(target, "wb") as dst:
                            shutil.copyfileobj(src, dst)
                        add(target, f"{display} > {member}")
            except zipfile.BadZipFile:
                found.append({"name": display, "path": path, "format": "zip", "unreadable": "corrupt zip"})
            return
        found.append({"name": display, "path": path, "format": ext or "unknown"})

    for p in paths:
        if os.path.isdir(p):
            for root, _, files in os.walk(p):
                for f in sorted(files):
                    if f.startswith(".") or f == "meta.json":
                        continue
                    full = os.path.join(root, f)
                    add(full, os.path.relpath(full, p))
        else:
            add(p, os.path.basename(p))
    return found


def read_document(doc, work_dir):
    """Fills doc['text'] and, for a legacy .doc, doc['converted_docx'] (the
    LibreOffice .docx conversion the filler can work on)."""
    fmt = doc["format"]
    text = ""
    if fmt == "pdf":
        text, err = pdf_text_or_error(doc["path"])
        if err:
            doc["unreadable"] = err
        else:
            try:
                text = ocr_scanned_pages(doc, text, work_dir)
            except Exception as e:  # experimental step: never worse than no OCR
                scan_ocr.system_alerts.record("ocr", f"OCR crashed on a tender document: {type(e).__name__}: {e}",
                                              "error", exc=e, context={"file": os.path.basename(doc["path"])})
                doc.pop("ocr", None)
                doc["_ocr_failure"] = f"OCR failed on it ({type(e).__name__})"
    elif fmt == "docx":
        try:
            text = docx_text(doc["path"])
        except Exception as e:  # corrupt / not really a docx
            doc["unreadable"] = f"python-docx could not open it: {e}"
    elif fmt == "doc":
        conv = convert_with_soffice(doc["path"], "docx", work_dir)
        if conv:
            doc["converted_docx"] = conv
            text = docx_text(conv)
        else:
            doc["unreadable"] = "LibreOffice could not convert this legacy .doc"
    elif fmt == "txt":
        text = open(doc["path"], encoding="utf-8", errors="replace").read()
    doc["text"] = text or ""
    if not doc.get("unreadable") and fmt in ("pdf", "docx", "doc") and len(doc["text"].strip()) < 50:
        tried = doc.pop("_ocr_failure", None)
        if doc.pop("ocr", None):
            tried = "OCR read only a few words from it"
        doc["unreadable"] = ("no text layer (scanned image)"
                             + (f", and {tried}" if tried else "") + " - read it by eye")
    doc.pop("_ocr_failure", None)
    return doc


# A page with fewer non-space characters than this in its text layer is
# offered to OCR. Not 0: a scan with a few typed-in fields on top has a tiny
# text layer (Tourism NB143's SBD 1 page: just "NB143/2026 20 OCT 2026 10H00"
# over a scanned form). The OCR text replaces the page's own only when it is
# clearly more (OCR_REPLACE_FACTOR x the text layer + 50 characters), so a
# short real page - a divider, "this page intentionally left blank" - keeps
# its text layer and isn't flagged as OCR.
OCR_PAGE_MIN_CHARS = 200
OCR_REPLACE_FACTOR = 2


def ocr_scanned_pages(doc, text, work_dir):
    """EXPERIMENTAL (engine/ocr/scan_ocr.py, docs/specs/OCR-Evaluation.md).
    Pages of a PDF with no text layer are OCR'd; each page whose OCR passes
    the quality gate gets the OCR text in its place (pages stay separated by
    form feeds, so page numbers still line up), and doc['ocr'] records which
    pages came from OCR so the manifest and the UI can say so. When OCR is
    unavailable or reads nothing usable the text is returned unchanged and
    read_document() refuses the document exactly as before OCR existed."""
    pages = text.split("\f")
    if pages and not pages[-1].strip():
        pages = pages[:-1]
    empty = [i + 1 for i, t in enumerate(pages) if len(re.sub(r"\s", "", t)) < OCR_PAGE_MIN_CHARS]
    if empty:
        with_images = scan_ocr.pages_with_images(doc["path"])
        if with_images is not None:
            empty = [n for n in empty if n in with_images]
    if not empty:
        return text
    fully_scanned = len(re.sub(r"\s", "", text)) < 50
    res = scan_ocr.ocr_pdf_pages(doc["path"], empty, work_dir=work_dir)
    used = []
    for n in res["read"]:
        native = len(re.sub(r"\s", "", pages[n - 1]))
        if len(re.sub(r"\s", "", res["pages"][n]["text"])) >= OCR_REPLACE_FACTOR * native + 50:
            pages[n - 1] = res["pages"][n]["text"]
            used.append(n)
    short = bool(res["read"]) and not used
    res["read"] = used
    if not used:
        if fully_scanned:
            doc["_ocr_failure"] = "OCR read only a few words from it" if short else scan_ocr.failure_sentence(res)
        return text
    summary = scan_ocr.summary(res, total_pages=len(pages))
    if not fully_scanned:
        # A blank page in a text PDF isn't news; only an unclear scan is.
        summary["pages_failed"] = [f for f in summary["pages_failed"] if "unclear" in f["reason"]]
    doc["ocr"] = summary
    return "\f".join(pages) + "\f"


# ---------------------------------------------------------------------------
# Standard bidding forms (SBD = National Treasury, MBD = municipal twin)
# ---------------------------------------------------------------------------

# code -> (title, what the engine can do with a .docx edition). Titles are the
# real form headings seen in the corpus.
FORMS = {
    "1": ("Invitation to bid", "fill"),
    "2": ("Tax clearance / tax compliance status", "attach"),
    "3": ("Pricing schedule", "by_hand"),  # price is tender-specific - never auto-filled
    "4": ("Bidder's disclosure / declaration of interest", "fill"),
    "5": ("National Industrial Participation Programme", "by_hand"),
    "6.1": ("Preference points claim (PPR 2022)", "fill"),
    "6.2": ("Declaration certificate for local production and content", "by_hand"),
    "7": ("Contract form", "by_hand"),
    "8": ("Declaration of bidder's past supply chain management practices", "fill"),
    "9": ("Certificate of independent bid determination", "fill"),
}

FORM_RE = re.compile(r"\b([SM])\s?B\s?D\s*[-–]?\s*(\d{1,2})(?:\s*[.(]\s*(\d)\s*\)?)?(?:\s*\(\s*[a-z]\s*\))?", re.IGNORECASE)
# "MDB 7.1" typo seen verbatim in Dawid Kruiper TN051 3.4
FORM_TYPO_RE = re.compile(r"\bMDB\s*(\d)(?:\.(\d))?", re.IGNORECASE)


def _form_code(major, minor):
    major = str(int(major))
    if major == "6" and minor in ("1", "2"):
        return f"6.{minor}"
    if major == "3":
        return "3"
    if major == "7":
        return "7"
    return major if major in FORMS else None


def forms_in_text(text):
    """{(suite, code): count} for every SBD/MBD reference in a text."""
    counts = {}
    for m in FORM_RE.finditer(text):
        code = _form_code(m.group(2), m.group(3))
        if code:
            key = ("SBD" if m.group(1).upper() == "S" else "MBD", code)
            counts[key] = counts.get(key, 0) + 1
    for m in FORM_TYPO_RE.finditer(text):
        code = _form_code(m.group(1), m.group(2))
        if code:
            counts[("MBD", code)] = counts.get(("MBD", code), 0) + 1
    return counts


# Field labels that make a document an EDITION of the form (not just a
# mention). Taken from the real Treasury / RAF / DSTI / SSA / SITA / Mthashana
# / Tourism editions (findings doc, round 3).
EDITION_MARKERS = {
    # SBD 4 / 6.1 editions also say "NAME OF BIDDER" / "invitation to bid", so
    # SBD 1 needs its own bidder-detail block, 3 of these 4.
    "1": ([r"VAT\s+REGISTRATION\s+NUMBER", r"TCS\s*PIN|CSD\s*No|SUPPLIER\s+COMPLIANCE\s+STATUS|TAX\s+COMPLIANCE\s+SYSTEM\s+PIN",
           r"ACCREDITED\s+REPRESENTATIVE|FOREIGN\s+BASED\s+SUPPLIER", r"POSTAL\s+ADDRESS"], 3),
    "4": ([r"employed\s+by\s+(?:the\s+state|an\s+organ\s+of\s+state)|EMPLOYMENT\s+BY\s+(?:AN\s+)?ORGAN\s+OF\s+STATE",
           r"relationship\s+with\s+any\s+person", r"interest\s+in\s+any\s+other\s+related",
           # MBD 4, the older municipal "Declaration of Interest" layout (Dawid Kruiper TN051 p.30)
           r"DECLARATION\s+OF\s+INTEREST", r"in\s+the\s+service\s+of\s+the\s+state"], 2),
    "6.1": ([r"Name\s+of\s+company/firm", r"Company\s+registration\s+number", r"POINTS\s+CLAIMED|points\s+claimed"], 2),
    # The CSD supplier leaflet also names both registers - SBD 8 needs its heading.
    "8": ([r"PAST\s+SUPPLY\s+CHAIN\s+MANAGEMENT\s+PRACTICES", r"Database\s+of\s+Restricted\s+Suppliers|Register\s+for\s+Tender\s+Defaulters"], 2),
    # Heading + the Competition Act wording every SBD/MBD 9 carries: a contract
    # form that merely LISTS "Certificate of Independent Bid Determination"
    # among its annexures (Ratlou LM's MBD 7.2, live OCDS 2026-10-17) is not one.
    "9": ([r"INDEPENDENT\s+BID\s+DETERMINATION", r"bid\s+rigging|collusive\s+bid|Competition\s+Act"], 2),
}


def edition_forms(text):
    out = []
    for code, (pats, need) in EDITION_MARKERS.items():
        if sum(1 for p in pats if re.search(p, text, re.IGNORECASE)) >= need:
            out.append(code)
    return out


# ---------------------------------------------------------------------------
# Supporting documents ("returnables"). Each key was seen in the round-3
# corpus (counts per tender in the findings doc); vault = the companies column
# that already answers it (migrations 0001/0014), upload = the user attaches a
# copy for this tender when no vault field exists.
# ---------------------------------------------------------------------------

RETURNABLES = [
    {"key": "csd_report", "label": "Central Supplier Database (CSD) registration / summary report",
     "re": r"central\s+supplier\s+database|\bCSD\s+(?:registration|report|summary|number|supplier)", "vault": "csd"},
    {"key": "tax_compliance", "label": "Tax compliance status (SARS TCS PIN)",
     "re": r"tax\s+compliance\s+status|TCS\s*PIN|tax\s+clearance\s+certificate", "vault": "tax"},
    {"key": "bbbee_proof", "label": "B-BBEE certificate or sworn affidavit",
     "re": r"B-?BBEE\s+(?:status\s+level\s+)?(?:verification\s+)?certificate|sworn\s+affidavit", "vault": "bbbee"},
    {"key": "coid", "label": "COID letter of good standing",
     "re": r"\bCOID\b|compensation\s+(?:for\s+)?occupational\s+injuries|letter\s+of\s+good\s+standing", "vault": "coid"},
    {"key": "cidb", "label": "CIDB registration / grading",
     "re": r"\bCIDB\b", "vault": "cidb"},
    {"key": "cipc_documents", "label": "CIPC company registration documents",
     "re": r"\bCIPC\b|company\s+registration\s+(?:documents|certificate)", "vault": None},
    {"key": "director_ids", "label": "Certified ID copies of directors / members",
     "re": r"(?:certified\s+)?cop(?:y|ies)\s+of\s+(?:the\s+)?(?:directors'?\s+|members'?\s+)?(?:ID|identity)", "vault": None},
    {"key": "municipal_account", "label": "Municipal rates account / proof not in arrears",
     "re": r"municipal\s+(?:rates|account|services)|rates\s+(?:and\s+taxes|clearance)|not\s+in\s+arrears", "vault": None},
    {"key": "bank_letter", "label": "Bank confirmation letter / bank rating",
     "re": r"bank\s+(?:confirmation\s+)?letter|bank\s+rating|letter\s+from\s+(?:the\s+|your\s+)?bank", "vault": "bank"},
    {"key": "public_liability", "label": "Public liability insurance",
     "re": r"public\s+liability", "vault": None},
    {"key": "pi_insurance", "label": "Professional indemnity insurance",
     "re": r"professional\s+indemnity", "vault": "insurance"},
    {"key": "psira", "label": "PSIRA registration certificate",
     "re": r"\bPSIRA\b", "vault": None},
    {"key": "financial_statements", "label": "Audited / annual financial statements",
     "re": r"audited\s+(?:annual\s+)?financial\s+statements|annual\s+financial\s+statements", "vault": None},
    {"key": "she_plan", "label": "Health & safety (SHE) plan / file",
     "re": r"health\s+and\s+safety\s+(?:plan|file)|SHE\s+(?:plan|file)|safety\s+plan", "vault": None},
    {"key": "jv_agreement", "label": "Joint venture / consortium agreement (if bidding as a JV)",
     "re": r"joint\s+venture\s+agreement|consortium\s+agreement", "vault": None, "conditional": "jv"},
]

REQUIRED_CONTEXT_RE = re.compile(
    r"\bmust\b|\bshall\b|attach|submit|compulsory|mandatory|non-?\s*responsive|disqualif|requir|proof\s+of|provide",
    re.IGNORECASE)


# ---------------------------------------------------------------------------
# Per-item confidence (Sarel, 2026-10-20: "indicate the confidence level ...
# train to improve the confidence"). A pure function so it's unit-testable
# without running the whole extractor (test_requirements.py).
#
# Score = a base per HOW the item was recognised, nudged by how many
# independent mentions agree, then pulled down if the text behind it came
# from OCR (scan_ocr.py's own mean_confidence - never re-derived, just
# folded in). Every signal already exists elsewhere in this module; nothing
# here is a fabricated number.
#
#   form_edition    - a real document FILE was matched to a form code (the
#                      strongest signal: the form itself exists, not just a
#                      mention of its name).
#   form_code       - an anchored SBD/MBD form-number regex matched (FORM_RE/
#                      FORM_TYPO_RE) but no edition file was found.
#   strong_context  - a returnable/role whose surrounding text matched an
#                      explicit requirement verb (REQUIRED_CONTEXT_RE /
#                      REQUIREMENT_VERB_RE: "must", "shall", "attach"...).
#   weak_context    - same pattern matched, but only the looser topic regex
#                      (ROLE_CONTEXT_RE) - plausible, not explicit.
#   anchored_regex  - a narrow, specific pattern with its own capture group
#                      (CIDB_RE / BRIEFING_RE), not a bag of synonyms.
CONFIDENCE_BASE = {
    "form_edition": 92,
    "form_code": 72,
    "strong_context": 82,
    "weak_context": 58,
    "anchored_regex": 85,
}

CONFIDENCE_LABEL = {
    "form_edition": "a document file matching this form was found",
    "form_code": "the form's number was named in the text",
    "strong_context": "matched next to an explicit requirement (\"must\"/\"shall\"/\"attach\")",
    "weak_context": "matched a plausible context, but no explicit requirement wording nearby",
    "anchored_regex": "matched a specific, narrow pattern",
}


def score_confidence(kind, mentions=1, ocr_confidence=None):
    """Returns (score 0-100, reason). kind is one of CONFIDENCE_BASE's keys.
    mentions = how many independent, consistent hits backed this item (more
    agreement = higher confidence; a single hit is neither boosted nor
    penalised - every base score above already assumes exactly one).
    ocr_confidence = scan_ocr.py's mean_confidence (0-100) when any of the
    text behind this item was read by OCR, else None - OCR can misread
    characters, so it always pulls the score down, scaled by how good the
    OCR read itself was (a near-perfect OCR barely moves it; a poor one pulls
    hard)."""
    score = CONFIDENCE_BASE[kind]
    reason = CONFIDENCE_LABEL[kind]
    if mentions >= 5:
        score += 8
        reason += f"; {mentions} consistent mentions"
    elif mentions >= 2:
        score += 4
        reason += f"; {mentions} consistent mentions"
    score = min(score, 100)
    if ocr_confidence is not None:
        factor = 0.5 + 0.5 * (ocr_confidence / 100)
        score = round(score * factor)
        reason += f"; read from an OCR'd scan at {ocr_confidence:.0f}/100 OCR confidence, reduced"
    return max(0, min(100, int(round(score)))), reason


# ---------------------------------------------------------------------------
# "Train to improve the confidence" (Sarel). Honesty check first: this
# module has NO LLM call anywhere in it - every extraction above is regex/
# heuristic. So there is no prompt to fold few-shot examples into. What a
# handful of past admin corrections (requirement_extraction_reviews,
# migration 0074) CAN honestly do here: when a new item looks like one a
# human already flagged as wrong (same item type + same stable identifying
# field - a form code, a returnable key, a role code... - the SOURCE TEXT
# will differ tender to tender, so matching on the quote itself would almost
# never hit), knock its confidence down and say why, so the UI shows it as
# less trustworthy instead of silently repeating the same mistake with full
# confidence. It never changes WHAT gets extracted - no regex adapts itself
# from a correction - only how confidently that one item is reported this
# time. `past_corrections` is passed in by the caller (process_general_jobs.py
# reads up to 5 recent rows from requirement_extraction_reviews - "most
# recent N corrections globally", the cheap heuristic the task scope allows
# when nothing smarter - buyer/form matching - is available); this stays a
# plain function of its input so it's unit-testable without any database.
def _item_identity(item_type, item):
    """The one field that identifies "the same kind of item" across
    different tenders/documents (never the quote - that's tender-specific
    free text and would never recur)."""
    return {
        "forms": item.get("code"), "buyer_forms": item.get("key"),
        "returnables": item.get("key"), "roles": item.get("role_code"),
        "person_registrations": item.get("body"),
        "cidb": "cidb", "briefing": "briefing",
    }.get(item_type)


def _apply_past_corrections(item_type, item, past_corrections):
    """Mutates item['confidence']/['confidence_reason'] in place when a past
    correction matches this item's identity (see module comment above).
    Up to 5 corrections are expected in past_corrections; only the first
    match (most recent) is applied."""
    if not past_corrections or item.get("confidence") is None:
        return
    ident = _item_identity(item_type, item)
    if ident is None:
        return
    for c in past_corrections:
        if c.get("item_type") != item_type:
            continue
        if _item_identity(item_type, c.get("item_snapshot") or {}) != ident:
            continue
        if c.get("verdict") == "false_positive":
            item["confidence"] = max(5, item["confidence"] - 25)
            item["confidence_reason"] += ("; a platform admin previously marked a similar "
                                          f"{item_type} extraction a false positive - treat with extra caution")
        elif c.get("verdict") == "corrected":
            item["confidence"] = max(5, item["confidence"] - 12)
            item["confidence_reason"] += ("; a platform admin previously corrected a similar "
                                          f"{item_type} extraction - double-check this one")
        break


def _apply_past_corrections_to_manifest(manifest, past_corrections):
    if not past_corrections:
        return
    for item_type in ("forms", "buyer_forms", "returnables", "roles", "person_registrations"):
        for item in manifest.get(item_type) or []:
            _apply_past_corrections(item_type, item, past_corrections)
    for item_type in ("cidb", "briefing"):
        item = manifest.get(item_type)
        if item:
            _apply_past_corrections(item_type, item, past_corrections)


def _ocr_mean(docs):
    """{document name -> scan_ocr.py's mean_confidence} for every document
    some of whose text came from OCR (see ocr_scanned_pages())."""
    return {d["name"]: d["ocr"]["mean_confidence"] for d in docs if d.get("ocr") and d["ocr"].get("mean_confidence") is not None}


def _quote(text, start, end, width=110):
    a = max(0, start - min(width, 70))
    b = min(len(text), end + width + 40)
    s = re.sub(r"\s+", " ", text[a:b]).strip()
    return s[:240]


def find_returnables(docs, ocr_mean=None):
    ocr_mean = ocr_mean or {}
    out = []
    for r in RETURNABLES:
        rx = re.compile(r["re"], re.IGNORECASE)
        best = None
        mentions = 0
        for d in docs:
            t = d.get("text") or ""
            for m in rx.finditer(t):
                mentions += 1
                ctx = t[max(0, m.start() - 200): m.end() + 200]
                strong = bool(REQUIRED_CONTEXT_RE.search(ctx))
                if best is None or (strong and not best[0]):
                    best = (strong, d["name"], _quote(t, m.start(), m.end()))
                if best[0]:
                    break
            if best and best[0]:
                break
        if best:
            confidence, confidence_reason = score_confidence(
                "strong_context" if best[0] else "weak_context", mentions, ocr_mean.get(best[1]))
            out.append({
                "key": r["key"], "label": r["label"], "vault": r["vault"],
                "conditional": r.get("conditional"),
                "strength": "required" if best[0] else "mentioned",
                "mentions": mentions, "document": best[1], "quote": best[2],
                "confidence": confidence, "confidence_reason": confidence_reason,
            })
    return out


# ---------------------------------------------------------------------------
# Key-person roles (migration 0016's non-SANRAL tender_roles rows). Each
# regex was checked against the tender its role row cites.
# ---------------------------------------------------------------------------

ROLE_PATTERNS = [
    ("contractor.contract_manager", r"\bcontracts?\s+manager\b"),
    ("contractor.site_agent", r"\bsite\s+agent\b"),
    ("contractor.foreman", r"\bforem[ae]n\b"),
    # NEC3's "Employer's / Client's Project Manager" is the buyer's own agent, not a bidder role.
    ("contractor.project_manager", r"(?<!employer's\s)(?<!employer’s\s)(?<!client's\s)\bproject\s+manager\b"),
    ("hs.safety_officer", r"(?<!purchaser's\s)(?<!purchaser’s\s)(?<!employer's\s)\bsafety\s+officer\b"),
    ("psp.professional_engineer", r"\bprofessional\s+engineer\b(?!ing)|\bPr\.?\s?Eng\b"),
    ("psp.professional_architect", r"\bprofessional\s+architect\b"),
    ("psp.architectural_technologist", r"\barchitectural\s+technologist\b"),
    ("psp.quantity_surveyor", r"\bquantity\s+surveyor\b"),
    ("psp.town_planner", r"\btown\s+(?:and\s+regional\s+)?planner\b"),
    ("psp.property_valuer", r"\b(?:property\s+)?valuer\b|\bSACPVP\b"),
    ("trades.mechanic", r"\bmechanic\s+trade\s+test|\bmechanical\s+engineering\s+diploma"),
    ("trades.auto_electrician", r"\bauto\s+electrician\b|\bmillwright\b"),
    ("trades.carpenter", r"\bcarpenter\b"),
    ("services.pest_control_operator", r"(?<!assistant\s)\bpest\s+control\s+operator\b"),
    ("services.assistant_pest_control_operator", r"\bassistant\s+pest\s+control\s+operator\b"),
    ("services.security_officer", r"\bsecurity\s+officers?\b|\bPSIRA[- ]registered\b"),
]


ROLE_CONTEXT_RE = re.compile(
    r"key\s+(?:personnel|staff)|\bCVs?\b|curriculum|experience|qualifi|registered|registration|certificate|"
    r"years|\bmust\b|\bshall\b|required|appoint|organogram|team",
    re.IGNORECASE)
# "Pr.CPM   Professional Construction Project Manager" - an abbreviations
# list (Treasury RT83's glossary) names roles nobody is asked to field.
GLOSSARY_LINE_RE = re.compile(r"^\s*[A-Z][A-Za-z.()/&]{0,14}\s{2,}[A-Z]")


def _line_at(text, pos):
    a = text.rfind("\n", 0, pos) + 1
    b = text.find("\n", pos)
    return text[a: b if b != -1 else len(text)]


# Finding 47: "Contract Manager"/"Project Manager" preceded by the TENDER'S
# OWN BUYER NAME ("the Transnet Contract Manager") is the buyer's own named
# administrator, not a bidder role - confirmed on a real Transnet SOC Ltd
# tender (both real mentions read "...for review and acceptance by the
# Transnet Contract Manager and SHEQ Department", not a bidder ask). This is
# the same bug class NEC3's "Employer's/Client's Project Manager" lookbehind
# already excludes (ROLE_PATTERNS above), just with the buyer's own NAME
# instead of a generic legal title - a blanket ROLE_CONTEXT_RE change was
# tried and rejected first (deleted 3 other genuine detections, see
# QA-Edge-Case-Findings.md Finding 47), so this is deliberately narrow:
# only the tender's own buyer name, only directly before the match, never a
# change to the shared context-word check every other role pattern depends
# on.
_BUYER_STOPWORDS = {
    "the", "of", "and", "soc", "ltd", "pty", "limited", "department", "dept",
    "municipality", "metropolitan", "local", "district", "council", "board",
    "agency", "authority", "corporation", "company", "services", "group",
    "national", "provincial", "government", "republic", "south", "africa",
    "province", "city", "state", "owned", "holdings", "incorporated", "inc",
}


def _buyer_core_words(buyer):
    """The buyer name's own distinctive words ("Transnet" from "Transnet SOC
    Ltd", "ESKOM" from "ESKOM Holdings SOC Ltd") - generic legal/government
    suffixes stripped so they never themselves count as a "buyer name"
    match (a document saying plain "the Department Project Manager" isn't
    naming ANY specific buyer). Words under 4 characters are dropped too -
    too short to be a reliable, specific signal on their own."""
    if not buyer:
        return []
    words = re.findall(r"[A-Za-z][A-Za-z'-]*", buyer)
    return sorted({w for w in words if len(w) >= 4 and w.lower() not in _BUYER_STOPWORDS}, key=len, reverse=True)


def _is_buyers_own_person(text, match_start, buyer_words):
    """Is the tender's own buyer named directly before this match
    ("the Transnet Contract Manager")? A short, immediately-preceding
    window only (not the whole ±250-char context window) - this is about
    what qualifies THIS specific mention, not the document's general
    subject matter."""
    if not buyer_words:
        return False
    window = text[max(0, match_start - 60): match_start]
    return any(re.search(r"\b" + re.escape(w) + r"\b", window, re.IGNORECASE) for w in buyer_words)


def find_roles(docs, ocr_mean=None, buyer=None):
    ocr_mean = ocr_mean or {}
    buyer_words = _buyer_core_words(buyer)
    out = []
    for code, pat in ROLE_PATTERNS:
        rx = re.compile(pat, re.IGNORECASE)
        hit = None
        n = 0
        for d in docs:
            t = d.get("text") or ""
            for m in rx.finditer(t):
                if GLOSSARY_LINE_RE.match(_line_at(t, m.start())):
                    continue
                if not ROLE_CONTEXT_RE.search(t[max(0, m.start() - 250): m.end() + 250]):
                    continue
                if _is_buyers_own_person(t, m.start(), buyer_words):
                    continue
                n += 1
                if hit is None:
                    hit = (d["name"], _quote(t, m.start(), m.end()))
        if hit:
            confidence, confidence_reason = score_confidence("weak_context", n, ocr_mean.get(hit[0]))
            out.append({"role_code": code, "mentions": n, "document": hit[0], "quote": hit[1],
                        "confidence": confidence, "confidence_reason": confidence_reason})
    return out


# ---------------------------------------------------------------------------
# Per-person professional registrations a tender asks for that are NOT on
# every person's record (migration 0040, audit Q5): SACNASP (Pr Sci.Nat - e.g.
# an engineering geologist, environmental scientist) and SAGC (land /
# engineering surveyors; PLATO is the council's former name); and (held via the 0050 registrations catalog,
# audit O4) SACAP, SACPLAN, SACQSP, SACPVP and PSIRA. Sarel: "don't
# specifically ask for it unless the tender requires it" - so this list is
# what makes the People page show those fields and the workspace name who
# holds them. Same context rule as the roles: a glossary line or a passing
# mention doesn't count.
# ---------------------------------------------------------------------------

PERSON_REGISTRATION_PATTERNS = [
    ("SACNASP", r"\bSACNASP\b|\bPr\.?\s?Sci\.?\s?Nat\b|natural\s+scientific\s+professions|professional\s+natural\s+scientist"),
    ("SAGC", r"\bSAGC\b|geomatics\s+council|\bPLATO\b|professional\s+land\s+surveyor"),
    # 0050 catalog bodies (audit O4) - the other statutory councils, by acronym, council
    # name or registered title. Read in the 123-tender corpus: ACSA COR8138
    # (SACAP / SACQSP), DLRRD NC 001 and TDR11 (SACPLAN), Maluti TVET RFQ168
    # (SACPVP), and six guarding / alarm-response tenders (PSIRA).
    # "Landscape Architectural Profession" is SACLAP, not SACAP.
    ("SACAP", r"\bSACAP\b|council\s+for\s+the\s+architectural\s+profession|(?<!landscape\s)architectural\s+profession\s+act|"
              r"\bPr\.?\s?Arch\b|professional\s+(?:senior\s+)?architectural\s+(?:technologist|draughtsperson)|professional\s+architect\b"),
    ("SACPLAN", r"\bSACPLAN\b|council\s+for\s+planners|planning\s+professions?\s+act|professional\s+planner\b|\bPr\.?\s?Pln\b"),
    ("SACQSP", r"\bSACQSP\b|council\s+for\s+the\s+quantity\s+surveying\s+profession|quantity\s+surveying\s+profession\s+act|"
               r"professional\s+quantity\s+surveyor|\bPr\.?\s?QS\b"),
    ("SACPVP", r"\bSACPVP\b|council\s+for\s+the\s+property\s+valuers\s+profession|property\s+valuers\s+profession\s+act|"
               r"professional\s+(?:associated\s+)?valuer\b"),
    ("PSIRA", r"\bPSIRA\b|private\s+security\s+industry\s+regulat"),
]
# A council NAMED, not asked for: a definitions clause or a statute table
# ("statutory council: a) South African Council for the Architectural
# Profession, established by the Architectural Profession Act of 2000 ...",
# the CIDB standard's "Categories of registration ... Act" table, read in
# CDC44126's gazette annexure) lists every council whether or not the tender
# needs any of them. Such a mention is skipped unless a requirement verb sits
# right next to it.
DEFINITION_CONTEXT_RE = re.compile(r"established\s+by|\(Act\s+No|Act\s+No\.?\s*\d|Act\s+of\s+(?:19|20)\d\d|categor(?:y|ies)\s+of\s+registration",
                                   re.IGNORECASE)
REQUIREMENT_VERB_RE = re.compile(r"\bmust\b|\bshall\b|required|registered\s+with|proof\s+of\s+registration|certificate|be\s+registered",
                                 re.IGNORECASE)


def find_person_registrations(docs, ocr_mean=None):
    ocr_mean = ocr_mean or {}
    out = []
    for body, pat in PERSON_REGISTRATION_PATTERNS:
        rx = re.compile(pat, re.IGNORECASE)
        hit, n, strong = None, 0, False
        for d in docs:
            t = d.get("text") or ""
            for m in rx.finditer(t):
                if GLOSSARY_LINE_RE.match(_line_at(t, m.start())):
                    continue
                if not ROLE_CONTEXT_RE.search(t[max(0, m.start() - 250): m.end() + 250]):
                    continue
                near = t[max(0, m.start() - 160): m.end() + 160]
                if DEFINITION_CONTEXT_RE.search(near) and not REQUIREMENT_VERB_RE.search(near):
                    continue
                n += 1
                if REQUIREMENT_VERB_RE.search(near):
                    strong = True
                if hit is None:
                    hit = (d["name"], _quote(t, m.start(), m.end()))
        if hit:
            confidence, confidence_reason = score_confidence(
                "strong_context" if strong else "weak_context", n, ocr_mean.get(hit[0]))
            out.append({"body": body, "mentions": n, "document": hit[0], "quote": hit[1],
                        "confidence": confidence, "confidence_reason": confidence_reason})
    return out


# ---------------------------------------------------------------------------
# Buyer-specific forms: things a bidder must complete that are not an SBD/MBD
# form. The engine has no template for these - they are listed honestly as
# "complete by hand" (findings doc: ESKOM Form A / declarations / NEC3,
# municipal Form of Quotation, Sentech consent form, QCTO CV template).
# ---------------------------------------------------------------------------

BUYER_FORM_NAME_RE = re.compile(
    r"declaration|form\b|questionnaire|affidavit|consent|template|acknowledg|annexure|schedule|certificate|"
    r"returnable|method\s+statement|quality\s+(?:control|plan)|attendance",
    re.IGNORECASE)
REFERENCE_NAME_RE = re.compile(
    r"conditions\s+of\s+(?:contract|tender)|\bGCC\b|manual|leaflet|fact\s*sheet|read\s*me|help|exchange\s+rates|"
    r"instruction|specification|scope\s+of\s+work|terms\s+of\s+reference|\bToR\b|strategy|drawings?|fraud\s+notice|"
    r"contact\s+details|explanation",
    re.IGNORECASE)
IN_TEXT_BUYER_FORMS = [
    ("form_of_quotation", "Form of Quotation", r"FORM\s+OF\s+QUOTATION"),
    ("form_of_offer", "Form of Offer and Acceptance", r"FORM\s+OF\s+OFFER(?:\s+AND\s+ACCEPTANCE)?"),
    ("nec3_contract", "NEC3 contract (ESKOM / Transnet contract family)", r"\bNEC\s*3\b"),
    ("compulsory_enterprise_questionnaire", "Compulsory Enterprise Questionnaire", r"COMPULSORY\s+ENTERPRISE\s+QUESTIONNAIRE"),
    ("record_of_addenda", "Record of Addenda", r"RECORD\s+OF\s+ADDENDA"),
    ("certificate_of_authority", "Certificate / resolution of authority for signatory",
     r"(?:CERTIFICATE|RESOLUTION)\s+OF\s+AUTHORITY"),
    ("declaration_of_shareholding", "Declaration of Shareholding and Directorship", r"DECLARATION\s+OF\s+SHAREHOLDING"),
    # ESKOM's own Supplier Quality Management form - confirmed from a real
    # ESKOM tender (E3703CXMWP, 2026-10-21): "complete, sign and return Form A
    # with the other returnables", repeated across 6 documents, never
    # surfaced before this - requirements.py's own docstring already named
    # "Form A" as one of ESKOM's own forms, but nothing matched it. "Form A"
    # alone is too generic a string to match safely (any buyer could have an
    # unrelated "Form A") - anchored on the parenthetical that names what it
    # actually is, which recurs on every real mention in the corpus.
    ("eskom_form_a", "Form A (Supplier Quality Management)",
     r"FORM\s+A\s*\(\s*(?:ENQUIRY|TENDER)\s*(?:AND|/)?\s*CONTRACT\s*(?:AND|/)?\s*QUALITY\s+REQUIREMENTS\s+FOR\s+SUPPLIER\s+QUALITY\s+MANAGEMENT"),
]


# ---------------------------------------------------------------------------
# Finding 42 (QA-Edge-Case-Findings.md, 2026-10-21): some buyers (confirmed:
# City of Tshwane's "Notice N of <FY>" bulletins) publish ONE PDF covering
# SEVERAL unrelated tenders - a numbered table, one row per tender, each with
# its own reference, closing date, CIDB grading, preference system. Reading
# this as if it were one tender's own document attributes every other
# tender's facts to whichever one is being processed (confirmed: GPM
# 08-2026/27, a property lease with no CIDB angle at all, came back with a
# CIDB grading requirement that actually belongs to a different row, GPM 18,
# in the same PDF). Per Sarel's own call: detect this shape and decline to
# extract tender-specific facts from it at all (option (a) - safer than
# trying to isolate the right row, which nothing in this engine does today),
# rather than keep attributing another tender's facts to this one.
#
# Detection, buyer-agnostic (no reference-number format assumed): a numbered
# list item ("N. ") immediately followed by an uppercase reference-like token
# (letters then digits - "HD 03", "GPM 18", "WSBU 01"), with a month name
# within the next 200 characters (the row's own closing date, wrapped PDF
# text and all). Checked against every other real document already in this
# corpus (templates/general_samples/, plus 26 further real tenders sampled
# for Finding 40/42 themselves) - zero false positives at a threshold of 3;
# the one real bulletin matched 12 times, not a borderline case.
_BULLETIN_MONTHS = (r"January|February|March|April|May|June|July|August|September|October|"
                    r"November|December")
BULLETIN_ROW_RE = re.compile(
    r"(?:^|\n)\s*\d{1,2}\.\s+[A-Z]{2,8}\s?\d{1,4}[\s\S]{0,200}?\b(?:" + _BULLETIN_MONTHS + r")\b")
BULLETIN_MIN_ROWS = 3


def is_combined_bulletin(text):
    """True if `text` looks like several different tenders' own rows bundled
    into one document, not one tender's own document - see module comment
    above BULLETIN_ROW_RE for the evidence and why 3 is the threshold."""
    return len(BULLETIN_ROW_RE.findall(text or "")) >= BULLETIN_MIN_ROWS


def classify_document(doc, mentioned_forms):
    """kind: 'standard_form' (an SBD/MBD edition), 'buyer_form', 'tender_document'
    (the main bundle / RFQ), 'reference' (conditions, manuals, specs), 'unreadable'."""
    name = os.path.basename(doc["name"])
    text = doc.get("text") or ""
    if REFERENCE_NAME_RE.search(name) and not forms_in_text(name):
        doc["edition_of"] = []
        return "reference"
    eds = edition_forms(text)
    # A single-form file whose name AND own heading name the same form is an
    # edition of that form, whatever else its body mentions (Ratlou's
    # "MBD 7.2.pdf" was read as an MBD 9 edition and put into the pack as
    # "MBD 9" while the real "MBD 9.pdf" sat next to it).
    named = forms_in_text(os.path.basename(doc["name"]))
    if len(named) == 1:
        (_, code), = named.keys()
        if any(c == code for (_, c) in forms_in_text(text[:1500])):
            eds = [code]
    doc["edition_of"] = eds
    in_doc = forms_in_text(text) if text else {}
    name_forms = forms_in_text(name)
    if doc.get("unreadable") and not name_forms:
        return "unreadable"
    big = len(text) > 60000
    if eds and (name_forms or not big):
        return "standard_form"
    if eds and big:
        return "tender_document"  # a bundle that contains forms (Mthashana, SSA, Dawid Kruiper)
    if name_forms:
        return "standard_form"
    if REFERENCE_NAME_RE.search(name):
        return "reference"
    if BUYER_FORM_NAME_RE.search(name) and len(text) < 60000:
        return "buyer_form"
    return "tender_document"


# ---------------------------------------------------------------------------
# Questions only a human can answer for THIS tender. Each one is tied to a
# real field on a real form edition (quoted labels from the corpus), and is
# only asked when that form (or condition) is actually required.
# ---------------------------------------------------------------------------

YES_NO = ["no", "yes"]

QUESTIONS = [
    # --- SBD/MBD 1 -------------------------------------------------------
    {"key": "signatory_person_id", "form": "1", "group": "Signatory", "type": "person",
     "label": "Who signs this bid?",
     "help": "Fills SBD 1 'capacity under which this bid is signed', SBD 4 'Designation' and SBD 6.1 'Surname and name'. "
             "Their capacity comes from the person's record (People page)."},
    {"key": "total_bid_price", "form": "1", "group": "Price", "type": "text", "optional": True,
     "label": "Total bid price (incl. VAT)",
     "help": "SBD 1 'TOTAL BID PRICE'. Tender-specific - Tender Tool never computes it. Leave blank to write it by hand."},
    {"key": "sbd1.accredited_representative", "form": "1", "group": "SBD 1 questions", "type": "yes_no",
     "label": "Are you the accredited representative in South Africa for the goods/services/works offered?",
     "help": "SBD 1 Part B. If yes, the form says to enclose proof."},
    {"key": "sbd1.foreign_based", "form": "1", "group": "SBD 1 questions", "type": "yes_no", "derived_default": "no",
     "label": "Are you a foreign-based supplier for the goods/services/works offered?",
     "help": "SBD 1 Part B. Defaults to No for a CIPC-registered South African company."},
    # --- SBD/MBD 4 -------------------------------------------------------
    {"key": "sbd4.employed_by_state", "form": "4", "group": "Declarations (SBD 4)", "type": "yes_no",
     "label": "Is the bidder, or any director / shareholder / member / person with a controlling interest, employed by the state?",
     "help": "SBD 4 2.1 (Treasury: 'employed by the state'; RAF/SALB/QCTO: 'DECLARATION ON EMPLOYMENT BY ORGAN OF STATE')."},
    {"key": "sbd4.relationship_with_employee", "form": "4", "group": "Declarations (SBD 4)", "type": "yes_no",
     "label": "Does anyone connected with the bidder have a relationship with an employee of THIS buyer?",
     "help": "SBD 4 2.2 - specific to this buyer, so it's asked per tender."},
    {"key": "sbd4.interest_in_related_enterprise", "form": "4", "group": "Declarations (SBD 4)", "type": "yes_no",
     "label": "Does the bidder or any director / shareholder / person with a controlling interest have an interest in any other related enterprise, whether or not it is bidding for this contract?",
     "help": "SBD 4 'interest in any other related enterprise' (2.3 in older editions, 2.4 in Treasury's current one). "
             "The form asks about ANY related enterprise, not only ones bidding here. If yes, Treasury's current "
             "edition requires every CSD-registered company linked to the directors (MAAA number and status) - "
             "'failure to disclose ... will lead to disqualification'. That table is completed by hand."},
    # --- SBD/MBD 6.1 -----------------------------------------------------
    {"key": "sbd6_1.points_claimed", "form": "6.1", "group": "Preference points (SBD 6.1)", "type": "text", "optional": True,
     "label": "Specific-goal points you claim, per this buyer's table",
     "help": "Each buyer sets its own specific-goal categories (B-BBEE level, ownership, locality...). "
             "Written on the completion guide; the points table itself is ticked by hand."},
    # --- SBD/MBD 8 -------------------------------------------------------
    {"key": "sbd8.restricted_supplier", "form": "8", "group": "Past SCM practices (SBD 8)", "type": "yes_no",
     "label": "Is the bidder or any director listed on National Treasury's Database of Restricted Suppliers?",
     "help": "SBD 8 4.1."},
    {"key": "sbd8.tender_defaulter", "form": "8", "group": "Past SCM practices (SBD 8)", "type": "yes_no",
     "label": "Is the bidder or any director listed on the Register for Tender Defaulters?",
     "help": "SBD 8 4.2."},
    {"key": "sbd8.convicted", "form": "8", "group": "Past SCM practices (SBD 8)", "type": "yes_no",
     "label": "Was the bidder or any director convicted of fraud or corruption in the past five years?",
     "help": "SBD 8 4.3."},
    {"key": "sbd8.contract_terminated", "form": "8", "group": "Past SCM practices (SBD 8)", "type": "yes_no",
     "label": "Was any contract with an organ of state terminated in the past five years for failure to perform?",
     "help": "SBD 8 4.4."},
    # --- conditional ---------------------------------------------------------
    {"key": "bidding_as_jv", "condition": "jv", "group": "Bid structure", "type": "yes_no",
     "label": "Are you bidding as a joint venture / consortium?",
     "help": "This tender mentions a JV agreement. If yes, upload the signed agreement below."},
    {"key": "briefing_attended", "condition": "briefing", "group": "Briefing", "type": "yes_no",
     "label": "Did you attend the compulsory briefing / site meeting?",
     "help": "Missing a compulsory briefing usually makes the bid non-responsive. Upload the attendance certificate if one was issued."},
]

# Company facts SBD 1 asks for that 0001's companies table had no column for
# (added by 0017). Asked once on the workspace, saved to the company.
COMPANY_FIELDS = [
    {"key": "company.telephone", "column": "telephone", "form": "1", "label": "Company telephone number",
     "help": "SBD 1 'TELEPHONE NUMBER' (bidder section)."},
    {"key": "company.cellphone", "column": "cellphone", "form": "1", "label": "Company cellphone number", "optional": True,
     "help": "SBD 1 'CELLPHONE NUMBER'."},
    {"key": "company.email", "column": "email", "form": "1", "label": "Company e-mail address",
     "help": "SBD 1 'E-MAIL ADDRESS' (bidder section)."},
    {"key": "company.cidb_grading", "column": "cidb_grading", "condition": "cidb", "label": "CIDB grading (e.g. 6CE)",
     "help": "This tender names a CIDB grading. Leave blank if you have none."},
]


# ---------------------------------------------------------------------------
# Manifest
# ---------------------------------------------------------------------------

def extract_requirements(paths, work_dir=None, source_label=None, keep_text=False, past_corrections=None, buyer=None):
    """paths: files and/or directories (zips are unpacked). Returns the manifest
    dict (JSON-safe). Document dicts keep local paths under `_local` so the pack
    generator can reuse the downloaded copies; strip_local() removes them before
    the manifest is stored.

    past_corrections: up to a handful of past platform-admin verdicts from
    requirement_extraction_reviews (migration 0074), each shaped like
    {"item_type": ..., "item_snapshot": {...}, "verdict": "corrected" |
    "false_positive", "corrected_value": ...}. Optional, and never fetched by
    this function itself (it stays a pure function of its inputs - no
    database access here) - see the module comment above
    _apply_past_corrections_to_manifest for exactly what this can and can't
    do.

    buyer: the tender's own buyer/client name (e.g. "Transnet SOC Ltd"), if
    known - optional, and only ever used to tell find_roles() apart a
    "<buyer name> Contract/Project Manager" mention (the buyer's own agent)
    from a genuine bidder-role ask (Finding 47). Never fetched here; the
    caller already has it (process_general_jobs.py reads it straight off
    the tender row being processed)."""
    own_tmp = work_dir is None
    work_dir = work_dir or tempfile.mkdtemp(prefix="tt_req_")
    docs = [read_document(d, work_dir) for d in collect_documents(paths, work_dir)]
    for d in docs:
        d["is_bulletin"] = (not d.get("unreadable")) and is_combined_bulletin(d.get("text"))
    # Finding 42: a bulletin's facts belong to whichever tender its own row is
    # for, never knowably this one - excluded from every signal below (suite,
    # forms, returnables, roles, CIDB, briefing, locality, experience) the
    # same way an unreadable document already is, not just hidden from the UI.
    usable_docs = [d for d in docs if not d["is_bulletin"]]
    all_text = "\n".join(d.get("text") or "" for d in usable_docs)

    mentioned = {}
    for d in usable_docs:
        for k, n in forms_in_text(d.get("text") or "").items():
            mentioned[k] = mentioned.get(k, 0) + n
        for k, n in forms_in_text(os.path.basename(d["name"])).items():
            mentioned[k] = mentioned.get(k, 0) + n

    ocr_docs = {d["name"] for d in docs if d.get("ocr")}
    ocr_mean = _ocr_mean(docs)
    doc_entries = []
    for d in docs:
        kind = "combined_bulletin" if d["is_bulletin"] else classify_document(d, mentioned)
        editable = d["format"] == "docx" or bool(d.get("converted_docx"))
        doc_entries.append({
            "name": d["name"], "format": d["format"], "kind": kind,
            "forms": d.get("edition_of", []), "editable": editable and kind in ("standard_form", "tender_document"),
            "converted_from_doc": bool(d.get("converted_docx")),
            "unreadable": d.get("unreadable"),
            "chars": len(d.get("text") or ""),
            # EXPERIMENTAL: set when some/all of the text came from OCR of a
            # scanned page (engine/ocr/scan_ocr.py) - pages, confidence, notice.
            "ocr": d.get("ocr"),
            "_local": d.get("converted_docx") or d["path"],
        })

    # A single stray cross-reference to the other suite (a municipal buyer's
    # boilerplate note pointing bidders at "the National Treasury SBD 4
    # guideline" while every real form in the bundle is MBD, the OCPO-style
    # glossary annexure) used to flip the whole tender's suite on mere
    # presence - `if "MBD" in suites and "SBD" not in suites` required SBD
    # to be COMPLETELY ABSENT for MBD to win, so one single, real SBD
    # mention anywhere outweighed ten real MBD mentions. Reproduced directly
    # with a synthetic MBD bundle (MBD 1/4/6.1/7/8/9, 10 real mentions) plus
    # one boilerplate "SBD 4 guideline" reference in an annexure - every form
    # label in the manifest came back "SBD 1", "SBD 4", etc. for a genuinely
    # municipal tender (QA-Edge-Case-Findings.md, 2026-10-20). Majority of
    # mentions decides now; a tie still defaults to SBD, same as before.
    mbd_total = sum(n for (s, _), n in mentioned.items() if s == "MBD")
    sbd_total = sum(n for (s, _), n in mentioned.items() if s == "SBD")
    if mbd_total > sbd_total:
        suite = "MBD"
    elif sbd_total > 0:
        suite = "SBD"
    else:
        suite = None

    forms = []
    # Finding 35 (2026-10-20): a stray cross-reference to the LOSING suite's
    # own form code (e.g. one boilerplate "this tender does not use SBD 9"
    # sentence in an MBD-dominant bundle) used to still list that code as
    # required - `{c for (_, c) in mentioned}` unioned mentions across BOTH
    # suites with no suite-consistency check, the same root-cause class
    # Finding 29 already fixed for the suite label itself. Only a mention
    # under the WINNING suite's own prefix counts now; a genuine document
    # EDITION found for a code (doc_entries' own "forms") is a different,
    # suite-independent signal (a real file exists) and stays unconditional.
    codes = sorted({c for (s, c) in mentioned if s == suite} | {c for e in doc_entries for c in e["forms"]},
                   key=lambda c: float(c))
    for code in codes:
        title, capability = FORMS[code]
        editions = [e for e in doc_entries if code in e["forms"]]
        fillable = [e for e in editions if e["editable"]]
        prefix = suite or "SBD"
        n_mentions = mentioned.get((prefix, code), 0) + mentioned.get(("SBD" if prefix == "MBD" else "MBD", code), 0)
        quote, quote_doc = None, None
        for d in usable_docs:
            m = re.search(rf"\b{prefix[0]}\s?B\s?D\s*{re.escape(code).replace(r'\.', r'\s*[.(]\s*')}", d.get("text") or "", re.IGNORECASE)
            if m:
                quote, quote_doc = _quote(d["text"], m.start(), m.end(), 80), d["name"]
                break
        overlay = _overlay_edition(code, editions, usable_docs) if capability == "fill" and not fillable else None
        if capability == "fill" and fillable:
            status, reason = "auto", f"Filled from your data in {fillable[0]['name']}" + (
                " (converted from .doc by LibreOffice - check the layout)" if fillable[0]["converted_from_doc"] else "")
        elif overlay:
            status, reason = "overlay", (
                f"Flat PDF ({overlay['document']}, p.{overlay['page']}) - Tender Tool types your details onto the printed form "
                "next to each label it can find; anything it can't place safely stays on the completion guide to write by hand.")
        elif capability == "fill" and editions and editions[0]["name"] in ocr_docs:
            status, reason = "by_hand", (f"Only a scanned PDF edition was published ({editions[0]['name']}, read by OCR) - "
                                         "nothing can be typed onto a scan, so it's completed by hand from the completion guide.")
        elif capability == "fill" and editions:
            status, reason = "by_hand", (f"Only a {editions[0]['format'].upper()} edition was published ({editions[0]['name']}) - "
                                         "flat PDFs have no fillable fields, so it's completed by hand from the completion guide.")
        elif capability == "fill":
            status, reason = "by_hand", ("Named in the tender documents but no copy of the form was found in them - "
                                         "it's usually inside the main bundle; complete it by hand from the completion guide.")
        elif capability == "attach":
            status, reason = "attach", "Satisfied by the tax compliance status (TCS PIN) - attached from the vault."
        else:
            status, reason = "by_hand", {
                "3": "Pricing is tender-specific - Tender Tool never fills prices.",
                "5": "National Industrial Participation applies only to large imported content - complete by hand if it applies.",
                "6.2": "Local-content declaration needs this bid's own product data - complete by hand.",
                "7": "Contract form - signed by hand (some buyers want it with the bid, most only after award).",
            }.get(code, "Complete by hand.")
        confidence, confidence_reason = score_confidence(
            "form_edition" if editions else "form_code", n_mentions,
            ocr_mean.get(quote_doc) if quote_doc in ocr_docs else None)
        forms.append({"code": code, "label": f"{prefix} {code}", "title": title, "status": status, "reason": reason,
                      **({"overlay": overlay} if overlay else {}),
                      "mentions": n_mentions, "quote": quote,
                      **({"from_ocr": True} if quote_doc in ocr_docs else {}),
                      "confidence": confidence, "confidence_reason": confidence_reason,
                      "editions": [{"document": e["name"], "format": e["format"], "editable": e["editable"]} for e in editions]})

    buyer_forms = []
    for e in doc_entries:
        if e["kind"] == "buyer_form":
            buyer_forms.append({"key": "doc:" + os.path.basename(e["name"]), "label": os.path.basename(e["name"]),
                                "document": e["name"], "format": e["format"],
                                "reason": "Buyer-specific form - no Tender Tool template for it; complete by hand and attach."})
    for key, label, pat in IN_TEXT_BUYER_FORMS:
        for d in usable_docs:
            m = re.search(pat, d.get("text") or "", re.IGNORECASE)
            if m:
                confidence, confidence_reason = score_confidence("form_code", 1, ocr_mean.get(d["name"]))
                entry = {"key": key, "label": label, "document": d["name"], "format": d["format"],
                         "quote": _quote(d["text"], m.start(), m.end(), 80),
                         "confidence": confidence, "confidence_reason": confidence_reason,
                         "reason": "Buyer-specific form inside the tender document - complete by hand."}
                ov = _overlay_buyer_form(key, d)
                if ov:
                    entry["overlay"] = ov
                    entry["reason"] = (f"Found on p.{ov['page']} of {d['name']} - Tender Tool types your company details "
                                       f"({ov['fills']}) onto it; {ov['by_hand']} stay by hand.")
                buyer_forms.append(entry)
                break

    cidb_m = CIDB_RE.search(all_text)
    # Finding, 2026-10-24: CIDB_RE.search() only ever returns the FIRST
    # mention in the whole document set, so an addendum correcting the
    # grading requirement (confirmed real and reachable - exactly the same
    # shape as Finding 61's SANRAL closing-time/delivery-address ambiguity,
    # just for CIDB grading) is silently outvoted by the ORIGINAL, now-stale
    # value: a 2-document reproduction (original notice "6CE", addendum
    # "the CIDB grading requirement for this contract is corrected to
    # 8CE") picked up "6CE", never "8CE". Submitting against the wrong
    # minimum grading is exactly the kind of compliance-critical silent
    # error this project exists to prevent. `cidb_ambiguous` flags it
    # instead of guessing which one is right - no reliable "the later
    # mention wins" signal exists here either (plain text concatenation
    # order, not a dated revision history).
    cidb_gradings = {m.group(1).replace(" ", "").upper() for m in CIDB_RE.finditer(all_text)}
    cidb_ambiguous = len(cidb_gradings) > 1
    brief_m = BRIEFING_RE.search(all_text)
    non_compulsory = bool(re.search(r"non[- ]compulsory\s+(?:clarification\s+)?brief", all_text, re.IGNORECASE))
    outlook = estimate_score_outlook(all_text, "full_document" if all_text.strip() else "no_extractable_text",
                                     {"disciplines": [], "bee_level": None, "offices": []})
    loc = extract_locality(all_text) if all_text.strip() else {"rows": [], "mandatory_local_office": None}
    cidb_confidence = score_confidence("anchored_regex", 1) if cidb_m else (None, None)
    briefing_confidence = score_confidence("anchored_regex", 1) if brief_m else (None, None)
    returnables = find_returnables(usable_docs, ocr_mean)
    roles = find_roles(usable_docs, ocr_mean, buyer=buyer)
    person_registrations = find_person_registrations(usable_docs, ocr_mean)
    # EXPERIMENTAL OCR: anything read out of OCR'd text says so, so the UI
    # can show it isn't from a real text layer.
    for item in returnables + roles + buyer_forms + person_registrations:
        if item.get("document") in ocr_docs:
            item["from_ocr"] = True
    conditions = {
        "cidb": bool(cidb_m) or any(r["key"] == "cidb" and r["strength"] == "required" for r in returnables),
        "briefing": bool(brief_m) and not non_compulsory,
        "jv": any(r["key"] == "jv_agreement" for r in returnables),
    }
    required_form_codes = {f["code"] for f in forms}
    questions = [q for q in QUESTIONS
                 if (q.get("form") and q["form"] in required_form_codes) or (q.get("condition") and conditions.get(q["condition"]))]
    company_fields = [c for c in COMPANY_FIELDS
                      if (c.get("form") and c["form"] in required_form_codes) or (c.get("condition") and conditions.get(c["condition"]))]

    manifest = {
        "version": MANIFEST_VERSION,
        "extracted_at": datetime.datetime.now(datetime.timezone.utc).isoformat(timespec="seconds"),
        "source": source_label,
        "form_suite": suite,
        "documents": doc_entries,
        "forms": forms,
        "buyer_forms": buyer_forms,
        "returnables": returnables,
        "roles": roles,
        # 0040: only present when the tender asks (People page shows the
        # SACNASP / SAGC fields because of it)
        **({"person_registrations": person_registrations} if person_registrations else {}),
        # cidb/briefing are read off all_text (every document concatenated),
        # so no single source document can be attributed for an OCR discount
        # - score_confidence() still applies the anchored_regex base, just
        # with ocr_confidence left unknown rather than guessed.
        "cidb": {"grading": cidb_m.group(1).replace(" ", "").upper(), "quote": _quote(all_text, cidb_m.start(), cidb_m.end()),
                 "confidence": cidb_confidence[0], "confidence_reason": cidb_confidence[1],
                 "ambiguous": cidb_ambiguous} if cidb_m else None,
        "briefing": {"compulsory": conditions["briefing"], "quote": _quote(all_text, brief_m.start(), brief_m.end()),
                     "confidence": briefing_confidence[0], "confidence_reason": briefing_confidence[1]} if brief_m else None,
        "scoring": {
            "preference_system": outlook["preference_system"],
            "functionality": outlook["functionality"],
            "specific_goal_bases": outlook["specific_goals"].get("bases", []),
        },
        "experience_evidence": extract_experience_evidence(all_text) if all_text.strip() else {"source": "not_found", "rules": []},
        "locality": {"rows": loc["rows"], "mandatory_local_office": loc["mandatory_local_office"]},
        # the tender's own C3.1.11(d)-style key-person minimums table (audit
        # Q5) - a manually loaded SANRAL tender's Book 2/3 carries it
        "key_person_minimums": extract_key_person_minimums(all_text) if all_text.strip()
        else {"source": "not_found", "clause": None, "rows": [], "unparsed": [], "reason": "No document text."},
        "questions": questions,
        "company_fields": company_fields,
        "gaps": _gaps(doc_entries, forms, suite, buyer_forms, _depth(all_text)) + _key_personnel_gap(usable_docs),
        "document_depth": _depth(all_text),
    }
    _apply_past_corrections_to_manifest(manifest, past_corrections)
    if own_tmp:
        manifest["_work_dir"] = work_dir
    if keep_text:
        # for the intake step (engine/intake/classify_tender.py): fit scoring
        # reads the same text instead of extracting every document twice.
        manifest["_all_text"] = all_text
    return manifest


# Flat PDF editions: engine/general/fill_flat_pdf.py types onto the printed
# form, anchored to each label in the PDF's own text layer (findings doc,
# "Round 4": fixed x/y positions differ per municipality, the label wording
# doesn't). Here we only check that the form's page is really in the document,
# so the manifest can promise it; what actually got placed is reported per
# field when the pack is built.
OVERLAY_BUYER_FORMS = {
    "compulsory_enterprise_questionnaire": ("ceq", "name, VAT, CSD, company registration and tax numbers, contact details",
                                            "the CIDB number, partner/principal particulars, banking account number and the signature"),
    "form_of_offer": ("form_of_offer", "signatory name, capacity, company name, address and telephone",
                      "the offered price, the signature and the witnesses"),
}


def _overlay_pages(doc, code):
    try:
        import fill_flat_pdf
    except ImportError:  # pdfplumber/pypdf missing on this host - by hand, as before
        return []
    pages = fill_flat_pdf.find_form_pages((doc.get("text") or "").split("\f"), code)
    # The overlay places values next to labels in the PDF's own text layer; a
    # page read by OCR has none, so its form stays by hand.
    ocr_pages = {n - 1 for n in ((doc.get("ocr") or {}).get("pages_ocr") or [])}
    return [p for p in pages if p not in ocr_pages]


def _overlay_edition(code, editions, docs):
    if code not in ("1", "4", "9"):
        return None
    for e in editions:
        d = next((x for x in docs if x["name"] == e["name"]), None)
        if d and d["format"] == "pdf" and not d.get("unreadable"):
            pages = _overlay_pages(d, code)
            if pages:
                return {"document": d["name"], "page": pages[0] + 1}
    return None


def _overlay_buyer_form(key, doc):
    if key not in OVERLAY_BUYER_FORMS or doc["format"] != "pdf" or doc.get("unreadable"):
        return None
    code, fills, by_hand = OVERLAY_BUYER_FORMS[key]
    pages = _overlay_pages(doc, code)
    return {"code": code, "document": doc["name"], "page": pages[0] + 1, "fills": fills, "by_hand": by_hand} if pages else None


def _depth(all_text):
    """assess_tender.document_depth() - the same advert-vs-full-document test
    Discovery uses - on everything that was read."""
    from assess_tender import document_depth
    return document_depth(all_text or "")


def _gaps(doc_entries, forms, suite, buyer_forms, depth="full_document"):
    """Plain statements of what this engine can't do for this tender - shown
    in the UI and the pack index instead of being hidden."""
    gaps = []
    bulletin = [e["name"] for e in doc_entries if e["kind"] == "combined_bulletin"]
    if bulletin:
        # Finding 42: this document's own facts (closing date, CIDB grading,
        # preference system) belong to whichever tender its own row is for,
        # never knowably this one - excluded from every signal above, so say
        # so plainly rather than let the generic "couldn't be read" (implies
        # a parsing failure, not true here) or "no SBD/MBD forms" (implies
        # this document was actually checked) wording stand in for it.
        gaps.append(
            "Covers multiple tenders in one document, not just this one: " + "; ".join(bulletin[:5]) +
            ". Tender Tool hasn't tried to work out which part is this tender's own - nothing from it is shown "
            "below. Find this tender's own row/section in it and read it by eye, or get its own separate document "
            "from the buyer if one exists.")
    non_bulletin = [e for e in doc_entries if e["kind"] != "combined_bulletin"]
    readable = [e for e in non_bulletin if e["chars"] > 0 and not e["unreadable"]]
    if not non_bulletin:
        pass  # every document was a bulletin - already said so above, nothing more to add
    elif not readable:
        # Nothing was read, so nothing can be concluded about the forms - the
        # old wording ("uses no SBD/MBD standard forms") was a confident
        # statement made from zero evidence.
        gaps.append("None of this tender's documents could be read, so what it requires is unknown - nothing below "
                    "was checked against its documents. Upload a text PDF or Word copy.")
    elif suite is None and any(e["unreadable"] for e in non_bulletin):
        gaps.append("No SBD/MBD standard forms were found in the documents that could be read - but some couldn't be "
                    "read (listed below), and the forms may be in those.")
    elif suite is None and depth != "full_document":
        # 2026-10-17 scenario pass: Stellenbosch BSM 25/27 published only its
        # newspaper advert on eTenders ("...Adverts only.pdf"); the manifest
        # used to say "uses no SBD/MBD standard forms" and the pack built as
        # if that were the whole tender.
        gaps.append("Only an advert or short notice was found, not the full tender document - so whether it uses "
                    "SBD/MBD forms (and everything else it requires) is unknown. Get the full bid document from the buyer "
                    "and load it (Tenders > Load a tender manually, or Add documents).")
    elif suite is None:
        gaps.append("This tender uses no SBD/MBD standard forms - its returnable forms are the buyer's own, which "
                    "Tender Tool has no templates for. The pack orders your supporting documents; the buyer's forms are "
                    "completed by hand.")
    for f in forms:
        if f["status"] == "by_hand" and FORMS[f["code"]][1] == "fill":
            gaps.append(f"{f['label']} must be completed by hand: {f['reason']}")
        elif f["status"] == "overlay":
            gaps.append(f"{f['label']} is a flat PDF: your details are typed onto it, but check the placement before you sign "
                        "(fields Tender Tool couldn't place are listed in the pack's completion guide).")
    unread = [e["name"] for e in doc_entries if e["kind"] == "unreadable"]
    if unread:
        gaps.append("Couldn't read: " + "; ".join(unread[:5]) + " - check these documents by eye.")
    gaps.extend(_ocr_gaps(doc_entries))
    return gaps


def _ocr_gaps(doc_entries):
    """EXPERIMENTAL OCR: one plain line per document whose text (some of it or
    all) came from OCR, so nobody mistakes it for a clean read."""
    out = []
    for e in doc_entries:
        o = e.get("ocr")
        if not o:
            continue
        pages = o["pages_ocr"]
        where = (f"all {len(pages)} page(s)" if o.get("total_pages") and len(pages) == o["total_pages"]
                 else "page(s) " + ", ".join(map(str, pages)))
        line = (f"{e['name']}: {where} read by OCR from a scan (no text layer), average OCR confidence "
                f"{o['mean_confidence']:.0f}/100. OCR can misread characters and numbers and can drop a line - check "
                "anything listed from this document against the scan. Flat-PDF typing is not offered on OCR'd pages.")
        if o.get("pages_failed"):
            line += " Not read: " + "; ".join(f"p.{f['page']} ({f['reason']})" for f in o["pages_failed"][:4]) + "."
        if o.get("pages_skipped"):
            line += f" Pages {o['pages_skipped'][0]}-{o['pages_skipped'][-1]} were not OCR'd (page limit) - read them by eye."
        out.append(line)
    return out


# A tender that asks for key personnel by any role. find_roles() only
# recognises the catalogued roles (ROLE_PATTERNS - the ones met in the sample
# tenders), so a role outside that list (an environmental control officer, a
# land surveyor, an electrician...) would otherwise just not appear, and the
# Roles card would read as if it were the whole requirement.
KEY_PERSONNEL_RE = re.compile(r"key\s+(?:personnel|staff|persons?)\b|curriculum\s+vitae|\bCVs?\s+(?:of|for)\b", re.IGNORECASE)
# A real find (2026-10-21, phase-2 sampling): "5. Change Key Personnel on the
# Service" was a false positive - one item in a list of things the SERVICE
# PROVIDER may not do without the buyer's written consent DURING THE
# CONTRACT ("deviate from the programme", "deviate from the scope", "change
# key personnel") - a contract-administration clause, not a bid-time ask for
# key personnel. The whole-document text this came from genuinely has no
# other key-personnel mention at all, so returning it as "this tender asks
# for key personnel" was actively wrong, not just imprecise.
_KEY_PERSONNEL_FALSE_CONTEXT_RE = re.compile(r"\b(?:change|changing|substitut\w*|replac\w*)\s*$", re.IGNORECASE)


def _key_personnel_gap(docs):
    for d in docs:
        t = d.get("text") or ""
        for m in KEY_PERSONNEL_RE.finditer(t):
            if _KEY_PERSONNEL_FALSE_CONTEXT_RE.search(t[max(0, m.start() - 20):m.start()]):
                continue  # "change/substitute/replace key personnel" - contract administration, not a bid-time ask
            return [f"This tender asks for key personnel (\"{_quote(t, m.start(), m.end())[:160]}\", {d['name']}). "
                    f"Tender Tool recognises {len(ROLE_PATTERNS)} common roles (listed under Roles when found); "
                    "check the tender's own list for any role it doesn't show."]
    return []


def strip_local(manifest):
    m = json.loads(json.dumps(manifest))
    m.pop("_work_dir", None)
    m.pop("_all_text", None)
    for d in m.get("documents", []):
        d.pop("_local", None)
    return m


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("paths", nargs="+")
    ap.add_argument("--json", help="write the manifest here")
    args = ap.parse_args()
    m = extract_requirements(args.paths)
    clean = strip_local(m)
    if args.json:
        json.dump(clean, open(args.json, "w"), indent=2)
    print(f"form suite: {clean['form_suite']}")
    for d in clean["documents"]:
        print(f"  doc [{d['kind']}] {d['name']} ({d['format']}{', editable' if d['editable'] else ''}) forms={d['forms']}"
              + (f" UNREADABLE: {d['unreadable']}" if d["unreadable"] else ""))
    for f in clean["forms"]:
        print(f"  form {f['label']:8} {f['status']:8} x{f['mentions']} - {f['reason'][:110]}")
    for b in clean["buyer_forms"]:
        print(f"  buyer form: {b['label']}")
    for r in clean["returnables"]:
        print(f"  returnable {r['key']:22} {r['strength']:9} x{r['mentions']}: {r['quote'][:100]}")
    for r in clean["roles"]:
        print(f"  role {r['role_code']:38} x{r['mentions']}: {r['quote'][:90]}")
    print(f"  cidb: {clean['cidb'] and clean['cidb']['grading']}  briefing: {clean['briefing'] and clean['briefing']['compulsory']}")
    print(f"  pref: {clean['scoring']['preference_system']['system']}  functionality: {clean['scoring']['functionality'].get('threshold')}"
          f"  evidence rules: {len(clean['experience_evidence']['rules'])}  locality rows: {len(clean['locality']['rows'])}")
    for b in clean["buyer_forms"]:
        if b.get("overlay"):
            print(f"  buyer form typed onto: {b['key']} p.{b['overlay']['page']}")
    print(f"  questions: {[q['key'] for q in clean['questions']]}")
    print(f"  company fields: {[c['key'] for c in clean['company_fields']]}")
    for g in clean["gaps"]:
        print(f"  GAP: {g[:160]}")
    if m.get("_work_dir"):
        shutil.rmtree(m["_work_dir"], ignore_errors=True)


if __name__ == "__main__":
    main()
