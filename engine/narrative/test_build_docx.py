import os
import sys
import tempfile

import docx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from build_docx import COLOR_THEMES, build_docx

FAILS = []


def check(name, cond, detail=""):
    print(("  ok   " if cond else "  FAIL ") + name + (f"  ({detail})" if detail and not cond else ""))
    if not cond:
        FAILS.append(name)


SECTIONS = [
    {"heading": "Introduction", "kind": "prose", "paragraphs": ["We are Mahlangu Civils.", "We tender often."], "bullets": []},
    {"heading": "Key People", "kind": "bullet_list", "paragraphs": [], "bullets": ["Thabo Mahlangu, Site Agent", "Jane Doe, QA Manager"]},
]


def _build(tmp_name, **kwargs):
    with tempfile.TemporaryDirectory() as tmp:
        out_path = os.path.join(tmp, tmp_name)
        build_docx(SECTIONS, out_path, title="Methodology Statement", **kwargs)
        return docx.Document(out_path)


def test_builds_a_real_readable_docx():
    doc = _build("out.docx")
    all_text = "\n".join(p.text for p in doc.paragraphs)
    check("title text is present", "Methodology Statement" in all_text)
    check("prose section heading is present", "Introduction" in all_text)
    check("prose paragraph text is present", "We are Mahlangu Civils." in all_text)
    check("bullet section heading is present", "Key People" in all_text)
    check("bullet text is present", "Thabo Mahlangu, Site Agent" in all_text)


def test_bullets_use_the_list_bullet_style_not_plain_prose():
    doc = _build("out.docx")
    bullet_paragraphs = [p for p in doc.paragraphs if p.text == "Thabo Mahlangu, Site Agent"]
    check("the bullet paragraph exists", len(bullet_paragraphs) == 1)
    check("it uses the List Bullet style", bullet_paragraphs[0].style.name == "List Bullet",
          bullet_paragraphs[0].style.name)


def test_subtitle_is_included_and_italicized_when_given():
    doc = _build("out.docx", subtitle="N.004-010-2026/1F - SANRAL")
    subtitle_paragraphs = [p for p in doc.paragraphs if "N.004-010-2026/1F" in p.text]
    check("the subtitle text is present", len(subtitle_paragraphs) == 1)
    check("the subtitle run is italicized", subtitle_paragraphs[0].runs[0].italic is True)


def test_color_theme_is_actually_applied_to_the_heading_style():
    for theme_name, expected_rgb in COLOR_THEMES.items():
        doc = _build("out.docx", color_theme=theme_name)
        heading_style = doc.styles["Heading 1"]
        check(f"heading color matches the '{theme_name}' theme",
              heading_style.font.color.rgb == expected_rgb,
              (heading_style.font.color.rgb, expected_rgb))


def test_unknown_color_theme_falls_back_to_default_rather_than_crashing():
    doc = _build("out.docx", color_theme="not-a-real-theme")
    heading_style = doc.styles["Heading 1"]
    check("an unknown theme name falls back to the default color, not an error",
          heading_style.font.color.rgb == COLOR_THEMES["navy"])


def test_visual_rich_adds_a_rule_plain_does_not():
    visual_rich_doc = _build("out.docx", visual_style="visual_rich")
    plain_doc = _build("out.docx", visual_style="plain")
    rich_text = "\n".join(p.text for p in visual_rich_doc.paragraphs)
    plain_text = "\n".join(p.text for p in plain_doc.paragraphs)
    check("visual_rich includes the decorative rule", "━" in rich_text)
    check("plain does not include the decorative rule", "━" not in plain_text)


def test_control_characters_are_stripped_not_crashed_on():
    print("a control character / NULL byte anywhere in the document (title, subtitle, "
          "heading, paragraph or bullet) must not crash the builder")
    # Finding (2026-10-23): python-docx (via lxml) raises ValueError on a
    # control character/NULL byte - the sibling generator (engine/cv/
    # generate_cv.py's own _safe_text()) already guards every string it
    # writes for exactly this; build_docx.py had no such guard at all.
    # subtitle is a direct, un-sanitized passthrough of the tender's own
    # contract_no (read off a real buyer's page/portal - this codebase has
    # already found OCR/scrape corruption in fields exactly like this one),
    # so this crashes deterministically from real data, not only from a
    # hypothetical LLM echo of a company's own free-text "how we work".
    sections = [
        {"heading": "Our Approach\x01", "kind": "prose", "paragraphs": ["We build \x00 roads safely."], "bullets": []},
        {"heading": "Key People", "kind": "bullet_list", "paragraphs": [], "bullets": ["Jane \x1f Doe"]},
    ]
    with tempfile.TemporaryDirectory() as tmp:
        out_path = os.path.join(tmp, "out.docx")
        try:
            build_docx(sections, out_path, title="Cover\x02 Letter", subtitle="N.011-010\x00-2027/1F")
        except Exception as e:
            check("build_docx does not crash on a control character anywhere in its input", False, str(e))
            return
        doc = docx.Document(out_path)
        all_text = "\n".join(p.text for p in doc.paragraphs)
        check("build_docx does not crash on a control character anywhere in its input", True)
        check("the title is stripped of its control character, not dropped entirely", "Cover Letter" in all_text)
        check("the subtitle (a direct passthrough of contract_no) is stripped, not dropped entirely",
              "N.011-010-2027/1F" in all_text)
        check("a heading's control character is stripped", "Our Approach" in all_text)
        check("a paragraph's control character is stripped", "We build  roads safely." in all_text)
        check("a bullet's control character is stripped", "Jane  Doe" in all_text)


def main():
    test_builds_a_real_readable_docx()
    test_bullets_use_the_list_bullet_style_not_plain_prose()
    test_subtitle_is_included_and_italicized_when_given()
    test_color_theme_is_actually_applied_to_the_heading_style()
    test_unknown_color_theme_falls_back_to_default_rather_than_crashing()
    test_visual_rich_adds_a_rule_plain_does_not()
    test_control_characters_are_stripped_not_crashed_on()
    print(f"\n{len(FAILS)} failure(s)")
    sys.exit(1 if FAILS else 0)


if __name__ == "__main__":
    main()
