"""
Focused test for fill_docx_forms.py's C.1.2.3 key-person row handling - added
during this session's enterprise-readiness audit (Sarel, 2026-10-20: "make
sure it is designed and build to catch all ... scenarios"). No dedicated
test file existed for this module before this.

Finding: `_row_code(label)` returns None for two different reasons - a
genuinely optional "Alternate to ... (if applicable)" row, and a real
key-position row whose title simply doesn't match anything in
POSITION_TITLES (e.g. a future SANRAL tender using a title this project's
own dict doesn't recognise yet). fill_contract_data() used to treat both
identically - silently leaving the row blank with nothing in the
provenance report telling a human to fill it by hand. This would hit a
brand-new company's very first tender with no warning at all, since
admin_harness.py's own baseline-diff treats a first run as "clean,
baseline stored" unconditionally.

Run: python3 engine/test_fill_docx_forms.py
"""
import os
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import docx  # noqa: E402
import fill_docx_forms as fdf  # noqa: E402
from provenance import ProvenanceLog  # noqa: E402

FAILS = []


def check(name, cond, detail=""):
    print(("  ok   " if cond else "  FAIL ") + name + (f"  ({detail})" if detail and not cond else ""))
    if not cond:
        FAILS.append(name)


def _build_template(src_dir):
    """Minimal synthetic C.1.2.3_Contract_Data.docx: just enough structure
    for fill_contract_data() to run without erroring (3 tables - company
    info, key-persons, directors), not a full real-looking form."""
    d = docx.Document()
    t0 = d.add_table(rows=1, cols=1)
    t0.rows[0].cells[0].paragraphs[0].add_run("placeholder")

    t1 = d.add_table(rows=4, cols=2)
    t1.rows[0].cells[0].text = "Position"
    t1.rows[1].cells[0].text = "Project Leader"
    t1.rows[2].cells[0].text = "Alternate to Project Leader (if applicable)"
    t1.rows[3].cells[0].text = "Quantum Flux Coordinator"  # a bogus, unrecognized title

    d.add_table(rows=2, cols=4)  # directors table - header + one blank row

    d.save(os.path.join(src_dir, "C.1.2.3_Contract_Data.docx"))


def test_unrecognized_key_position_title_is_logged_not_silently_skipped():
    print("fill_contract_data: an unrecognized key-position title vs a genuinely optional 'Alternate' row")
    with tempfile.TemporaryDirectory() as src_dir, tempfile.TemporaryDirectory() as out_dir:
        _build_template(src_dir)
        comp = {"legal_name": "Test Co", "registered_address": "1 Test St"}
        signatory = {"name": "Jane Doe", "signatory_capacity": "Director"}
        log = ProvenanceLog(source="test")
        fdf.fill_contract_data(src_dir, out_dir, comp, signatory, key_persons=[], log=log)

        missing = {e["field_key"]: e["reason"] for e in log.entries if e["status"] == "missing"}
        check("a recognized title with no confirmed person is still logged missing (unchanged baseline)",
              "c123.key_person_row_1" in missing, missing)
        check("a genuinely optional 'Alternate to ...' row stays silently skipped - no log entry",
              "c123.key_person_row_2" not in missing, missing)
        check("an unrecognized title is NOW logged missing, with a clear, actionable reason (the fix)",
              "c123.key_person_row_3" in missing
              and "doesn't match any known SANRAL key-position title" in missing.get("c123.key_person_row_3", ""),
              missing)


def main():
    test_unrecognized_key_position_title_is_logged_not_silently_skipped()
    print(f"\n{len(FAILS)} failure(s)")
    sys.exit(1 if FAILS else 0)


if __name__ == "__main__":
    main()
